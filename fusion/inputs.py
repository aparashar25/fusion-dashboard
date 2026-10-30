"""All deal inputs (the blue cells on F1 Fusion Assumptions), plus an Excel loader."""
from __future__ import annotations
from dataclasses import dataclass, field, asdict
import datetime as dt
import json

TERMS_N = 4
AGES = 61          # loan ages 0..60 (F2 columns C:BK)
MONTHS = 120       # projection months 1..120 (F3 columns C:DR)
PLAN_MONTHS = 60   # purchase plan months 1..60 (F1 columns B:BI)


@dataclass
class Fee:
    active: bool = False
    name: str = ""
    paid_to: str = ""
    type: str = "Upfront"          # Upfront | Across life
    method: str = "%"              # % | Fixed $
    basis: str = "Net paid"        # Upfront: Net paid | Face ; Across life: Collections | Beginning balance | Ending balance | Average balance
    rate: float | None = 0.0
    fixed: float | None = None     # $ per loan (upfront) or $ per active loan per month (across life)


@dataclass
class Tier:
    active: bool = False
    investor_pct: float | None = None
    moic_target: float | None = None
    irr_target: float | None = None


@dataclass
class Inputs:
    # 1. Switches
    recycling: bool = True
    leverage: bool = True
    losses: bool = True
    prepay: bool = True
    payment_stop: bool = True
    promote: bool = True
    target_mode: str = "MOIC"            # MOIC | IRR
    outlay_basis: str = "Price + fee"    # Price + fee | Price only
    after_last_tier: str = "Continue split"   # Continue split | 100% to school
    capital_basis: str = "OCM %"         # OCM % | Fixed $
    # 2. Loans & credit (index 0..3 = 24/36/48/60 columns)
    terms: list = field(default_factory=lambda: [24, 36, 48, 60])
    apr: list = field(default_factory=lambda: [0.1495, 0.1695, 0.1795, 0.1795])
    mix: list = field(default_factory=lambda: [0.25, 0.25, 0.25, 0.25])
    curve_family: list = field(default_factory=lambda: ["G_short", "G_short", "G_long", "G_long"])
    loan_size: list = field(default_factory=lambda: [10000.0] * 4)
    loss_mult: float = 1.4
    prepay_mult: float = 1.0
    recovery_rate: float = 0.0
    recovery_lag: int = 0
    stop_share_short: list = field(default_factory=lambda: [0.0] * 4)   # share of payment paid 1..4 months before charge-off
    stop_share_long: list = field(default_factory=lambda: [0.0] * 4)
    curve_names: list = field(default_factory=lambda: ["G_short", "G_long"])
    loss_curves: list = field(default_factory=lambda: [[0.0] * AGES, [0.0] * AGES])   # cumulative loss % of face by age 0..60
    smm_curves: list = field(default_factory=lambda: [[0.0] * AGES, [0.0] * AGES])    # prepayment SMM by age 0..60
    # 3. Price & fees
    advance: list = field(default_factory=lambda: [0.3] * 4)
    fees: list = field(default_factory=lambda: [Fee() for _ in range(10)])
    # 4. Targets
    tiers: list = field(default_factory=lambda: [Tier() for _ in range(10)])
    # 5. Program
    first_month: str = "2027-01-01"
    purchase_amount: float = 300_000_000.0
    purchase_period: int = 36
    ramp_months: int = 4
    ramp_split: list = field(default_factory=lambda: [0.0] * PLAN_MONTHS)
    # 6a. Capital
    la_fixed: float = 2_500_000.0
    ocm_pct: float = 0.975
    # 6b. Leverage
    ltv: float = 0.5
    spread: float = 0.04
    ineligibles: float = 0.08
    facility_start: int = 7
    last_draw: int = 36
    sweep: float = 1.0
    sofr: list = field(default_factory=lambda: [0.04] * MONTHS)
    # 6c. Other SPV fees
    fee_years_high: int = 2
    fee_high: float = 1_500_000.0
    fee_low: float = 750_000.0
    # 6d. Promote
    hurdles: list = field(default_factory=lambda: [0.14, 0.20, 0.25])      # tier 2 pref, tier 3, tier 4
    la_tier_shares: list = field(default_factory=lambda: [0.05, 0.30, 0.50])  # tiers 3, 4, 5
    promote_method: str = "Absolute"    # Absolute | On top of pro rata
    # 7. School
    school_pv_rate: float = 0.03
    # 8. Sensitivity steps
    advance_step: float = 0.05
    split_step: float = 0.05

    # ---- serialisation ----
    def to_dict(self) -> dict:
        return asdict(self)

    def to_json(self) -> str:
        return json.dumps(self.to_dict(), indent=1)

    @staticmethod
    def from_dict(d: dict) -> "Inputs":
        d = dict(d)
        d["fees"] = [Fee(**f) for f in d.get("fees", [])]
        d["tiers"] = [Tier(**t) for t in d.get("tiers", [])]
        return Inputs(**d)

    @staticmethod
    def from_json(s: str) -> "Inputs":
        return Inputs.from_dict(json.loads(s))

    def copy(self) -> "Inputs":
        return Inputs.from_dict(self.to_dict())


# ------------------------------------------------------------------------- Excel loader
def _num(v, default=0.0):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else default


def load_inputs_from_excel(path: str) -> Inputs:
    """Read every input on 'F1 Fusion Assumptions' into an Inputs object."""
    import openpyxl
    from openpyxl.utils import get_column_letter as L
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["F1 Fusion Assumptions"]
    c = lambda a: ws[a].value
    onoff = lambda a: str(c(a)).strip().upper() == "ON"
    cols4 = ["B", "C", "D", "E"]
    row = lambda r, n, start=2: [ws.cell(r, start + i).value for i in range(n)]

    inp = Inputs()
    inp.recycling, inp.leverage, inp.losses, inp.prepay = onoff("B16"), onoff("B17"), onoff("B18"), onoff("B19")
    inp.payment_stop, inp.promote = onoff("B20"), onoff("B21")
    inp.target_mode, inp.outlay_basis = str(c("B22")), str(c("B23"))
    inp.after_last_tier, inp.capital_basis = str(c("B24")), str(c("B25"))
    inp.terms = [int(c(x + "28")) for x in cols4]
    inp.apr = [_num(c(x + "29")) for x in cols4]
    inp.mix = [_num(c(x + "30")) for x in cols4]
    inp.curve_family = [str(c(x + "31")) for x in cols4]
    inp.loan_size = [_num(c(x + "36")) for x in cols4]
    inp.loss_mult, inp.prepay_mult = _num(c("B40")), _num(c("B41"))
    inp.recovery_rate, inp.recovery_lag = _num(c("B42")), int(_num(c("B43")))
    inp.stop_share_short = [_num(v) for v in row(47, 4)]
    inp.stop_share_long = [_num(v) for v in row(48, 4)]
    inp.curve_names = [str(c("A52")), str(c("A53"))]
    inp.loss_curves = [[_num(v) for v in row(52, 61)], [_num(v) for v in row(53, 61)]]
    inp.smm_curves = [[_num(v) for v in row(54, 61)], [_num(v) for v in row(55, 61)]]
    inp.advance = [_num(c(x + "71")) for x in cols4]
    fees = []
    for r in range(78, 88):
        fees.append(Fee(active=onoff(f"B{r}"), name=str(c(f"C{r}") or ""), paid_to=str(c(f"D{r}") or ""),
                        type=str(c(f"E{r}") or ""), method=str(c(f"F{r}") or ""), basis=str(c(f"G{r}") or ""),
                        rate=c(f"H{r}") if isinstance(c(f"H{r}"), (int, float)) else None,
                        fixed=c(f"I{r}") if isinstance(c(f"I{r}"), (int, float)) else None))
    inp.fees = fees
    tiers = []
    for r in range(101, 111):
        g = lambda a: c(a) if isinstance(c(a), (int, float)) else None
        tiers.append(Tier(active=onoff(f"B{r}"), investor_pct=g(f"C{r}"), moic_target=g(f"E{r}"), irr_target=g(f"F{r}")))
    inp.tiers = tiers
    fm = c("B117")
    inp.first_month = fm.date().isoformat() if isinstance(fm, dt.datetime) else str(fm)
    inp.purchase_amount, inp.purchase_period = _num(c("B118")), int(_num(c("B119")))
    inp.ramp_months = int(_num(c("B125")))
    inp.ramp_split = [_num(v) for v in row(126, 60)]
    inp.la_fixed, inp.ocm_pct = _num(c("B129")), _num(c("B130"))
    inp.ltv, inp.spread, inp.ineligibles = _num(c("B136")), _num(c("B137")), _num(c("B138"))
    inp.facility_start, inp.last_draw, inp.sweep = int(_num(c("B140"))), int(_num(c("B141"))), _num(c("B142"))
    inp.sofr = [_num(v) for v in row(145, 120)]
    inp.fee_years_high, inp.fee_high, inp.fee_low = int(_num(c("B148"))), _num(c("B149")), _num(c("B150"))
    inp.hurdles = [_num(c("C154")), _num(c("D154")), _num(c("E154"))]
    inp.la_tier_shares = [_num(c("D155")), _num(c("E155")), _num(c("F155"))]
    inp.promote_method = str(c("B158"))
    inp.school_pv_rate = _num(c("B161"))
    inp.advance_step, inp.split_step = _num(c("B164")), _num(c("B165"))
    return inp
