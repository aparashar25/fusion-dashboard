"""Build an investor workbook with live Excel formulas that reproduces the dashboard.

Only the Assumptions tab holds typed-in values (blue). Every other number is a formula that traces back to it.
build_workbook(inputs) returns the .xlsx as bytes.
"""
from __future__ import annotations
import datetime as dt
import io
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from .inputs import Inputs

# ------------------------------------------------------------------ house style
F_PCT = '0.00%'
F_X = '0.000"x"'
F_X2 = '0.00"x"'
F_USD = '"$"#,##0;-"$"#,##0'
F_USD3 = '"$"#,##0.000;-"$"#,##0.000'
F_INT = '#,##0'
F_NUM = '#,##0.000'
F_DATE = 'mmm-yy'

NAVY = "1F3A56"
BLUE_IN = "0000FF"
FILL_IN = PatternFill("solid", fgColor="FFF7D6")
FILL_KEY = PatternFill("solid", fgColor="DCE6F1")
FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_SEC = PatternFill("solid", fgColor="EEF2F7")
THIN = Side(style="thin", color="C9D1DC")

AGE0 = 3      # loan age 0 sits in column C; ages 0..60 -> C..BK
MON1 = 3      # month 1 sits in column C; months 1..120 -> C..DR
NA = 61
NM = 120

AS, LE, GN, PL, SP, PD, CH, SA, SS, CK, OV, RM = ("Assumptions", "Loan economics", "Gross vs net", "Pool cash flows",
                                                "SPV and waterfall", "Program detail", "Charts", "Sens - Advance rate",
                                                "Sens - Investor split", "Checks", "Overview", "Read me")


def q(sheet):
    return f"'{sheet}'!"


def ac(c):   # column letter for loan age a (0..60)
    return L(AGE0 + c)


def mc(m):   # column letter for month m (1..120)
    return L(MON1 + m - 1)


class W:
    """Thin helper over a worksheet: labels, rows of formulas, formats."""

    def __init__(self, ws, label_width=58):
        self.ws = ws
        ws.column_dimensions["A"].width = label_width
        ws.sheet_view.showGridLines = False

    def title(self, text, sub=None):
        c = self.ws.cell(1, 1, text)
        c.font = Font(bold=True, size=15, color=NAVY)
        if sub:
            s = self.ws.cell(2, 1, sub)
            s.font = Font(italic=True, size=9, color="5B6675")
        self.ws.freeze_panes = None

    def section(self, r, text, width=8):
        c = self.ws.cell(r, 1, text)
        c.font = Font(bold=True, color="FFFFFF")
        for j in range(1, width + 1):
            self.ws.cell(r, j).fill = FILL_HEAD
        self.ws.cell(r, 1).font = Font(bold=True, color="FFFFFF")

    def label(self, r, text, key=False, note=None, col=1):
        c = self.ws.cell(r, col, text)
        c.font = Font(bold=key)
        if note:
            n = self.ws.cell(r, 2, note)
            n.font = Font(italic=True, size=9, color="5B6675")
        return c

    def put(self, r, c, v, fmt=None, key=False, inp=False, bold=False):
        cell = self.ws.cell(r, c, v)
        if fmt:
            cell.number_format = fmt
        if inp:
            cell.font = Font(color=BLUE_IN)
            cell.fill = FILL_IN
        if key or bold:
            cell.font = Font(bold=True, color=BLUE_IN if inp else None)
        if key:
            cell.fill = FILL_KEY
        return cell

    def key_row(self, r, c1, c2):
        for c in range(c1, c2 + 1):
            cell = self.ws.cell(r, c)
            cell.fill = FILL_KEY
            cell.font = Font(bold=True)

    def header_row(self, r, c1, values, fmt=None):
        for j, v in enumerate(values):
            cell = self.ws.cell(r, c1 + j, v)
            cell.font = Font(bold=True, color=NAVY)
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(horizontal="right")
            if fmt:
                cell.number_format = fmt

    def widths(self, c1, c2, w):
        for c in range(c1, c2 + 1):
            self.ws.column_dimensions[L(c)].width = w


class Book:
    def __init__(self):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.names = {}
        self.ref = {}      # logical name -> absolute reference (with sheet)

    def sheet(self, name, label_width=58):
        return W(self.wb.create_sheet(name), label_width)

    def name(self, nm, sheet, cell):
        """Define a workbook name so formulas read like Recycling, LossMultEff, NTiers ..."""
        col = "".join(ch for ch in cell if ch.isalpha())
        row = "".join(ch for ch in cell if ch.isdigit())
        self.wb.defined_names[nm] = DefinedName(nm, attr_text=f"{q(sheet)}${col}${row}")
        self.names[nm] = f"{q(sheet)}${col}${row}"


def dv_list(ws, cell, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(cell)


# ============================================================================ Assumptions
def build_assumptions(B: Book, inp: Inputs):
    w = B.sheet(AS, 60)
    ws = w.ws
    w.title("Assumptions", "Blue cells on a yellow background are inputs. Everything else in the workbook is a formula.")
    R = {}
    on = lambda b: "ON" if b else "OFF"

    w.section(4, "1. Switches")
    sw = [("Recycling", "Recycling: collections fund new purchases (ON / OFF)", on(inp.recycling), ["ON", "OFF"]),
          ("Leverage", "Leverage facility (ON / OFF)", on(inp.leverage), ["ON", "OFF"]),
          ("LossCurves", "Loss curves (ON / OFF); OFF also turns off recoveries", on(inp.losses), ["ON", "OFF"]),
          ("PrepayCurves", "Prepayment curves (ON / OFF)", on(inp.prepay), ["ON", "OFF"]),
          ("PayStop", "Payment stop before charge-off (ON / OFF)", on(inp.payment_stop), ["ON", "OFF"]),
          ("Promote", "Promote waterfall OCM / Loan America (ON / OFF)", on(inp.promote), ["ON", "OFF"]),
          ("TargetMode", "Target mode for the investor / school split", inp.target_mode, ["MOIC", "IRR"]),
          ("OutlayBasis", "Investor outlay used for the target test", inp.outlay_basis, ["Price + fee", "Price only"]),
          ("AfterLast", "After the last active tier", inp.after_last_tier, ["Continue split", "100% to school"]),
          ("CapitalBasis", "Loan America commitment basis", inp.capital_basis, ["OCM %", "Fixed $"])]
    for i, (nm, lab, val, opts) in enumerate(sw):
        r = 5 + i
        w.label(r, lab)
        w.put(r, 2, val, inp=True)
        dv_list(ws, f"B{r}", opts)
        B.name(nm, AS, f"B{r}")

    # ---- 2. loans and credit
    w.section(16, "2. Loans and credit")
    w.header_row(17, 2, [f"{t}m" for t in inp.terms] + ["Total / WA"])
    R["term"], R["apr"], R["mix"], R["fam"], R["size"] = 18, 19, 20, 21, 22
    rows = [("Term (months)", inp.terms, F_INT), ("APR (annual)", inp.apr, F_PCT), ("Term mix (% of face)", inp.mix, F_PCT),
            ("Curve family", inp.curve_family, None), ("Average loan size, face ($)", inp.loan_size, F_USD)]
    for (lab, vals, fmt), r in zip(rows, [18, 19, 20, 21, 22]):
        w.label(r, lab)
        for k in range(4):
            w.put(r, 2 + k, vals[k], fmt, inp=True)
    for k in range(4):
        dv_list(ws, f"{L(2 + k)}21", list(inp.curve_names))
    R["rate"], R["pmt"], R["fidx"], R["nloans"] = 23, 24, 25, 26
    w.label(23, "Monthly rate (APR / 12)")
    w.label(24, "Monthly payment per $1 of face (level amortisation)")
    w.label(25, "Curve row used (1 = first curve, 2 = second)")
    w.label(26, "Number of loans in the program")
    for k in range(4):
        c = L(2 + k)
        w.put(23, 2 + k, f"={c}19/12", F_PCT)
        w.put(24, 2 + k, f"=IF({c}23=0,1/{c}18,{c}23/(1-(1+{c}23)^-{c}18))", F_USD3)
        w.put(25, 2 + k, f"=IFERROR(MATCH({c}21,$A$40:$A$41,0),1)", F_INT)
        w.put(26, 2 + k, f"=TotalFace*{c}20/{c}22", F_INT)
    w.put(18, 6, "=SUMPRODUCT(B18:E18,B20:E20)", F_NUM)
    w.put(19, 6, "=SUMPRODUCT(B19:E19,B20:E20)", F_PCT)
    w.put(20, 6, "=SUM(B20:E20)", F_PCT, key=True)
    w.put(22, 6, "=IFERROR(1/SUMPRODUCT(B20:E20/B22:E22),0)", F_USD)
    w.put(26, 6, "=SUM(B26:E26)", F_INT)
    w.label(27, "Check: term mix sums to 100%")
    w.put(27, 2, '=IF(ABS(SUM(B20:E20)-1)<0.000001,"OK","CHECK")')
    B.name("ChkMix", AS, "B27")

    w.label(29, "Loss multiplier (x)")
    w.put(29, 2, inp.loss_mult, F_X2, inp=True); B.name("LossMult", AS, "B29")
    w.label(30, "Prepayment multiplier (x)")
    w.put(30, 2, inp.prepay_mult, F_X2, inp=True); B.name("PrepayMult", AS, "B30")
    w.label(31, "Recovery rate (% of booked charge-offs)")
    w.put(31, 2, inp.recovery_rate, F_PCT, inp=True); B.name("RecRate", AS, "B31")
    w.label(32, "Recovery lag (months)")
    w.put(32, 2, inp.recovery_lag, F_INT, inp=True); B.name("RecLag", AS, "B32")
    w.label(33, "Effective loss multiplier used (0 when loss curves are OFF)")
    w.put(33, 2, '=IF(LossCurves="ON",LossMult,0)', F_X2); B.name("LossMultEff", AS, "B33")
    w.label(34, "Effective prepayment multiplier used (0 when prepayment is OFF)")
    w.put(34, 2, '=IF(PrepayCurves="ON",PrepayMult,0)', F_X2); B.name("PrepayMultEff", AS, "B34")
    w.label(35, "Payment stop: share of the scheduled payment still paid ... months before charge-off")
    w.header_row(35, 2, ["1 month", "2 months", "3 months", "4 months"])
    w.label(36, "   short curve (terms up to 36 months)")
    w.label(37, "   long curve (terms over 36 months)")
    for j in range(4):
        w.put(36, 2 + j, inp.stop_share_short[j], F_PCT, inp=True)
        w.put(37, 2 + j, inp.stop_share_long[j], F_PCT, inp=True)

    w.label(39, "Base curves by loan age (cumulative loss = % of original face; prepayment = monthly SMM)")
    w.header_row(39, 2, list(range(61)))
    names = list(inp.curve_names)
    for i in range(2):
        ws.cell(40 + i, 1, names[i]).font = Font(color=BLUE_IN)
        ws.cell(40 + i, 1).fill = FILL_IN
        ws.cell(42 + i, 1, f"{names[i]} prepayment SMM")
        for a in range(61):
            w.put(40 + i, 2 + a, float(inp.loss_curves[i][a]), '0.0000%', inp=True)
            w.put(42 + i, 2 + a, float(inp.smm_curves[i][a]), '0.0000%', inp=True)
    w.label(45, "Effective curves by term (after the multipliers and the term cut-off)")
    w.header_row(45, 2, list(range(61)))
    for k in range(4):
        tc = L(2 + k)
        w.label(46 + k, f"=\"Cumulative loss, \"&{tc}$18&\"m\"")
        w.label(50 + k, f"=\"Prepayment SMM, \"&{tc}$18&\"m\"")
        for a in range(61):
            col = L(2 + a)
            w.put(46 + k, 2 + a, f"=MIN(1,LossMultEff*INDEX($B$40:$BJ$41,{tc}$25,MIN({col}$45,{tc}$18)+1))", '0.0000%')
            w.put(50 + k, 2 + a, f"=IF({col}$45<={tc}$18,PrepayMultEff*INDEX($B$42:$BJ$43,{tc}$25,{col}$45+1),0)", '0.0000%')
    w.label(54, "Check: every effective loss curve is non-decreasing")
    w.put(54, 2, '=IF(SUMPRODUCT(--(C46:BJ49<B46:BI49))=0,"OK","CHECK")')
    B.name("ChkCurves", AS, "B54")
    R.update(loss_eff=46, smm_eff=50)

    # ---- 3. price and fees
    w.section(56, "3. Price and fees")
    w.header_row(57, 2, [f"{t}m" for t in inp.terms] + ["Blended"])
    w.label(58, "Advance rate: cash paid to the school per $1 of face")
    for k in range(4):
        w.put(58, 2 + k, inp.advance[k], F_PCT, inp=True)
    w.put(58, 6, "=SUMPRODUCT(B58:E58,$B$20:$E$20)", F_PCT, key=True); B.name("AdvBlended", AS, "F58")
    w.header_row(60, 1, ["Fee table (set Active to ON to use a row)", "Active", "Fee name", "Paid to", "Type", "Method",
                         "Basis (if %)", "Rate", "Fixed $ per loan", "Row check", "Upfront per $1 (blended)"])
    for i, f in enumerate(inp.fees):
        r = 61 + i
        ws.cell(r, 1, f"Fee {i + 1}")
        w.put(r, 2, "ON" if f.active else "OFF", inp=True); dv_list(ws, f"B{r}", ["ON", "OFF"])
        w.put(r, 3, f.name or None, inp=True)
        w.put(r, 4, f.paid_to or None, inp=True)
        w.put(r, 5, f.type or None, inp=True); dv_list(ws, f"E{r}", ["Upfront", "Across life"])
        w.put(r, 6, f.method or None, inp=True); dv_list(ws, f"F{r}", ["%", "Fixed $"])
        w.put(r, 7, f.basis or None, inp=True)
        dv_list(ws, f"G{r}", ["Net paid", "Face", "Collections", "Beginning balance", "Ending balance", "Average balance"])
        w.put(r, 8, f.rate, F_PCT, inp=True)
        w.put(r, 9, f.fixed, F_USD3, inp=True)
        w.put(r, 10, f'=IF(B{r}="ON",IF(AND(OR(E{r}="Upfront",E{r}="Across life"),OR(AND(F{r}="%",ISNUMBER(H{r}),N(H{r})>=0,'
                     f'IF(E{r}="Upfront",OR(G{r}="Net paid",G{r}="Face"),OR(G{r}="Collections",G{r}="Beginning balance",'
                     f'G{r}="Ending balance",G{r}="Average balance"))),AND(F{r}="Fixed $",ISNUMBER(I{r}),N(I{r})>=0))),"OK","CHECK"),"")')
        w.put(r, 11, f'=IF(AND(J{r}="OK",E{r}="Upfront"),IF(F{r}="%",N(H{r})*IF(G{r}="Face",1,AdvBlended),N(I{r})/$F$22),0)', F_USD3)
    ok = '($J$61:$J$70="OK")'
    w.label(72, "Upfront fees per $1 of face (paid by the investor at month 0)")
    w.label(73, "Investor outlay per $1 of face at month 0 (advance + upfront fees)")
    w.label(74, "Investor outlay per $1 used for the target test")
    for k in range(4):
        c = L(2 + k)
        w.put(72, 2 + k, f'=SUMPRODUCT({ok}*($E$61:$E$70="Upfront")*($F$61:$F$70="%")*$H$61:$H$70*(($G$61:$G$70="Face")+($G$61:$G$70<>"Face")*{c}58))'
                         f'+IFERROR(SUMPRODUCT({ok}*($E$61:$E$70="Upfront")*($F$61:$F$70="Fixed $")*$I$61:$I$70)/{c}22,0)', F_USD3)
        w.put(73, 2 + k, f"={c}58+{c}72", F_USD3)
        w.put(74, 2 + k, f'=IF(OutlayBasis="Price only",{c}58,{c}73)', F_USD3)
    for r in (72, 73, 74):
        w.put(r, 6, f"=SUMPRODUCT(B{r}:E{r},$B$20:$E$20)", F_USD3, key=(r == 73))
    B.name("UpfrontBlended", AS, "F72"); B.name("OutlayBlended", AS, "F73")
    w.label(75, "Upfront %-of-net-paid rates in use (used by the advance-rate sensitivity)")
    w.put(75, 2, f'=SUMPRODUCT({ok}*($E$61:$E$70="Upfront")*($F$61:$F$70="%")*($G$61:$G$70<>"Face")*$H$61:$H$70)', F_PCT)
    B.name("UpfrontNetPaidRate", AS, "B75")
    coef = [("FeeColl", "Across-life: % of monthly collections", '"Collections"', F_PCT),
            ("FeeBeg", "Across-life: annual % of beginning balance (charged / 12)", '"Beginning balance"', F_PCT),
            ("FeeEnd", "Across-life: annual % of ending balance (charged / 12)", '"Ending balance"', F_PCT),
            ("FeeAvg", "Across-life: annual % of average balance (charged / 12)", '"Average balance"', F_PCT)]
    for i, (nm, lab, basis, fmt) in enumerate(coef):
        r = 76 + i
        w.label(r, lab)
        w.put(r, 2, f'=SUMPRODUCT({ok}*($E$61:$E$70="Across life")*($F$61:$F$70="%")*($G$61:$G$70={basis})*$H$61:$H$70)', fmt)
        B.name(nm, AS, f"B{r}")
    w.label(80, "Across-life: fixed $ per active loan per month")
    w.put(80, 2, f'=SUMPRODUCT({ok}*($E$61:$E$70="Across life")*($F$61:$F$70="Fixed $")*$I$61:$I$70)', F_USD3)
    B.name("FeeFixed", AS, "B80")
    w.label(81, "Number of active fees")
    w.put(81, 2, '=COUNTIF(J61:J70,"OK")', F_INT)
    w.label(82, "Check: every active fee row valid; loan sizes positive")
    w.put(82, 2, '=IF(AND(MIN(B22:E22)>0,COUNTIF(J61:J70,"CHECK")=0),"OK","CHECK")')
    B.name("ChkFees", AS, "B82")

    # ---- 4. split targets
    w.section(84, "4. Investor / school split targets")
    w.header_row(85, 1, ["Tier (set Active to ON to use a row)", "Active", "Investor % of net collections", "MOIC target (x outlay)",
                         "IRR target (annual)", "School %", "Row check", "Active count", "Order #", "Investor % in force",
                         "Target in force"])
    for i, t in enumerate(inp.tiers):
        r = 86 + i
        ws.cell(r, 1, f"Tier {i + 1}")
        w.put(r, 2, "ON" if t.active else "OFF", inp=True); dv_list(ws, f"B{r}", ["ON", "OFF"])
        w.put(r, 3, t.investor_pct, F_PCT, inp=True)
        w.put(r, 4, t.moic_target, F_X, inp=True)
        w.put(r, 5, t.irr_target, F_PCT, inp=True)
        w.put(r, 6, f'=IF(ISNUMBER(C{r}),1-C{r},"")', F_PCT)
        w.put(r, 7, f'=IF(B{r}="ON",IF(AND(ISNUMBER(C{r}),N(C{r})>=0,N(C{r})<=1,ISNUMBER(IF(TargetMode="MOIC",D{r},E{r})),'
                    f'N(IF(TargetMode="MOIC",D{r},E{r}))>0),"OK","CHECK"),"")')
        w.put(r, 8, f'=COUNTIF($G$86:G{r},"OK")', F_INT)
        w.put(r, 9, i + 1, F_INT)
        w.put(r, 10, f"=IF(I{r}<=NTiers,INDEX($C$86:$C$95,MATCH(I{r},$H$86:$H$95,0)),0)", F_PCT)
        w.put(r, 11, f'=IF(I{r}<=NTiers,IF(TargetMode="MOIC",INDEX($D$86:$D$95,MATCH(I{r},$H$86:$H$95,0)),'
                     f'INDEX($E$86:$E$95,MATCH(I{r},$H$86:$H$95,0))),0)', F_NUM)
    w.label(97, "Number of active tiers")
    w.put(97, 2, '=COUNTIF(G86:G95,"OK")', F_INT); B.name("NTiers", AS, "B97")
    w.label(98, "After the last active tier: investor % (continue = last tier's split; 100% to school = 0%)")
    w.put(98, 2, '=IF(NTiers=0,0,IF(AfterLast="Continue split",INDEX(J86:J95,NTiers),0))', F_PCT); B.name("AfterPct", AS, "B98")
    w.label(99, "Check: at least one active tier, every active row valid")
    w.put(99, 2, '=IF(AND(NTiers>0,COUNTIF(G86:G95,"CHECK")=0),"OK","CHECK")'); B.name("ChkTiers", AS, "B99")
    w.label(100, "Single constant split (needed by the two sensitivity tabs)")
    w.put(100, 2, '=IF(AND(NTiers>0,SUMPRODUCT((J86:J95<>J86)*(I86:I95<=NTiers))=0,AfterPct=J86),"OK","CHECK")')
    B.name("SingleSplit", AS, "B100")
    w.label(101, "Investor % used by the sensitivity tabs")
    w.put(101, 2, "=J86", F_PCT); B.name("SplitPct", AS, "B101")

    # ---- 5. program
    w.section(103, "5. Program plan and ramp-up")
    w.label(104, "First purchase month")
    fm = dt.date.fromisoformat(str(inp.first_month)[:10])
    w.put(104, 2, fm, F_DATE, inp=True); B.name("FirstMonth", AS, "B104")
    w.label(105, "Purchase amount: cash paid to schools for the program ($)")
    w.put(105, 2, inp.purchase_amount, F_USD, inp=True); B.name("PurchaseAmt", AS, "B105")
    w.label(106, "Purchase period (months, 5 to 60)")
    w.put(106, 2, inp.purchase_period, F_INT, inp=True); B.name("Period", AS, "B106")
    w.label(107, "Ramp-up period (months)")
    w.put(107, 2, inp.ramp_months, F_INT, inp=True); B.name("RampMonths", AS, "B107")
    w.label(108, "Total face purchased ($)")
    w.put(108, 2, "=SUM(B113:BI113)", F_USD, key=True); B.name("TotalFace", AS, "B108")
    w.label(109, "Month of the program")
    w.header_row(109, 2, list(range(1, 61)))
    w.label(110, "Ramp-up purchase split inputs (% of purchase amount; only the ramp-up months are used)")
    w.label(111, "Monthly purchase split (% of purchase amount)")
    w.label(112, "Monthly purchase cash paid to schools ($)")
    w.label(113, "Monthly face purchased ($) = cash / blended advance rate")
    for m in range(1, 61):
        c = L(1 + m)
        w.put(110, 1 + m, float(inp.ramp_split[m - 1]) if m - 1 < len(inp.ramp_split) else 0.0, F_PCT, inp=True)
        w.put(111, 1 + m, f'=IF({c}109<=RampMonths,{c}110,IF(AND({c}109<=Period,Period>RampMonths),'
                          f'(1-SUMIF($B$109:$BI$109,"<="&RampMonths,$B$110:$BI$110))/(Period-RampMonths),0))', F_PCT)
        w.put(112, 1 + m, f"=PurchaseAmt*{c}111", F_USD)
        w.put(113, 1 + m, f"=IF(AdvBlended>0,{c}112/AdvBlended,0)", F_USD)
    w.label(114, "Check: purchase split sums to 100%, period 5 to 60 months")
    w.put(114, 2, '=IF(AND(ABS(SUM(B111:BI111)-1)<0.000001,Period>=5,Period<=60),"OK","CHECK")'); B.name("ChkSplit", AS, "B114")
    w.label(115, "Check: monthly purchase cash adds up to the program amount")
    w.put(115, 2, '=IF(ABS(SUM(B112:BI112)-PurchaseAmt)<1,"OK","CHECK")'); B.name("ChkCash", AS, "B115")
    w.label(116, "Check: ramp-up shorter than the purchase period; ramp-up % at most 100%")
    w.put(116, 2, '=IF(AND(RampMonths>=0,RampMonths<Period,SUMIF($B$109:$BI$109,"<="&RampMonths,$B$110:$BI$110)<=1),"OK","CHECK")')
    B.name("ChkRamp", AS, "B116")

    # ---- 6. SPV structure
    w.section(118, "6. SPV structure")
    rows6 = [(119, "LAFixed", "Loan America fixed commitment ($) (used when the basis is Fixed $)", inp.la_fixed, F_USD),
             (120, "OCMPct", "OCM share of capital (used when the basis is OCM %)", inp.ocm_pct, F_PCT),
             (122, "LTV", "Leverage: advance rate on eligible cost basis (LTV)", inp.ltv, F_PCT),
             (123, "Spread", "Leverage: spread over SOFR", inp.spread, F_PCT),
             (124, "Inelig", "Leverage: ineligibles (% of cost basis)", inp.ineligibles, F_PCT),
             (125, "FacStart", "Leverage: facility start month", inp.facility_start, F_INT),
             (126, "LastDraw", "Leverage: last draw month (the facility amortises after it)", inp.last_draw, F_INT),
             (127, "Sweep", "Leverage: cash sweep in amortisation", inp.sweep, F_PCT),
             (133, "FeeYearsHigh", "Other SPV fees: years at the higher fee", inp.fee_years_high, F_INT),
             (134, "FeeHigh", "Other SPV fees per year, first years ($)", inp.fee_high, F_USD),
             (135, "FeeLow", "Other SPV fees per year, thereafter ($)", inp.fee_low, F_USD),
             (137, "Hurdle2", "Promote: tier 2 preferred return (IRR)", inp.hurdles[0], F_PCT),
             (138, "Hurdle3", "Promote: tier 3 runs until OCM reaches this IRR", inp.hurdles[1], F_PCT),
             (139, "Hurdle4", "Promote: tier 4 runs until OCM reaches this IRR", inp.hurdles[2], F_PCT),
             (140, "LAShare3", "Promote: Loan America share in tier 3", inp.la_tier_shares[0], F_PCT),
             (141, "LAShare4", "Promote: Loan America share in tier 4", inp.la_tier_shares[1], F_PCT),
             (142, "LAShare5", "Promote: Loan America share in tier 5 (residual)", inp.la_tier_shares[2], F_PCT),
             (143, "PromoteMethod", "Promote method (Absolute / On top of pro rata)", inp.promote_method, None)]
    for r, nm, lab, val, fmt in rows6:
        w.label(r, lab)
        w.put(r, 2, val, fmt, inp=True)
        B.name(nm, AS, f"B{r}")
    dv_list(ws, "B143", ["Absolute", "On top of pro rata"])
    w.label(128, "Maximum facility ($) = LTV x purchase amount")
    w.put(128, 2, "=LTV*PurchaseAmt", F_USD); B.name("FacilityMax", AS, "B128")
    w.label(129, "Check: last draw month on or after the start month")
    w.put(129, 2, '=IF(LastDraw>=FacStart,"OK","CHECK")'); B.name("ChkLev", AS, "B129")
    w.label(130, "Warning: last draw month covers the purchase period")
    w.put(130, 2, '=IF(OR(Leverage="OFF",LastDraw>=Period),"OK","WARNING")'); B.name("WarnDraw", AS, "B130")
    w.label(131, "SOFR (annual) by month of the projection")
    w.header_row(131, 2, list(range(1, 121)))
    for m in range(1, 121):
        w.put(132, 1 + m, float(inp.sofr[m - 1]), F_PCT, inp=True)
    w.label(132, "SOFR (annual)")

    w.section(145, "7. School and 8. Sensitivity settings")
    w.label(146, "School discount rate for PV of school cash flows (annual)")
    w.put(146, 2, inp.school_pv_rate, F_PCT, inp=True); B.name("SchoolPV", AS, "B146")
    w.label(147, "Advance-rate step size (advance-rate sensitivity)")
    w.put(147, 2, inp.advance_step, F_PCT, inp=True); B.name("AdvStep", AS, "B147")
    w.label(148, "Investor / school split step size (split sensitivity)")
    w.put(148, 2, inp.split_step, F_PCT, inp=True); B.name("SplitStep", AS, "B148")

    w.section(150, "Overall input check")
    w.label(151, "All input checks (term mix, curves, fees, tiers, purchase split and cash, ramp-up, leverage months)", key=True)
    w.put(151, 2, '=IF(COUNTIF(B152:B159,"CHECK")>0,"CHECK",IF(WarnDraw="WARNING","WARNING","OK"))', key=True)
    B.name("InputStatus", AS, "B151")
    for i, nm in enumerate(["ChkMix", "ChkCurves", "ChkFees", "ChkTiers", "ChkSplit", "ChkCash", "ChkRamp", "ChkLev"]):
        w.label(152 + i, f"   {nm.replace('Chk', '')}")
        w.put(152 + i, 2, f"={nm}")
    w.widths(2, 12, 14)
    ws.freeze_panes = "B4"
    B.ref["AS"] = R
    return R


# ============================================================================ split engine rows (shared by Loan economics and Gross vs net)
def write_split(w: W, start: int, net_row: int, outlay_ref: str, label_prefix=""):
    """Rows of the investor / school split for one term: 10 need rows, 11 breakpoints, 10 allocations, after-last,
    investor, school, tiers met, tier in force, investor % in force. Returns a dict of row numbers."""
    ws = w.ws
    R = {}
    r = start
    tier_pct = lambda k: f"{q(AS)}$J${85 + k}"
    tier_tgt = lambda k: f"{q(AS)}$K${85 + k}"
    R["need"] = r
    for k in range(1, 11):
        w.label(r, f"{label_prefix}Cash still needed to finish tier {k} (start of month)")
        ws.cell(r, 2, k).font = Font(color="8A94A3", size=8)
        for a in range(NA):
            c, p = ac(a), ac(a - 1) if a else None
            if a == 0:
                f = f'=IF({k}>NTiers,0,IF(TargetMode="MOIC",{tier_tgt(k)}*{outlay_ref},{outlay_ref}))'
            else:
                f = (f'=IF({k}>NTiers,0,IF(TargetMode="MOIC",{p}{r}-{p}{start + 32},'
                     f'({p}{r}-{p}{start + 32})*(1+{tier_tgt(k)})^(1/12)))')
            w.put(r, AGE0 + a, f, F_USD3)
        r += 1
    R["bp"] = r
    w.label(r, f"{label_prefix}Breakpoint 0")
    for a in range(NA):
        w.put(r, AGE0 + a, 0, F_USD3)
    r += 1
    for k in range(1, 11):
        w.label(r, f"{label_prefix}Breakpoint {k}: net collections needed to finish tiers 1-{k}")
        need, prev_need = R["need"] + k - 1, R["need"] + k - 2
        for a in range(NA):
            c = ac(a)
            inc = f"MAX(0,{c}{need})" if k == 1 else f"MAX(0,MAX({c}{need},0)-MAX({c}{prev_need},0))"
            f = f"={c}{r - 1}+IF({k}>NTiers,0,IF({inc}=0,0,IF({tier_pct(k)}=0,1E+12,{inc}/{tier_pct(k)})))"
            w.put(r, AGE0 + a, f, F_USD3)
        r += 1
    R["alloc"] = r
    for k in range(1, 11):
        w.label(r, f"{label_prefix}Net collections allocated to tier {k}")
        b0, b1 = R["bp"] + k - 1, R["bp"] + k
        for a in range(NA):
            c = ac(a)
            w.put(r, AGE0 + a, f"=MAX(0,MIN({c}{net_row},{c}{b1})-{c}{b0})", F_USD3)
        r += 1
    R["after"] = r
    w.label(r, f"{label_prefix}Net collections after the last tier")
    for a in range(NA):
        c = ac(a)
        w.put(r, AGE0 + a, f"=MAX(0,{c}{net_row}-{c}{R['bp'] + 10})", F_USD3)
    r += 1
    R["inv"] = r
    assert r == start + 32
    w.label(r, f"{label_prefix}Investor share of net collections", key=True)
    for a in range(NA):
        c = ac(a)
        w.put(r, AGE0 + a, f"=SUMPRODUCT({q(AS)}$J$86:$J$95,{c}{R['alloc']}:{c}{R['alloc'] + 9})+AfterPct*{c}{R['after']}", F_USD3)
    r += 1
    R["school"] = r
    w.label(r, f"{label_prefix}School share of net collections")
    for a in range(NA):
        c = ac(a)
        w.put(r, AGE0 + a, f"={c}{net_row}-{c}{R['inv']}", F_USD3)
    r += 1
    R["met"] = r
    w.label(r, f"{label_prefix}Tiers met so far")
    for a in range(NA):
        c = ac(a)
        w.put(r, AGE0 + a, f"=SUMPRODUCT(({c}{R['need']}:{c}{R['need'] + 9}-{c}{R['inv']}<=1E-9)*($B${R['need']}:$B${R['need'] + 9}<=NTiers))", F_INT)
    r += 1
    R["tif"] = r
    w.label(r, f"{label_prefix}Tier in force this month (NTiers + 1 = after the last tier)")
    for a in range(NA):
        c = ac(a)
        f = "=MIN(NTiers+1,1)" if a == 0 else f"=MIN(NTiers+1,1+{ac(a - 1)}{R['met']})"
        w.put(r, AGE0 + a, f, F_INT)
    r += 1
    R["pct"] = r
    w.label(r, f"{label_prefix}Investor % in force")
    for a in range(NA):
        c = ac(a)
        w.put(r, AGE0 + a, f"=IF({c}{R['tif']}>NTiers,AfterPct,INDEX({q(AS)}$J$86:$J$95,{c}{R['tif']}))", F_PCT)
    r += 1
    R["end"] = r
    return R


# ============================================================================ Loan economics
def build_loan_economics(B: Book, inp: Inputs):
    w = B.sheet(LE, 62)
    ws = w.ws
    w.title("Loan economics", "Per $1 of loan face, by term and blended. Month 0 is the purchase month.")
    blocks = {}
    BLOCK0, SIZE = 120, 112
    for k in range(4):
        blocks[k] = term_block(w, BLOCK0 + k * SIZE, k, inp)
    bl = blended_block(w, 72, blocks)
    B.ref["LE_sum"] = summary(w, 4, blocks, bl)
    tier_summary(w, 44, blocks, inp)
    w.widths(2, 2, 13)
    w.widths(3, 3 + NA, 11)
    ws.freeze_panes = "C4"
    B.ref["LE"] = {"blocks": blocks, "bl": bl}
    return blocks, bl


def term_block(w: W, s: int, k: int, inp: Inputs):
    ws = w.ws
    tc = L(2 + k)
    A_ = lambda row: f"{q(AS)}${tc}${row}"
    T, rr, pmt, size = A_(18), A_(23), A_(24), A_(22)
    w.section(s, f"Term block: {inp.terms[k]}-month loans (per $1 of face)", width=12)
    w.label(s + 1, "Loan age (months since purchase)")
    w.header_row(s + 1, AGE0, list(range(NA)))
    w.label(s + 2, "Term (months)"); w.put(s + 2, 2, f"={T}", F_INT)
    w.label(s + 3, "Monthly rate"); w.put(s + 3, 2, f"={rr}", F_PCT)
    w.label(s + 4, "Monthly payment per $1 of face"); w.put(s + 4, 2, f"={pmt}", F_USD3)
    w.label(s + 5, "Mix weight (% of face)"); w.put(s + 5, 2, f"={A_(20)}", F_PCT)
    Tc, rc, pc = f"$B${s + 2}", f"$B${s + 3}", f"$B${s + 4}"
    R = {}
    names = ["S", "fac", "loss", "smm", "open", "co", "surv", "int", "sch", "pre", "end", "active", "rec", "missed", "pps",
             "s104", "defbal", "defface", "removed", "gross", "fees", "net"]
    labels = {"S": "Scheduled balance (no losses or prepayments)", "fac": "Scheduled principal factor (share of balance amortising)",
              "loss": "Cumulative loss curve (% of face)", "smm": "Prepayment SMM", "open": "Opening performing balance",
              "co": "Charge-offs", "surv": "Surviving balance after charge-offs", "int": "Interest collected",
              "sch": "Scheduled principal collected", "pre": "Prepaid principal", "end": "Ending performing balance",
              "active": "Active loans (share of the original count)", "rec": "Recoveries",
              "missed": "Payments missed before charge-off (payment stop)", "pps": "Prepayment survival",
              "s104": "Survivors as if payment stop were OFF (caps charge-offs)", "defbal": "Balance of one defaulting loan at charge-off (per $1 defaulted face)",
              "defface": "Defaulted face this month", "removed": "Balance removed from the performing pool at charge-off",
              "gross": "Gross collections", "fees": "Across-life fees (off the top)", "net": "Net collections"}
    for i, nm in enumerate(names):
        R[nm] = s + 6 + i
        w.label(R[nm], labels[nm], key=nm in ("gross", "net"))
    rows = lambda nm: R[nm]
    age = lambda a: f"{ac(a)}${s + 1}"
    stop = lambda j: f"INDEX({q(AS)}$B$36:$E$37,IF({Tc}<=36,1,2),{j})"
    Srow = f"${ac(0)}${R['S']}:${ac(60)}${R['S']}"
    COrow = f"${ac(0)}${R['co']}:${ac(60)}${R['co']}"
    DFrow = f"${ac(0)}${R['defface']}:${ac(60)}${R['defface']}"
    for a in range(NA):
        c = ac(a)
        p = ac(a - 1) if a else None
        put = lambda nm, f, fmt=F_USD3: w.put(R[nm], AGE0 + a, f, fmt)
        put("S", f"=IF({a}<={Tc},IF({rc}=0,1-{a}/{Tc},MAX(0,(1+{rc})^{a}-{pc}*((1+{rc})^{a}-1)/{rc})),0)")
        put("fac", "=0" if a == 0 else f"=IF(AND({a}<={Tc},{p}{R['S']}>0),({p}{R['S']}-{c}{R['S']})/{p}{R['S']},0)", F_PCT)
        put("loss", f"={q(AS)}{L(2 + a)}${46 + k}", F_PCT)
        put("smm", f"={q(AS)}{L(2 + a)}${50 + k}", F_PCT)
        put("open", "=1" if a == 0 else f"={p}{R['end']}")
        if a == 0:
            put("co", f"=MIN({c}{R['open']},{c}{R['loss']})")
            put("pps", f"=1-{c}{R['smm']}", F_PCT)
            put("s104", f'=IF(PayStop="ON",IF({c}{R["open"]}>0,1-{c}{R["co"]}/{c}{R["open"]},0),0)', F_PCT)
            put("surv", f'={c}{R["open"]}-IF(PayStop="ON",{c}{R["defface"]},{c}{R["co"]})')
            put("int", "=0"); put("sch", "=0"); put("active", "=1", F_PCT); put("missed", "=0")
        else:
            base = f"{p}{R['S']}*{p}{R['s104']}*{p}{R['pps']}"
            put("co", f'=MIN(MAX(0,{c}{R["loss"]}-{p}{R["loss"]}),IF(PayStop="ON",IF({a}<={Tc},{base},0),{c}{R["open"]}))')
            put("pps", f"=IF({a}<={Tc},{p}{R['pps']}*(1-{c}{R['smm']}),{p}{R['pps']})", F_PCT)
            put("s104", f'=IF(PayStop="ON",IF({a}<={Tc},IF({base}>0,{p}{R["s104"]}*(1-{c}{R["co"]}/({base})),0),{p}{R["s104"]}),0)', F_PCT)
            put("surv", f'={c}{R["open"]}-IF(PayStop="ON",{c}{R["defface"]}*{p}{R["S"]},{c}{R["co"]})')
            put("int", f"={c}{R['surv']}*{rc}")
            put("sch", f"={c}{R['surv']}*{c}{R['fac']}")
            put("active", f"=IF({p}{R['S']}>0,{c}{R['open']}/{p}{R['S']},0)", F_PCT)
            terms = "+".join(f"(1-{stop(j)})*IF({a}+{j}<={Tc},INDEX({DFrow},{a}+{j}+1),0)" for j in range(1, 5))
            put("missed", f'=IF(AND(PayStop="ON",{a}<{Tc}),{pc}*({terms}),0)')
        put("pre", f"={c}{R['smm']}*({c}{R['surv']}-{c}{R['sch']})")
        put("end", f"={c}{R['surv']}-{c}{R['sch']}-{c}{R['pre']}")
        put("rec", f"=IF({a}-RecLag>=0,RecRate*INDEX({COrow},{a}-RecLag+1),0)")
        parts = "".join(f"-IF({a}-{j}>=1,{stop(j)}*({pc}-{rc}*INDEX({Srow},{a}-{j})),0)" for j in range(1, 5))
        put("defbal", f'=IF(AND(PayStop="ON",{a}<={Tc}),INDEX({Srow},MAX(0,{a}-5)+1){parts},0)')
        put("defface", f'=IF(AND(PayStop="ON",{a}<={Tc},{c}{R["defbal"]}>0),{c}{R["co"]}/{c}{R["defbal"]},0)')
        put("removed", f"={c}{R['open']}-{c}{R['surv']}")
        put("gross", f"={c}{R['int']}+{c}{R['sch']}+{c}{R['pre']}+{c}{R['rec']}-{c}{R['missed']}")
        fix = f"+FeeFixed/{size}*{c}{R['active']}"
        put("fees", f"=MIN({c}{R['gross']},FeeColl*{c}{R['gross']}" +
                    ("" if a == 0 else f"+(FeeBeg*{c}{R['open']}+FeeEnd*{c}{R['end']}+FeeAvg*({c}{R['open']}+{c}{R['end']})/2)/12{fix}") + ")")
        put("net", f"={c}{R['gross']}-{c}{R['fees']}")
    w.key_row(R["net"], 1, 1)
    sp = write_split(w, R["net"] + 2, R["net"], f"{q(AS)}${tc}$74")
    R.update({f"sp_{k_}": v for k_, v in sp.items()})
    r = sp["end"] + 1
    lab = [("cf", "Investor cash flow (month 0 = investor share - outlay)"), ("cum", "Investor cumulative cash received"),
           ("moic", "Investor running MOIC"), ("irr", "Investor running IRR (annualised)"),
           ("scf", "School cash flow (month 0 includes the advance)"), ("scum", "School cumulative cash"),
           ("chk", "Check: investor + school - net collections (0)")]
    for i, (nm, t) in enumerate(lab):
        R[nm] = r + i
        w.label(R[nm], t, key=nm in ("moic", "irr"))
    for a in range(NA):
        c = ac(a)
        inv = f"{c}{sp['inv']}"
        w.put(R["cf"], AGE0 + a, f"={inv}-{q(AS)}${tc}$73" if a == 0 else f"={inv}", F_USD3)
        w.put(R["cum"], AGE0 + a, f"=SUM(${ac(0)}{sp['inv']}:{c}{sp['inv']})", F_USD3)
        w.put(R["moic"], AGE0 + a, f"={c}{R['cum']}/{q(AS)}${tc}$73", F_X)
        w.put(R["irr"], AGE0 + a, "" if a == 0 else f'=IFERROR((1+IRR(${ac(0)}{R["cf"]}:{c}{R["cf"]},0.01))^12-1,"")', F_PCT)
        w.put(R["scf"], AGE0 + a, f"={c}{sp['school']}+{q(AS)}${tc}$58" if a == 0 else f"={c}{sp['school']}", F_USD3)
        w.put(R["scum"], AGE0 + a, f"=SUM(${ac(0)}{R['scf']}:{c}{R['scf']})", F_USD3)
        w.put(R["chk"], AGE0 + a, f"=ROUND({inv}+{c}{sp['school']}-{c}{R['net']},9)", F_USD3)
    for nm in ("gross", "net", "co", "int", "sch", "pre", "rec", "missed", "removed", "fees", "cf", "scf", "chk"):
        w.put(R[nm], 2, f"=SUM({ac(0)}{R[nm]}:{ac(60)}{R[nm]})", F_USD3, bold=True)
    w.put(sp["inv"], 2, f"=SUM({ac(0)}{sp['inv']}:{ac(60)}{sp['inv']})", F_USD3, bold=True)
    w.put(sp["school"], 2, f"=SUM({ac(0)}{sp['school']}:{ac(60)}{sp['school']})", F_USD3, bold=True)
    for k_ in range(10):
        rr_ = sp["alloc"] + k_
        w.put(rr_, 2, f"=SUM({ac(0)}{rr_}:{ac(60)}{rr_})", F_USD3)
    w.put(sp["after"], 2, f"=SUM({ac(0)}{sp['after']}:{ac(60)}{sp['after']})", F_USD3)
    R["principal_chk"] = R["chk"] + 1
    w.label(R["principal_chk"], "Check: removed + scheduled + prepaid + final balance - 1 (0)")
    w.put(R["principal_chk"], 2, f"=ROUND(B{R['removed']}+B{R['sch']}+B{R['pre']}+{ac(60)}{R['end']}-1,9)", F_USD3)
    w.ws.cell(s + 1, 2, "Total").font = Font(bold=True, color=NAVY)
    R["T"], R["start"] = Tc, s
    return R


def blended_block(w: W, s: int, blocks: dict):
    w.section(s, "Blended (mix-weighted across the four terms, per $1 of face)", width=12)
    w.label(s + 1, "Loan age")
    w.header_row(s + 1, AGE0, list(range(NA)))
    w.ws.cell(s + 1, 2, "Total").font = Font(bold=True, color=NAVY)
    items = [("open", "Opening performing balance", "open"), ("co", "Charge-offs", "co"), ("int", "Interest collected", "int"),
             ("sch", "Scheduled principal collected", "sch"), ("pre", "Prepaid principal", "pre"), ("end", "Ending performing balance", "end"),
             ("rec", "Recoveries", "rec"), ("missed", "Payments missed before charge-off", "missed"),
             ("removed", "Balance removed at charge-off", "removed"), ("gross", "Gross collections", "gross"),
             ("fees", "Across-life fees", "fees"), ("net", "Net collections", "net"), ("inv", "Investor share of net collections", "sp_inv"),
             ("school", "School share of net collections", "sp_school"), ("cf", "Investor cash flow", "cf")]
    R = {}
    mix = lambda k: f"$B${blocks[k]['start'] + 5}"
    r = s + 2
    for nm, lab, src in items:
        R[nm] = r
        w.label(r, lab, key=nm in ("net", "inv"))
        for a in range(NA):
            c = ac(a)
            w.put(r, AGE0 + a, "=" + "+".join(f"{mix(k)}*{c}{blocks[k][src]}" for k in range(4)), F_USD3)
        if nm not in ("open", "end"):
            w.put(r, 2, f"=SUM({ac(0)}{r}:{ac(60)}{r})", F_USD3, bold=True)
        r += 1
    extra = [("cum", "Investor cumulative cash received"), ("moic", "Investor running MOIC"), ("irr", "Investor running IRR (annualised)"),
             ("scf", "School cash flow (month 0 includes the advance)"), ("scum", "School cumulative cash"),
             ("cost", "Investor cost basis (advance rate x performing balance), for the leverage facility")]
    for nm, lab in extra:
        R[nm] = r
        w.label(r, lab, key=nm in ("moic", "irr"))
        r += 1
    for a in range(NA):
        c = ac(a)
        w.put(R["cum"], AGE0 + a, f"=SUM(${ac(0)}{R['inv']}:{c}{R['inv']})", F_USD3)
        w.put(R["moic"], AGE0 + a, f"={c}{R['cum']}/OutlayBlended", F_X)
        w.put(R["irr"], AGE0 + a, "" if a == 0 else f'=IFERROR((1+IRR(${ac(0)}{R["cf"]}:{c}{R["cf"]},0.01))^12-1,"")', F_PCT)
        w.put(R["scf"], AGE0 + a, f"={c}{R['school']}+AdvBlended" if a == 0 else f"={c}{R['school']}", F_USD3)
        w.put(R["scum"], AGE0 + a, f"=SUM(${ac(0)}{R['scf']}:{c}{R['scf']})", F_USD3)
        w.put(R["cost"], AGE0 + a, "=" + "+".join(f"{mix(k)}*{q(AS)}${L(2 + k)}$58*{c}{blocks[k]['end']}" for k in range(4)), F_USD3)
    w.put(R["scf"], 2, f"=SUM({ac(0)}{R['scf']}:{ac(60)}{R['scf']})", F_USD3, bold=True)
    return R


def summary(w: W, s: int, blocks: dict, bl: dict):
    ws = w.ws
    w.section(s, "Summary (per $1 of face)", width=7)
    w.header_row(s + 1, 2, ["24m", "36m", "48m", "60m", "Blended"])
    for k in range(4):
        ws.cell(s + 1, 2 + k, f'={q(AS)}{L(2 + k)}18&"m"').font = Font(bold=True, color=NAVY)
    rows = [("term", "Term (months)", F_NUM), ("mix", "Mix weight (% of face)", F_PCT), ("apr", "APR", F_PCT),
            ("adv", "Advance paid to the school (price)", F_USD3), ("up", "Upfront fees", F_USD3), ("out", "Investor outlay at month 0", F_USD3),
            ("gross", "Gross collections", F_USD3), ("fees", "Across-life fees", F_USD3), ("net", "Net collections", F_USD3),
            ("co", "Charge-offs (% of face)", F_PCT), ("missed", "Payments missed before charge-off", F_USD3),
            ("inv", "Investor cash received", F_USD3), ("moic", "Investor MOIC", F_X), ("irr", "Investor IRR (annualised)", F_PCT),
            ("hit", "Month tier 1 hit (loan age)", F_INT), ("s_adv", "School: received at day 0 (per $1 of face)", F_USD3),
            ("s_coll", "School: received through collections, undiscounted", F_USD3), ("s_tot", "School: total received, undiscounted", F_USD3),
            ("s_pvc", "School: PV of collections", F_USD3), ("s_pv", "School: total received, PV", F_USD3),
            ("s_disc", "Effective discount taken by the school, undiscounted (1 - total received)", F_PCT),
            ("s_discpv", "Effective discount taken by the school, PV (1 - PV received)", F_PCT),
            ("pv", "Discount rate used for the PV (annual)", F_PCT), ("chk", "Checks (split and balance identities = 0)", None)]
    R = {}
    for i, (nm, lab, fmt) in enumerate(rows):
        r = s + 2 + i
        R[nm] = r
        key = nm in ("net", "inv", "moic", "irr", "s_tot", "s_pv")
        w.label(r, lab, key=key)
        for j in range(5):
            col = 2 + j
            if j < 4:
                b = blocks[j]; tc = L(2 + j)
                As = lambda row: f"{q(AS)}${tc}${row}"
                src = {"term": f"={As(18)}", "mix": f"={As(20)}", "apr": f"={As(19)}", "adv": f"={As(58)}", "up": f"={As(72)}",
                       "out": f"={As(73)}", "gross": f"=$B${b['gross']}", "fees": f"=$B${b['fees']}", "net": f"=$B${b['net']}",
                       "co": f"=$B${b['co']}", "missed": f"=$B${b['missed']}", "inv": f"=$B${b['sp_inv']}",
                       "moic": f"=$B${b['sp_inv']}/{As(73)}", "irr": f'=IFERROR((1+IRR({ac(0)}{b["cf"]}:{ac(60)}{b["cf"]},0.01))^12-1,"n.a.")',
                       "hit": f'=IF(COUNTIF({ac(0)}{b["sp_met"]}:{ac(60)}{b["sp_met"]},">=1")=0,"Not hit",COUNTIF({ac(0)}{b["sp_met"]}:{ac(60)}{b["sp_met"]},"<1"))',
                       "s_adv": f"={As(58)}", "s_coll": f"=$B${b['sp_school']}",
                       "s_pvc": f"={ac(0)}{b['sp_school']}+NPV((1+SchoolPV)^(1/12)-1,{ac(1)}{b['sp_school']}:{ac(60)}{b['sp_school']})",
                       "chk": f'=IF(AND(ABS($B${b["chk"]})<0.000001,ABS($B${b["principal_chk"]})<0.000001),"OK","CHECK")'}
            else:
                src = {"term": "=SUMPRODUCT(B{0}:E{0},B{1}:E{1})".format(s + 2, s + 3), "mix": f"=SUM(B{s + 3}:E{s + 3})",
                       "apr": f"=SUMPRODUCT(B{s + 4}:E{s + 4},B{s + 3}:E{s + 3})", "adv": "=AdvBlended", "up": "=UpfrontBlended",
                       "out": "=OutlayBlended", "gross": f"=$B${bl['gross']}", "fees": f"=$B${bl['fees']}", "net": f"=$B${bl['net']}",
                       "co": f"=$B${bl['co']}", "missed": f"=$B${bl['missed']}", "inv": f"=$B${bl['inv']}",
                       "moic": f"=$B${bl['inv']}/OutlayBlended", "irr": f'=IFERROR((1+IRR({ac(0)}{bl["cf"]}:{ac(60)}{bl["cf"]},0.01))^12-1,"n.a.")',
                       "hit": '="n.a."', "s_adv": "=AdvBlended", "s_coll": f"=$B${bl['school']}",
                       "s_pvc": f"={ac(0)}{bl['school']}+NPV((1+SchoolPV)^(1/12)-1,{ac(1)}{bl['school']}:{ac(60)}{bl['school']})",
                       "chk": f'=IF(AND(COUNTIF(B{s + 2 + 23}:E{s + 2 + 23},"OK")=4,ABS(SUMPRODUCT(ABS({ac(0)}{bl["inv"]}:{ac(60)}{bl["inv"]}+{ac(0)}{bl["school"]}:{ac(60)}{bl["school"]}-{ac(0)}{bl["net"]}:{ac(60)}{bl["net"]})))<0.000001),"OK","CHECK")'}
            L_ = L(col)
            if nm == "s_tot":
                f = f"={L_}{R['s_adv']}+{L_}{R['s_coll']}"
            elif nm == "s_pv":
                f = f"={L_}{R['s_adv']}+{L_}{R['s_pvc']}"
            elif nm == "s_disc":
                f = f"=1-{L_}{R['s_tot']}"
            elif nm == "s_discpv":
                f = f"=1-{L_}{R['s_pv']}"
            elif nm == "pv":
                f = "=SchoolPV"
            else:
                f = src[nm]
            w.put(r, col, f, fmt, key=key)
    w.widths(2, 6, 13)
    w.ws.column_dimensions["B"].width = 13
    return R


def tier_summary(w: W, s: int, blocks: dict, inp: Inputs):
    ws = w.ws
    w.section(s, "Tier summary: loan age each target is hit, and cash within each tier", width=24)
    heads = ["Investor %", "Target"]
    for k in range(4):
        heads += [f"{inp.terms[k]}m age hit", f"{inp.terms[k]}m investor $", f"{inp.terms[k]}m school $", f"{inp.terms[k]}m MOIC at hit", f"{inp.terms[k]}m IRR at hit"]
    w.header_row(s + 1, 2, heads)
    for t in range(1, 11):
        r = s + 1 + t
        w.label(r, f"Tier {t}")
        w.put(r, 2, f'=IF({t}<=NTiers,{q(AS)}$J${85 + t},"")', F_PCT)
        w.put(r, 3, f'=IF({t}<=NTiers,{q(AS)}$K${85 + t},"")', F_NUM)
        for k in range(4):
            b = blocks[k]; c0 = 4 + 5 * k
            met = f"${ac(0)}${b['sp_met']}:${ac(60)}${b['sp_met']}"
            alloc = b["sp_alloc"] + t - 1
            age = f'IF(COUNTIF({met},">="&{t})=0,"Not hit",COUNTIF({met},"<"&{t}))'
            w.put(r, c0, f'=IF({t}<=NTiers,{age},"")', F_INT)
            w.put(r, c0 + 1, f'=IF({t}<=NTiers,$B{r}*SUM({ac(0)}{alloc}:{ac(60)}{alloc}),"")', F_USD3)
            w.put(r, c0 + 2, f'=IF({t}<=NTiers,(1-$B{r})*SUM({ac(0)}{alloc}:{ac(60)}{alloc}),"")', F_USD3)
            w.put(r, c0 + 3, f'=IF(ISNUMBER({L(c0)}{r}),INDEX(${ac(0)}${b["moic"]}:${ac(60)}${b["moic"]},{L(c0)}{r}+1),"")', F_X)
            w.put(r, c0 + 4, f'=IF(ISNUMBER({L(c0)}{r}),INDEX(${ac(0)}${b["irr"]}:${ac(60)}${b["irr"]},{L(c0)}{r}+1),"")', F_PCT)
    r = s + 12
    w.label(r, "After the last tier")
    w.put(r, 2, "=AfterPct", F_PCT)
    for k in range(4):
        b = blocks[k]; c0 = 4 + 5 * k
        w.put(r, c0 + 1, f"=AfterPct*$B${b['sp_after']}", F_USD3)
        w.put(r, c0 + 2, f"=(1-AfterPct)*$B${b['sp_after']}", F_USD3)
    r = s + 13
    w.label(r, "Total (ties to the investor and school shares)", key=True)
    for k in range(4):
        b = blocks[k]; c0 = 4 + 5 * k
        w.put(r, c0 + 1, f"=SUM({L(c0 + 1)}{s + 2}:{L(c0 + 1)}{s + 12})", F_USD3, key=True)
        w.put(r, c0 + 2, f"=SUM({L(c0 + 2)}{s + 2}:{L(c0 + 2)}{s + 12})", F_USD3, key=True)


# ============================================================================ Gross vs net
def build_gross_vs_net(B: Book, inp: Inputs):
    w = B.sheet(GN, 66)
    w.title("Gross vs net returns", "GROSS = no fees at all; the split is re-run on gross collections with targets on price only. "
                                    "NET = after all fees (from Loan economics). SPV costs are excluded from both.")
    le = B.ref["LE"]
    blocks = le["blocks"]
    G = {}
    s0 = 40
    for k in range(4):
        s = s0 + k * 52
        w.section(s, f"{inp.terms[k]}-month loans: split re-run on gross collections", width=12)
        w.label(s + 1, "Loan age"); w.header_row(s + 1, AGE0, list(range(NA)))
        w.label(s + 2, "Gross collections (no fees)", key=True)
        for a in range(NA):
            w.put(s + 2, AGE0 + a, f"={q(LE)}{ac(a)}{blocks[k]['gross']}", F_USD3)
        sp = write_split(w, s + 3, s + 2, f"{q(AS)}${L(2 + k)}$58", "Gross: ")
        cf = sp["end"]
        w.label(cf, "Investor cash flow, gross (month 0 = share - price)", key=True)
        for a in range(NA):
            c = ac(a)
            w.put(cf, AGE0 + a, f"={c}{sp['inv']}-{q(AS)}${L(2 + k)}$58" if a == 0 else f"={c}{sp['inv']}", F_USD3)
        G[k] = {"gross": s + 2, "inv": sp["inv"], "school": sp["school"], "cf": cf}
    bcf = s0 + 4 * 52 + 1
    w.label(bcf - 1, "Blended gross investor cash flow (mix-weighted; for the blended IRR)")
    w.label(bcf, "Investor cash flow, gross, blended", key=True)
    w.header_row(bcf - 1, AGE0, list(range(NA)))
    for a in range(NA):
        c = ac(a)
        w.put(bcf, AGE0 + a, "=" + "+".join(f"{q(AS)}${L(2 + k)}$20*{c}{G[k]['cf']}" for k in range(4)), F_USD3)
    w.put(bcf, 2, f"=SUM({ac(0)}{bcf}:{ac(60)}{bcf})", F_USD3, bold=True)
    # summary table at the top
    w.section(4, "Gross vs net, per $1 of face", width=7)
    w.header_row(5, 2, ["24m", "36m", "48m", "60m", "Blended"])
    rows = ["Mix weight (% of face)", "Price paid to the school (advance)", "Gross collections (no fees)", "Investor cash received, gross",
            "Investor MOIC, GROSS (before fees)", "Investor IRR, GROSS (before fees, annualised)", "Investor MOIC, NET (after fees)",
            "Investor IRR, NET (after fees)", "Fee drag on investor MOIC (gross - net)", "School total received, GROSS (% of face)",
            "School total received, NET (% of face)", "Check: gross investor + school - gross collections (0)"]
    fmts = [F_PCT, F_USD3, F_USD3, F_USD3, F_X, F_PCT, F_X, F_PCT, F_X, F_PCT, F_PCT, F_USD3]
    S = B.ref["LE_sum"]
    for i, lab in enumerate(rows):
        r = 6 + i
        key = "MOIC" in lab or "IRR" in lab
        w.label(r, lab, key=key)
        for j in range(5):
            col = 2 + j; cl = L(col)
            if j < 4:
                g = G[j]; tc = L(2 + j)
                gsum = lambda row: f"SUM({ac(0)}{row}:{ac(60)}{row})"
                f = [f"={q(AS)}${tc}$20", f"={q(AS)}${tc}$58", f"={gsum(g['gross'])}", f"={gsum(g['inv'])}",
                     f"={cl}9/{cl}7", f'=IFERROR((1+IRR({ac(0)}{g["cf"]}:{ac(60)}{g["cf"]},0.01))^12-1,"n.a.")',
                     f"={q(LE)}{cl}{S['moic']}", f"={q(LE)}{cl}{S['irr']}", f"={cl}10-{cl}12",
                     f"={cl}7+{gsum(g['school'])}", f"={q(LE)}{cl}{S['s_tot']}",
                     f"=ROUND({gsum(g['inv'])}+{gsum(g['school'])}-{gsum(g['gross'])},9)"][i]
            else:
                mixsum = lambda row: f"=SUMPRODUCT(B{row}:E{row},B6:E6)"
                f = ["=SUM(B6:E6)", mixsum(7), mixsum(8), mixsum(9), "=F9/F7",
                     f'=IFERROR((1+IRR({ac(0)}{bcf}:{ac(60)}{bcf},0.01))^12-1,"n.a.")', f"={q(LE)}F{S['moic']}",
                     f"={q(LE)}F{S['irr']}", "=F10-F12", mixsum(15), f"={q(LE)}F{S['s_tot']}", "=SUM(B17:E17)"][i]
            w.put(r, col, f, fmts[i], key=key)
    w.label(19, "Overall check")
    w.put(19, 2, '=IF(ABS(F17)<0.000001,"OK","CHECK")')
    B.name("ChkGross", GN, "B19")
    w.widths(2, 6, 13)
    w.widths(3, 3 + NA, 11)
    B.ref["GN"] = {"G": G, "bcf": bcf}


# ============================================================================ build
def build_book(inp: Inputs, saved: dict | None = None, sheet3=None, notes=None) -> Book:
    B = Book()
    build_assumptions(B, inp)
    build_loan_economics(B, inp)
    build_gross_vs_net(B, inp)
    build_pool(B, inp)
    build_spv(B, inp)
    build_program(B, inp)
    build_sens_advance(B, inp)
    build_sens_split(B, inp)
    build_checks(B, inp)
    build_overview(B, inp)
    build_charts(B, inp)
    build_readme(B, inp)
    build_compare(B, saved)
    build_notes(B, sheet3, notes)
    order = [RM, OV, AS, LE, GN, PL, SP, PD, CH, SA, SS, CK, "Compare scenarios", "Scenario notes", "Model notes"]
    B.wb._sheets = [B.wb[n] for n in order]
    B.wb.active = 1
    for ws in B.wb.worksheets:
        ws.sheet_properties.tabColor = {AS: "C8962E", CK: "2E7D4F", OV: "24527A"}.get(ws.title)
    B.wb.calculation.fullCalcOnLoad = True
    return B


def build_workbook(inp: Inputs, saved: dict | None = None, sheet3=None, notes=None) -> bytes:
    B = build_book(inp, saved, sheet3, notes)
    buf = io.BytesIO()
    B.wb.save(buf)
    return buf.getvalue()


# ============================================================================ Pool cash flows
def monthly_header(w: W, r: int, total=True):
    w.label(r, "Month of the projection")
    w.header_row(r, MON1, list(range(1, NM + 1)))
    if total:
        w.ws.cell(r, 2, "Total").font = Font(bold=True, color=NAVY)


def mrow(w: W, r: int, label, fn, fmt=F_USD, total="sum", key=False):
    """One monthly row: fn(m, col, prevcol) -> formula for month m."""
    w.label(r, label, key=key)
    for m in range(1, NM + 1):
        w.put(r, MON1 + m - 1, fn(m, mc(m), mc(m - 1) if m > 1 else None), fmt)
    if total == "sum":
        w.put(r, 2, f"=SUM({mc(1)}{r}:{mc(NM)}{r})", fmt, bold=True)
    elif total == "last":
        w.put(r, 2, f"={mc(NM)}{r}", fmt, bold=True)
    elif total == "abs":
        w.put(r, 2, f"=SUMPRODUCT(ABS({mc(1)}{r}:{mc(NM)}{r}))", fmt, bold=True)
    if key:
        w.key_row(r, 1, 2)


def build_pool(B: Book, inp: Inputs):
    w = B.sheet(PL, 64)
    ws = w.ws
    w.title("Pool cash flows", "Every monthly purchase is a vintage. Each vintage follows the blended per-$1 pattern from Loan economics "
                               "at its own age; the pool is the sum of all vintages. Then SPV fees, the debt facility and recycling.")
    bl = B.ref["LE"]["bl"]
    P = {}
    monthly_header(w, 3)
    w.label(4, "Month end")
    w.label(5, "Calendar year")
    for m in range(1, NM + 1):
        w.put(4, MON1 + m - 1, f"=EOMONTH(FirstMonth,{m - 1})", F_DATE)
        w.put(5, MON1 + m - 1, f"=YEAR({mc(m)}4)", "0")
    # helper rows (reversed per-$1 patterns) at the bottom
    H0 = 245
    hl = [("open", "open"), ("co", "co"), ("int", "int"), ("sch", "sch"), ("pre", "pre"), ("rec", "rec"), ("missed", "missed"),
          ("gross", "gross"), ("fees", "fees"), ("net", "net"), ("inv", "inv"), ("school", "school"), ("end", "end"),
          ("cost", "cost"), ("removed", "removed")]
    w.section(H0 - 1, "Helper rows: blended per-$1 pattern in reverse order (age 119 ... 0), used by the SUMPRODUCT stacking above", width=12)
    HR = {}
    for i, (nm, src) in enumerate(hl):
        r = H0 + i
        HR[nm] = r
        w.label(r, f"Reversed pattern: {nm}")
        for p in range(1, NM + 1):
            age = 120 - p
            w.put(r, MON1 + p - 1, f"={q(LE)}{ac(age)}{bl[src]}" if age <= 60 else 0, '0.000000')
    stack = lambda nm: (lambda m, c, p: f"=SUMPRODUCT($C$8:{c}$8,INDEX($C${HR[nm]}:$DR${HR[nm]},1,{121 - m}):$DR${HR[nm]})")

    w.section(7, "1. Purchases", width=12)
    mrow(w, 8, "Face purchased ($)", lambda m, c, p: f"=IF({m}<=Period,INDEX({q(AS)}$B$113:$BI$113,{m}),0)" if m <= 60 else "=0")
    mrow(w, 9, "Cash paid to schools (advance) ($)", lambda m, c, p: f"={c}8*AdvBlended")
    mrow(w, 10, "Upfront fees paid by the investor ($)", lambda m, c, p: f"={c}8*UpfrontBlended")
    mrow(w, 11, "Investor outlay for purchases ($)", lambda m, c, p: f"={c}9+{c}10", key=True)
    w.section(13, "2. Pool loan cash flows (all vintages and terms)", width=12)
    rows2 = [(14, "open", "Opening performing balance ($)", None), (15, "co", "Charge-offs ($)", "sum"), (16, "int", "Interest collected ($)", "sum"),
             (17, "sch", "Scheduled principal collected ($)", "sum"), (18, "pre", "Prepaid principal ($)", "sum"), (19, "rec", "Recoveries ($)", "sum"),
             (20, "missed", "Payments missed before charge-off ($)", "sum"), (21, "gross", "Gross collections ($)", "sum"),
             (22, "fees", "Across-life fees ($)", "sum"), (23, "net", "Net collections ($)", "sum"), (24, "inv", "Investor share of net collections ($)", "sum"),
             (25, "school", "School share of net collections ($)", "sum"), (26, "end", "Ending performing balance ($)", "last"),
             (27, "cost", "Investor cost basis ($) (advance rate x performing balance)", None), (28, "removed", "Balance removed at charge-off ($)", "sum")]
    for r, nm, lab, tot in rows2:
        mrow(w, r, lab, stack(nm), total=tot, key=nm in ("net", "inv"))
        P[nm] = r
    mrow(w, 29, "Check: gross = interest + scheduled + prepaid + recoveries - missed (0)",
         lambda m, c, p: f"=ROUND({c}21-({c}16+{c}17+{c}18+{c}19-{c}20),4)", total="abs")
    mrow(w, 30, "Check: investor + school - net collections (0)", lambda m, c, p: f"=ROUND({c}24+{c}25-{c}23,4)", total="abs")
    mrow(w, 31, "Check: balance roll (0)", lambda m, c, p: f"=ROUND({c}26-({'0' if m == 1 else p + '26'}+{c}8-{c}28-{c}17-{c}18),4)", total="abs")
    w.section(33, "3. School", width=12)
    mrow(w, 34, "Cash to school at purchase (advance) ($)", lambda m, c, p: f"={c}9")
    mrow(w, 35, "School share of net collections ($)", lambda m, c, p: f"={c}25")
    mrow(w, 36, "Total cash received by the school ($)", lambda m, c, p: f"={c}34+{c}35", key=True)
    mrow(w, 37, "Cumulative cash received by the school ($)", lambda m, c, p: f"=SUM($C$36:{c}36)", total="last")
    w.section(39, "4. Investor cash before the SPV structure (unlevered, no other fees)", width=12)
    mrow(w, 40, "Investor share of net collections ($)", lambda m, c, p: f"={c}24")
    mrow(w, 41, "Purchase outlay ($)", lambda m, c, p: f"=-{c}11")
    mrow(w, 42, "Net investor cash flow, unlevered ($)", lambda m, c, p: f"={c}40+{c}41", key=True)
    mrow(w, 43, "Cumulative net cash flow, unlevered ($)", lambda m, c, p: f"=SUM($C$42:{c}42)", total="last")
    w.section(45, "5. Other SPV fees", width=12)
    mrow(w, 46, "Pool active (1 = yes)", lambda m, c, p: f"=IF(OR({c}8>0,{c}26>0.5,{c}23>0.5),1,0)", F_INT)
    mrow(w, 47, "Other SPV fees ($) (annual / 12, capped at investor collections less interest; unpaid fees are not carried)",
         lambda m, c, p: f"=IF({c}63<=0,0,-({c}46*MIN(IF(ROUNDUP({m}/12,0)<=FeeYearsHigh,FeeHigh,FeeLow)/12,MAX(0,{c}63-{c}57))))")
    w.section(49, "6. Leverage facility", width=12)
    mrow(w, 50, "Facility mode", lambda m, c, p: f'=IF(Leverage="OFF","Off",IF({m}<FacStart,"Pre-raise",IF({m}<=LastDraw,"Origination","Amortization")))',
         None, total=None)
    mrow(w, 51, "Investor cost basis ($)", lambda m, c, p: f"={c}27", total=None)
    mrow(w, 52, "Eligible cost basis ($)", lambda m, c, p: f"={c}51*(1-Inelig)", total=None)
    mrow(w, 53, "Borrowing capacity ($)", lambda m, c, p: f'=IF(Leverage="ON",MIN(LTV*{c}52,FacilityMax),0)', total=None)
    mrow(w, 54, "SOFR (annual)", lambda m, c, p: f"=INDEX({q(AS)}$B$132:$DQ$132,{m})", F_PCT, total=None)
    mrow(w, 55, "Interest rate (SOFR + spread)", lambda m, c, p: f"={c}54+Spread", F_PCT, total=None)
    mrow(w, 56, "Opening debt ($)", lambda m, c, p: "=0" if m == 1 else f"={p}60", total=None)
    mrow(w, 57, "Interest on debt ($)", lambda m, c, p: f'=IF({c}50="Origination",({c}56+{c}56+{c}58-MAX(0,{c}56-{c}53))/2,{c}56)*{c}55/12')
    mrow(w, 58, "Debt draws ($)", lambda m, c, p: f'=IF({c}50="Origination",IF(Recycling="ON",MAX(0,{c}53-{c}56),MIN(MAX(0,{c}53-{c}56),{c}11)),0)')
    mrow(w, 59, "Debt repayments ($)", lambda m, c, p: f'=IF({c}50="Origination",MAX(0,{c}56-{c}53),IF({c}50="Amortization",MIN({c}56,MAX(0,{c}56-{c}53,Sweep*({c}63+{c}47-{c}57))),0))')
    mrow(w, 60, "Ending debt ($)", lambda m, c, p: f"={c}56+{c}58-{c}59", total="last")
    w.section(62, "7. SPV cash and recycling", width=12)
    mrow(w, 63, "Investor share of net collections ($)", lambda m, c, p: f"={c}24")
    mrow(w, 64, "Debt draws ($)", lambda m, c, p: f"={c}58")
    mrow(w, 65, "Purchase outlay ($)", lambda m, c, p: f"=-{c}11")
    mrow(w, 66, "Other SPV fees ($)", lambda m, c, p: f"={c}47")
    mrow(w, 67, "Interest on debt ($)", lambda m, c, p: f"=-{c}57")
    mrow(w, 68, "Debt repayments ($)", lambda m, c, p: f"=-{c}59")
    mrow(w, 69, "Net SPV cash before equity ($)", lambda m, c, p: f"=SUM({c}63:{c}68)")
    mrow(w, 70, "Operating cash (collections + fees - interest - repayments) ($)", lambda m, c, p: f"={c}63+{c}66+{c}67+{c}68")
    mrow(w, 71, "Equity called from investors ($)", lambda m, c, p: f'=IF(Recycling="ON",MAX(0,-{c}69),MAX(0,-({c}65+{c}64))+MAX(0,-{c}70))', key=True)
    mrow(w, 72, "Distributions to investors ($)", lambda m, c, p: f'=IF(Recycling="ON",MAX(0,{c}69),MAX(0,{c}70)+MAX(0,{c}64+{c}65))', key=True)
    mrow(w, 73, "Collections recycled into purchases ($)", lambda m, c, p: f'=IF(Recycling="ON",MIN(MAX(0,{c}70),MAX(0,-({c}65+{c}64))),0)')
    mrow(w, 74, "Check: distributions - equity - net SPV cash (0)", lambda m, c, p: f"=ROUND({c}72-{c}71-{c}69,4)", total="abs")
    mrow(w, 75, "Net cash flow to investors ($)", lambda m, c, p: f"={c}72-{c}71", key=True)
    mrow(w, 76, "Purchases funded by recycled collections ($)", lambda m, c, p: f"={c}73")
    mrow(w, 77, "Purchases funded by debt draws ($)", lambda m, c, p: f"=MIN(MAX(0,{c}58),MAX(0,{c}11-{c}73))")
    mrow(w, 78, "Purchases funded by new equity ($)", lambda m, c, p: f"=MAX(0,{c}11-{c}76-{c}77)")
    P.update(face=8, cash=9, upf=10, outlay=11, school_tot=36, unlev=42, active=46, fees_spv=47, mode=50, cap=53, debt_open=56,
             interest=57, draws=58, repay=59, debt_end=60, equity=71, dist=72, recycled=73, net_inv=75, fund_rec=76, fund_debt=77, fund_eq=78)
    # summary
    w.section(80, "8. Pool summary", width=4)
    sums = [("Total face purchased ($)", "=B8", F_USD), ("Cash paid to schools ($)", "=B9", F_USD), ("Upfront fees paid by the investor ($)", "=B10", F_USD),
            ("Gross collections ($)", "=B21", F_USD), ("Across-life fees ($)", "=B22", F_USD), ("Net collections ($)", "=B23", F_USD),
            ("Investor share of net collections ($)", "=B24", F_USD), ("School share of net collections ($)", "=B25", F_USD),
            ("Charge-offs ($)", "=B15", F_USD), ("Charge-offs (% of face)", "=IF(B8>0,B15/B8,0)", F_PCT),
            ("Unlevered investor MOIC", "=IF(B11>0,B24/B11,0)", F_X), ("Unlevered investor IRR", f'=IFERROR((1+IRR(C42:DR42,0.01))^12-1,"n.a.")', F_PCT),
            ("Tie-out: unlevered MOIC and IRR vs Loan economics, blended (0)", f"=ROUND(B91-{q(LE)}F{B.ref['LE_sum']['moic']},6)+ROUND(B92-{q(LE)}F{B.ref['LE_sum']['irr']},6)", F_NUM),
            ("School effective advance rate, undiscounted (% of face)", "=IF(B8>0,B36/B8,0)", F_PCT),
            ("School effective advance rate, PV (% of face)", "=IF(B8>0,(C36+NPV((1+SchoolPV)^(1/12)-1,D36:DR36))/B8,0)", F_PCT),
            ("Peak debt ($)", "=MAX(C60:DR60)", F_USD), ("Total debt draws ($)", "=B58", F_USD), ("Total interest ($)", "=B57", F_USD),
            ("Equity called ($)", "=B71", F_USD), ("Distributions ($)", "=B72", F_USD), ("Collections recycled into purchases ($)", "=B73", F_USD),
            ("Equity MOIC", "=IF(B71>0,B72/B71,0)", F_X), ("Equity IRR", '=IFERROR((1+IRR(C75:DR75,0.01))^12-1,"n.a.")', F_PCT)]
    PS = {}
    for i, (lab, f, fmt) in enumerate(sums):
        r = 81 + i
        key = any(x in lab for x in ("MOIC", "IRR", "Net collections", "Equity called", "Distributions ($)"))
        w.label(r, lab, key=key)
        w.put(r, 2, f, fmt, key=key)
        PS[lab] = r
    rchk = 81 + len(sums)
    w.label(rchk, "Overall pool check (identities, debt repaid, tie-out, vintage grids)", key=True)
    w.put(rchk, 2, '=IF(AND(B29+B30+B31+B74=0,ABS(DR60)<1,B93=0,ABS(B172)+ABS(B237)<1),"OK","CHECK")', key=True)
    B.name("ChkPool", PL, f"B{rchk}")
    # vintage grids
    for (title, r0, src, tot, chkrow) in [("9. Vintage view: investor share of net collections by monthly vintage ($)", 108, bl["inv"], 24, 172),
                                          ("10. Vintage view: school cash by vintage (advance at purchase + share of collections) ($)", 173, bl["scf"], 36, 237)]:
        w.section(r0, title, width=12)
        monthly_header(w, r0 + 1)
        for v in range(1, 61):
            r = r0 + 1 + v
            w.label(r, f"Vintage {v}")
            for m in range(1, NM + 1):
                f = f"=IF(AND({m}>={v},{m}-{v}<=60),{mc(v)}$8*{q(LE)}{ac(0)}{src}:{ac(60)}{src})" if False else \
                    (f"={mc(v)}$8*INDEX({q(LE)}${ac(0)}${src}:${ac(60)}${src},{m - v + 1})" if (m >= v and m - v <= 60) else 0)
                w.put(r, MON1 + m - 1, f, F_USD)
            w.put(r, 2, f"=SUM({mc(1)}{r}:{mc(NM)}{r})", F_USD)
        rt = r0 + 62
        w.label(rt, "Total of vintages ($)", key=True)
        for m in range(1, NM + 1):
            w.put(rt, MON1 + m - 1, f"=SUM({mc(m)}{r0 + 2}:{mc(m)}{r0 + 61})", F_USD)
        w.put(rt, 2, f"=SUM({mc(1)}{rt}:{mc(NM)}{rt})", F_USD, bold=True)
        w.key_row(rt, 1, 2)
        rc = rt + 1
        w.label(rc, "Check: total of vintages - pool row (0)")
        for m in range(1, NM + 1):
            w.put(rc, MON1 + m - 1, f"=ROUND({mc(m)}{rt}-{mc(m)}{tot},4)", F_USD)
        w.put(rc, 2, f"=SUMPRODUCT(ABS({mc(1)}{rc}:{mc(NM)}{rc}))", F_USD, bold=True)
        P[f"vin_{r0}"] = r0 + 2
    # fix check row numbers used in the overall check (investor grid check at 171, school grid check at 236)
    ws.cell(rchk, 2).value = '=IF(AND(B29+B30+B31+B74=0,ABS(DR60)<1,B93=0,B171+B236<1),"OK","CHECK")'
    w.widths(2, 2, 15)
    w.widths(3, 3 + NM, 13)
    ws.freeze_panes = "C6"
    B.ref["PL"] = P
    B.ref["PL_sum"] = PS
    B.ref["PL_vin"] = {"inv": 110, "school": 175}


# ============================================================================ SPV and waterfall
def build_spv(B: Book, inp: Inputs):
    w = B.sheet(SP, 66)
    ws = w.ws
    pl = f"{q(PL)}"
    w.title("SPV and waterfall", "Equity calls and distributions from Pool cash flows, split between OCM and Loan America by capital "
                                 "and the promote waterfall, with returns by party.")
    w.section(5, "1. Capital split and waterfall settings", width=5)
    w.label(6, "Capital basis"); w.put(6, 2, "=CapitalBasis")
    w.label(7, "Loan America share of capital (OCM % basis: 1 - OCM %; Fixed $ basis: commitment / total equity called)", key=True)
    w.put(7, 2, f'=IF(CapitalBasis="OCM %",1-OCMPct,IFERROR(MIN(1,LAFixed/{pl}B71),0))', F_PCT, key=True); B.name("LAShareCap", SP, "B7")
    w.label(8, "OCM share of capital"); w.put(8, 2, "=1-LAShareCap", F_PCT); B.name("OCMShareCap", SP, "B8")
    w.label(9, "Promote waterfall"); w.put(9, 2, "=Promote")
    w.label(10, "Monthly hurdle rates: tier 2 pref, tier 3, tier 4")
    for j, h in enumerate(["Hurdle2", "Hurdle3", "Hurdle4"]):
        w.put(10, 2 + j, f"=(1+{h})^(1/12)-1", '0.0000%')
    B.name("H2m", SP, "B10"); B.name("H3m", SP, "C10"); B.name("H4m", SP, "D10")
    w.label(11, "Effective Loan America share in tiers 3, 4, 5 (after the promote method)")
    for j, s_ in enumerate(["LAShare3", "LAShare4", "LAShare5"]):
        w.put(11, 2 + j, f'=IF(PromoteMethod="On top of pro rata",{s_}+(1-{s_})*LAShareCap,{s_})', F_PCT)
    B.name("S3e", SP, "B11"); B.name("S4e", SP, "C11"); B.name("S5e", SP, "D11")
    w.label(12, "Check: effective Loan America shares within 0-100%")
    w.put(12, 2, '=IF(AND(MIN(B11:D11)>=0,MAX(B11:D11)<=1),"OK","CHECK")'); B.name("ChkShares", SP, "B12")
    w.label(13, "Check: tiers add to distributions; OCM + Loan America = total", key=True)
    w.put(13, 2, '=IF(AND(ABS(SUM(B30:B34)-B18)<1,ABS(C18+D18-B18)<1,ABS(C17+D17-B17)<1,B61<0.001,B66<0.001),"OK","CHECK")', key=True)
    B.name("ChkWaterfall", SP, "B13")
    # monthly
    monthly_header(w, 37)
    w.label(38, "Month end")
    for m in range(1, NM + 1):
        w.put(38, MON1 + m - 1, f"={pl}{mc(m)}4", F_DATE)
    w.section(39, "3. Capital called (monthly)", width=12)
    mrow(w, 40, "Equity called from investors ($)", lambda m, c, p: f"={pl}{c}71", key=True)
    mrow(w, 41, "   OCM share ($)", lambda m, c, p: f"={c}40*OCMShareCap")
    mrow(w, 42, "   Loan America share ($)", lambda m, c, p: f"={c}40*LAShareCap")
    mrow(w, 43, "Cumulative equity called ($)", lambda m, c, p: f"=SUM($C$40:{c}40)", total="last")
    w.section(44, "4. Promote waterfall on investor distributions (monthly)", width=12)
    mrow(w, 45, "Distributable cash to investors ($)", lambda m, c, p: f"={pl}{c}72", key=True)
    mrow(w, 46, "Tier 1: return of capital ($)", lambda m, c, p: f"=MIN({c}45,MAX(0,{c}43-{'0' if m == 1 else p + '47'}))")
    mrow(w, 47, "Cumulative tier 1 paid ($)", lambda m, c, p: f"=SUM($C$46:{c}46)", total="last")
    mrow(w, 48, "Pref hurdle balance before distribution ($) (equity compounded at the pref, less distributions)",
         lambda m, c, p: f"={c}40" if m == 1 else f"={p}50*(1+H2m)+{c}40", total=None)
    mrow(w, 49, "Tiers 1 + 2: capital plus pref paid ($)", lambda m, c, p: f"=MAX({c}46,MIN({c}45,MAX(0,{c}48)))")
    mrow(w, 50, "Pref hurdle balance after distribution ($)", lambda m, c, p: f"={c}48-{c}45", total=None)
    mrow(w, 51, "Tier 2: preferred return ($)", lambda m, c, p: f"={c}49-{c}46")
    mrow(w, 52, "OCM share of tiers 1 + 2 ($)", lambda m, c, p: f"={c}49*OCMShareCap")
    mrow(w, 53, "OCM hurdle balance at the tier 3 IRR, before distribution ($)",
         lambda m, c, p: f"={c}41" if m == 1 else f"={p}60*(1+H3m)+{c}41", total=None)
    mrow(w, 54, "Tier 3: until OCM reaches the tier 3 IRR ($)",
         lambda m, c, p: f'=IF(Promote="ON",IF(S3e<1,MIN({c}45-{c}49,MAX(0,{c}53-{c}52)/(1-S3e)),{c}45-{c}49),0)')
    mrow(w, 55, "OCM hurdle balance at the tier 4 IRR, before distribution ($)",
         lambda m, c, p: f"={c}41" if m == 1 else f"={p}61*(1+H4m)+{c}41", total=None)
    mrow(w, 56, "Tier 4: until OCM reaches the tier 4 IRR ($)",
         lambda m, c, p: f'=IF(Promote="ON",IF(S4e<1,MIN({c}45-{c}49-{c}54,MAX(0,{c}55-{c}52-{c}54*(1-S3e))/(1-S4e)),{c}45-{c}49-{c}54),0)')
    mrow(w, 57, "Tier 5: residual ($)", lambda m, c, p: f"={c}45-{c}49-{c}54-{c}56")
    mrow(w, 58, "Distributions to OCM ($)", lambda m, c, p: f'=IF(Promote="ON",{c}52+{c}54*(1-S3e)+{c}56*(1-S4e)+{c}57*(1-S5e),{c}45*OCMShareCap)', key=True)
    mrow(w, 59, "Distributions to Loan America ($)", lambda m, c, p: f"={c}45-{c}58", key=True)
    mrow(w, 60, "OCM hurdle balance at the tier 3 IRR, after distribution ($)", lambda m, c, p: f"={c}53-{c}58", total=None)
    mrow(w, 61, "OCM hurdle balance at the tier 4 IRR, after distribution ($)", lambda m, c, p: f"={c}55-{c}58", total=None)
    mrow(w, 62, "Highest tier paid this month (0 = no distribution)",
         lambda m, c, p: f"=IF({c}45<=0,0,IF({c}57>0.005,5,IF({c}56>0.005,4,IF({c}54>0.005,3,IF({c}51>0.005,2,1)))))", F_INT, total=None)
    w.section(63, "5. Loan America distributions by tier", width=12)
    mrow(w, 64, "Loan America: tiers 1 + 2 (pro rata) ($)", lambda m, c, p: f"={c}49*LAShareCap")
    mrow(w, 65, "Loan America: tier 3 ($)", lambda m, c, p: f"={c}54*S3e")
    mrow(w, 66, "Loan America: tier 4 ($)", lambda m, c, p: f"={c}56*S4e")
    mrow(w, 67, "Loan America: tier 5 ($)", lambda m, c, p: f'={c}57*IF(Promote="ON",S5e,LAShareCap)')
    mrow(w, 68, "Check: Loan America by tier - Loan America total (0)", lambda m, c, p: f"=ROUND({c}64+{c}65+{c}66+{c}67-{c}59,4)", total="abs")
    w.section(69, "6. Net cash flows and running returns", width=12)
    mrow(w, 70, "OCM net cash flow ($)", lambda m, c, p: f"={c}58-{c}41", key=True)
    mrow(w, 71, "Loan America net cash flow ($)", lambda m, c, p: f"={c}59-{c}42", key=True)
    mrow(w, 72, "Total investor net cash flow ($)", lambda m, c, p: f"={pl}{c}75", key=True)
    mrow(w, 73, "Check: OCM + Loan America - total (0)", lambda m, c, p: f"=ROUND({c}70+{c}71-{c}72,4)", total="abs")
    mrow(w, 74, "OCM running MOIC", lambda m, c, p: f"=IF(SUM($C$41:{c}41)>0,SUM($C$58:{c}58)/SUM($C$41:{c}41),0)", F_X, total="last")
    mrow(w, 75, "Loan America running MOIC", lambda m, c, p: f"=IF(SUM($C$42:{c}42)>0,SUM($C$59:{c}59)/SUM($C$42:{c}42),0)", F_X, total="last")
    mrow(w, 76, "Total investor running MOIC", lambda m, c, p: f"=IF(SUM($C$40:{c}40)>0,SUM($C$45:{c}45)/SUM($C$40:{c}40),0)", F_X, total="last")
    mrow(w, 77, "OCM running IRR (annualised)", lambda m, c, p: f'=IFERROR((1+IRR($C$70:{c}70,0.01))^12-1,"")', F_PCT, total=None)
    mrow(w, 78, "Loan America running IRR (annualised)", lambda m, c, p: f'=IFERROR((1+IRR($C$71:{c}71,0.01))^12-1,"")', F_PCT, total=None)
    mrow(w, 79, "Total investor running IRR (annualised)", lambda m, c, p: f'=IFERROR((1+IRR($C$72:{c}72,0.01))^12-1,"")', F_PCT, total=None)
    mrow(w, 80, "Debt recap (memo): distributions in the facility start month funded by the first draw ($)",
         lambda m, c, p: f'=IF(Recycling="OFF",0,IF({m}=FacStart,MIN(MAX({c}45,0),MAX({pl}{c}58,0)),0))')
    # section 8, 9
    w.section(82, "8. Distributions over time: OCM vs Loan America (pro rata vs promote)", width=12)
    mrow(w, 83, "Distributions to OCM ($)", lambda m, c, p: f"={c}58")
    mrow(w, 84, "Loan America: pro rata share of distributions ($)", lambda m, c, p: f"={c}59-{c}85")
    mrow(w, 85, "Loan America: promote above pro rata ($)", lambda m, c, p: f"=MAX(0,{c}65+{c}66+{c}67-({c}54+{c}56+{c}57)*LAShareCap)")
    mrow(w, 86, "Loan America share of the month's distributions (%)", lambda m, c, p: f'=IF({c}45>0.5,{c}59/{c}45,"")', F_PCT, total=None)
    mrow(w, 87, "Cumulative distributions to OCM ($)", lambda m, c, p: f"=SUM($C$83:{c}83)", total="last")
    mrow(w, 88, "Cumulative distributions to Loan America ($)", lambda m, c, p: f"=SUM($C$59:{c}59)", total="last")
    w.section(90, "9. Equity calls and distributions by party (SPV / OCM / Loan America)", width=12)
    mrow(w, 91, "SPV: equity called (shown negative) ($)", lambda m, c, p: f"=-{c}40")
    mrow(w, 92, "SPV: distributions ($)", lambda m, c, p: f"={c}45")
    mrow(w, 93, "SPV: cumulative net cash flow ($)", lambda m, c, p: f"=SUM($C$91:{c}92)", total="last")
    mrow(w, 94, "OCM: equity called (shown negative) ($)", lambda m, c, p: f"=-{c}41")
    mrow(w, 95, "OCM: distributions ($)", lambda m, c, p: f"={c}58")
    mrow(w, 96, "OCM: cumulative net cash flow ($)", lambda m, c, p: f"=SUM($C$94:{c}95)", total="last")
    mrow(w, 97, "Loan America: equity called (shown negative) ($)", lambda m, c, p: f"=-{c}42")
    mrow(w, 98, "Loan America: distributions, pro rata share ($)", lambda m, c, p: f"={c}84")
    mrow(w, 99, "Loan America: distributions, promote ($)", lambda m, c, p: f"={c}85")
    mrow(w, 100, "Loan America: cumulative net cash flow ($)", lambda m, c, p: f"=SUM($C$97:{c}99)", total="last")
    # section 7 yearly
    w.section(102, "7. Yearly view: equity in and distributions out by party (calendar years)", width=12)
    w.label(103, "Calendar year")
    for j in range(10):
        w.put(103, 3 + j, f"=YEAR(FirstMonth)+{j}", "0", bold=True)
    w.ws.cell(103, 2, "Total").font = Font(bold=True, color=NAVY)
    yrow = lambda srow, sign="": (lambda j: f"={sign}SUMPRODUCT(({pl}$C$5:$DR$5={L(3 + j)}$103)*$C${srow}:$DR${srow})")
    ylines = [(104, "Total investors: equity called ($)", yrow(40, "-")), (105, "Total investors: distributions ($)", yrow(45)),
              (106, "Total investors: net cash flow ($)", lambda j: f"={L(3 + j)}104+{L(3 + j)}105"),
              (107, "Total investors: cumulative net cash flow ($)", lambda j: f"=SUM($C$106:{L(3 + j)}106)"),
              (108, "OCM: equity called ($)", yrow(41, "-")), (109, "OCM: distributions ($)", yrow(58)),
              (110, "OCM: net cash flow ($)", lambda j: f"={L(3 + j)}108+{L(3 + j)}109"),
              (111, "OCM: cumulative net cash flow ($)", lambda j: f"=SUM($C$110:{L(3 + j)}110)"),
              (112, "Loan America: equity called ($)", yrow(42, "-")), (113, "Loan America: distributions ($)", yrow(59)),
              (114, "Loan America: of which promote above pro rata ($)",
               lambda j: f"=SUMPRODUCT(({pl}$C$5:$DR$5={L(3 + j)}$103)*($C$65:$DR$65+$C$66:$DR$66+$C$67:$DR$67-($C$54:$DR$54+$C$56:$DR$56+$C$57:$DR$57)*LAShareCap))"),
              (115, "Loan America: net cash flow ($)", lambda j: f"={L(3 + j)}112+{L(3 + j)}113"),
              (116, "Loan America: cumulative net cash flow ($)", lambda j: f"=SUM($C$115:{L(3 + j)}115)"),
              (117, "Loan America share of the year's distributions", lambda j: f"=IF({L(3 + j)}105<>0,{L(3 + j)}113/{L(3 + j)}105,0)"),
              (118, "Check: OCM + Loan America - total (0)", lambda j: f"=ROUND({L(3 + j)}109+{L(3 + j)}113-{L(3 + j)}105,2)")]
    for r, lab, fn in ylines:
        key = "net cash flow" in lab and "cumulative" not in lab
        w.label(r, lab, key=key)
        fmt = F_PCT if "share" in lab else F_USD
        for j in range(10):
            w.put(r, 3 + j, fn(j), fmt)
        if "cumulative" in lab:
            w.put(r, 2, f"=L{r}", fmt, bold=True)
        elif "share" in lab:
            w.put(r, 2, "=IF(B105<>0,B113/B105,0)", fmt, bold=True)
        else:
            w.put(r, 2, f"=SUM(C{r}:L{r})", fmt, bold=True)
    # section 2 returns table (top)
    w.section(15, "2. Returns by party", width=5)
    w.header_row(16, 2, ["Total investors", "OCM", "Loan America"])
    sums = lambda row: f"SUM($C${row}:$DR${row})"
    irr = lambda row: f'IFERROR((1+IRR($C${row}:$DR${row},0.01))^12-1,"n.a.")'
    table = [(17, "Equity called ($)", [sums(40), sums(41), sums(42)], F_USD), (18, "Distributions ($)", [sums(45), sums(58), sums(59)], F_USD),
             (19, "Net profit ($)", ["B18-B17", "C18-C17", "D18-D17"], F_USD),
             (20, "MOIC", ["IF(B17>0,B18/B17,0)", "IF(C17>0,C18/C17,0)", "IF(D17>0,D18/D17,0)"], F_X),
             (21, "IRR (annualised)", [irr(72), irr(70), irr(71)], F_PCT),
             (22, "Share of capital", ["1", "OCMShareCap", "LAShareCap"], F_PCT),
             (23, "Share of distributions", ["1", "IF(B18>0,C18/B18,0)", "IF(B18>0,D18/B18,0)"], F_PCT),
             (24, "Debt recap ($): distributions in the facility start month funded by the first draw", [sums(80), "B24*OCMShareCap", "B24*LAShareCap"], F_USD),
             (25, "Adjusted MOIC (debt recap netted from both equity called and distributions)",
              ['IFERROR((B18-B24)/(B17-B24),"n.a.")', 'IFERROR((C18-C24)/(C17-C24),"n.a.")', 'IFERROR((D18-D24)/(D17-D24),"n.a.")'], F_X)]
    for r, lab, fs, fmt in table:
        key = lab.startswith(("Equity", "Distributions", "Net profit", "MOIC", "IRR", "Share of distributions", "Adjusted"))
        w.label(r, lab, key=key)
        for j, f in enumerate(fs):
            w.put(r, 2 + j, "=" + f, fmt, key=key)
    w.label(29, "Distributions by tier ($)")
    w.header_row(29, 2, ["Total investors", "OCM", "Loan America"])
    tiers = [(30, "1. Return of capital", sums(46), f"{sums(46)}*OCMShareCap", f"{sums(46)}*LAShareCap"),
             (31, "2. Preferred return", sums(51), f"{sums(51)}*OCMShareCap", f"{sums(51)}*LAShareCap"),
             (32, "3. Until OCM reaches the tier 3 IRR", sums(54), f"{sums(54)}*(1-S3e)", sums(65)),
             (33, "4. Until OCM reaches the tier 4 IRR", sums(56), f"{sums(56)}*(1-S4e)", sums(66)),
             (34, "5. Residual", sums(57), f'{sums(57)}*IF(Promote="ON",1-S5e,OCMShareCap)', sums(67))]
    for r, lab, a, b, c in tiers:
        w.label(r, lab)
        for j, f in enumerate([a, b, c]):
            w.put(r, 2 + j, "=" + f, F_USD)
    w.label(35, "Total of tiers ($)", key=True)
    for j in range(3):
        w.put(35, 2 + j, f"=SUM({L(2 + j)}30:{L(2 + j)}34)", F_USD, key=True)
    # check cells referenced by B13
    ws.cell(13, 2).value = '=IF(AND(ABS(B35-B18)<1,ABS(C18+D18-B18)<1,ABS(C17+D17-B17)<1,B68<0.001,B73<0.001),"OK","CHECK")'
    w.widths(2, 4, 16)
    w.widths(5, 3 + NM, 13)
    ws.freeze_panes = "C39"
    B.ref["SP"] = {"ret": 17, "tiers": 30}


# ============================================================================ Program detail
def build_program(B: Book, inp: Inputs):
    w = B.sheet(PD, 62)
    pl, sp, le = q(PL), q(SP), q(LE)
    w.title("Program detail", "Returns by party, the school's effective advance rate, a monthly view by party, and the recycling and leverage views.")
    w.section(4, "1. Deal returns by party (program level)", width=7)
    w.header_row(5, 2, ["OCM", "Loan America", "Total investors", "School", "Goal (servicer)"])
    rows = [(6, "Cash put in ($)", [f"={sp}C17", f"={sp}D17", f"={sp}B17", '=""', '=""'], F_USD),
            (7, "Cash received ($)", [f"={sp}C18", f"={sp}D18", f"={sp}B18", f"={pl}B36", f"={pl}B22"], F_USD),
            (8, "Net profit ($)", [f"={sp}C19", f"={sp}D19", f"={sp}B19", '=""', '=""'], F_USD),
            (9, "MOIC", [f"={sp}C20", f"={sp}D20", f"={sp}B20", '=""', '=""'], F_X),
            (10, "IRR (annualised)", [f"={sp}C21", f"={sp}D21", f"={sp}B21", '=""', '=""'], F_PCT),
            (11, "Share of capital", [f"={sp}C22", f"={sp}D22", "=1", '=""', '=""'], F_PCT),
            (12, "Loan America: origination / upfront fees received ($)", ['=""', f"={pl}B10", '=""', '=""', '=""'], F_USD)]
    for r, lab, fs, fmt in rows:
        key = lab.startswith(("Cash received", "Net profit", "MOIC", "IRR"))
        w.label(r, lab, key=key)
        for j, f in enumerate(fs):
            w.put(r, 2 + j, f, fmt, key=key)
    tiers_txt = "&".join([f'IF(NTiers>={k},"; T{k} "&TEXT({q(AS)}J{85 + k},"0.00%")&" to "&IF(TargetMode="MOIC",TEXT({q(AS)}K{85 + k},"0.000")&"x",TEXT({q(AS)}K{85 + k},"0.00%")),"")' for k in range(1, 11)])
    w.label(13, "Key settings")
    w.put(13, 2, f'="Targets: "&TargetMode&" | after last tier: "&AfterLast&" | capital: "&CapitalBasis&" | promote: "&Promote&" | recycling: "&Recycling&" | leverage: "&Leverage&" | tiers: "&MID({tiers_txt},3,2000)')
    w.section(15, "2. School: effective advance rate (program level and per $1 by term)", width=7)
    w.header_row(16, 2, ["Program"] + [f"{t}m" for t in inp.terms] + ["Blended per $1"])
    S = B.ref["LE_sum"]
    cols = ["B", "C", "D", "E", "F"]
    srows = [(17, "Face of loans sold ($ / per $1)", [f"={pl}B8"] + ["=1"] * 5, F_USD3),
             (18, "Upfront advance at month 0 (% of face)", [f"=IF({pl}B8>0,{pl}B9/{pl}B8,0)"] + [f"={le}{c}{S['s_adv']}" for c in cols], F_PCT),
             (19, "Share of collections, undiscounted (% of face)", [f"=IF({pl}B8>0,{pl}B25/{pl}B8,0)"] + [f"={le}{c}{S['s_coll']}" for c in cols], F_PCT),
             (20, "Effective advance rate, undiscounted (% of face)", [f"={pl}B94"] + [f"={le}{c}{S['s_tot']}" for c in cols], F_PCT),
             (21, "Effective advance rate, PV (% of face)", [f"={pl}B95"] + [f"={le}{c}{S['s_pv']}" for c in cols], F_PCT),
             (22, "Discount rate (annual)", ["=SchoolPV"] * 6, F_PCT)]
    for r, lab, fs, fmt in srows:
        key = "Effective" in lab
        w.label(r, lab, key=key)
        for j, f in enumerate(fs):
            w.put(r, 2 + j, f, fmt if not (r == 17 and j == 0) else F_USD, key=key)
    w.section(24, "3. Investor per $1 of face by term (before SPV fees, leverage and recycling)", width=7)
    w.header_row(25, 2, [f"{t}m" for t in inp.terms] + ["Blended"])
    prows = [(26, "Investor outlay per $1 (price + upfront fees)", "out", F_USD3), (27, "Investor cash received per $1", "inv", F_USD3),
             (28, "MOIC", "moic", F_X), (29, "IRR (annualised)", "irr", F_PCT), (30, "Month tier 1 hit (loan age)", "hit", F_INT)]
    for r, lab, key_, fmt in prows:
        key = key_ in ("moic", "irr")
        w.label(r, lab, key=key)
        for j, c in enumerate(cols):
            w.put(r, 2 + j, f"={le}{c}{S[key_]}", fmt, key=key)
    w.section(32, "4. Monthly view (program level): cash by party and running returns", width=12)
    monthly_header(w, 33)
    mv = [(34, "OCM net cash flow ($)", lambda m, c, p: f"={sp}{c}70", F_USD, "sum"),
          (35, "Loan America net cash flow ($)", lambda m, c, p: f"={sp}{c}71", F_USD, "sum"),
          (36, "Total investor net cash flow ($)", lambda m, c, p: f"={sp}{c}72", F_USD, "sum"),
          (37, "Investor distributions ($)", lambda m, c, p: f"={sp}{c}45", F_USD, "sum"),
          (38, "School cash received ($)", lambda m, c, p: f"={pl}{c}36", F_USD, "sum"),
          (39, "Goal / across-life fees ($)", lambda m, c, p: f"={pl}{c}22", F_USD, "sum"),
          (40, "OCM running MOIC", lambda m, c, p: f"={sp}{c}74", F_X, "last"),
          (41, "Loan America running MOIC", lambda m, c, p: f"={sp}{c}75", F_X, "last"),
          (42, "Total investor running MOIC", lambda m, c, p: f"={sp}{c}76", F_X, "last"),
          (43, "OCM running IRR", lambda m, c, p: f"={sp}{c}77", F_PCT, None),
          (44, "Loan America running IRR", lambda m, c, p: f"={sp}{c}78", F_PCT, None),
          (45, "Total investor running IRR", lambda m, c, p: f"={sp}{c}79", F_PCT, None),
          (46, "School cumulative cash (% of program face)", lambda m, c, p: f"=IF({pl}$B$8>0,{pl}{c}37/{pl}$B$8,0)", F_PCT, "last")]
    for r, lab, fn, fmt, tot in mv:
        mrow(w, r, lab, fn, fmt, total=tot, key="MOIC" in lab and "Total" in lab)
    w.section(48, "5. Recycling: how each month's loan purchases are funded", width=12)
    monthly_header(w, 49)
    rc = [(50, "Loan purchases ($)", lambda m, c, p: f"={pl}{c}11", F_USD), (51, "Funded by recycled collections ($)", lambda m, c, p: f"={pl}{c}76", F_USD),
          (52, "Funded by debt draws ($)", lambda m, c, p: f"={pl}{c}77", F_USD), (53, "Funded by new equity ($)", lambda m, c, p: f"={pl}{c}78", F_USD),
          (54, "Share funded by recycled collections (%)", lambda m, c, p: f'=IF({c}50>0,{c}51/{c}50,"")', F_PCT),
          (55, "Equity called in the month ($)", lambda m, c, p: f"={pl}{c}71", F_USD),
          (56, "Cumulative equity called ($)", lambda m, c, p: f"=SUM($C$55:{c}55)", F_USD),
          (57, "Cumulative purchases ($)", lambda m, c, p: f"=SUM($C$50:{c}50)", F_USD),
          (58, "Check: collections + debt + equity - purchases (0)", lambda m, c, p: f"=ROUND({c}51+{c}52+{c}53-{c}50,2)", F_USD),
          (59, "Helper: month if collections fund at least 50%", lambda m, c, p: f'=IF(AND(ISNUMBER({c}54),N({c}54)>=0.5),{m},999)', F_INT),
          (60, "Helper: month if no new equity was needed for purchases", lambda m, c, p: f"=IF(AND({c}50>0,{c}53<1),{m},999)", F_INT)]
    for r, lab, fn, fmt in rc:
        tot = "abs" if "Check" in lab else (None if ("Share" in lab or "Helper" in lab or "Cumulative" in lab) else "sum")
        mrow(w, r, lab, fn, fmt, total=tot)
    w.put(54, 2, "=IF(B50>0,B51/B50,0)", F_PCT, bold=True)
    w.label(62, "First month collections fund at least 50% of that month's purchases", key=True)
    w.put(62, 2, '=IF(MIN(C59:DR59)=999,"never",MIN(C59:DR59))', F_INT, key=True)
    w.label(63, "First month collections fund 100% of that month's purchases (no new equity for purchases)", key=True)
    w.put(63, 2, '=IF(MIN(C60:DR60)=999,"never",MIN(C60:DR60))', F_INT, key=True)
    w.label(64, "Total equity called / total purchases", key=True)
    w.put(64, 2, "=IF(B50>0,B55/B50,0)", F_PCT, key=True)
    w.section(66, "6. Leverage: debt facility month by month and its effect on equity and distributions", width=12)
    monthly_header(w, 67)
    lv = [(68, "Facility mode", lambda m, c, p: f"={pl}{c}50", None, None), (69, "Investor cost basis ($)", lambda m, c, p: f"={pl}{c}51", F_USD, None),
          (70, "Borrowing capacity ($)", lambda m, c, p: f"={pl}{c}53", F_USD, None), (71, "Opening debt ($)", lambda m, c, p: f"={pl}{c}56", F_USD, None),
          (72, "Draws ($)", lambda m, c, p: f"={pl}{c}58", F_USD, "sum"), (73, "Repayments ($)", lambda m, c, p: f"=-{pl}{c}59", F_USD, "sum"),
          (74, "Interest paid ($)", lambda m, c, p: f"={pl}{c}57", F_USD, "sum"), (75, "Ending debt ($)", lambda m, c, p: f"={pl}{c}60", F_USD, None),
          (76, "Equity called ($)", lambda m, c, p: f"={pl}{c}71", F_USD, "sum"), (77, "Distributions to investors ($)", lambda m, c, p: f"={sp}{c}45", F_USD, "sum"),
          (78, "Check: ending = opening + draws + repayments (0)", lambda m, c, p: f"=ROUND({c}75-({c}71+{c}72+{c}73),2)", F_USD, "abs"),
          (79, "Helper: month if there is a draw", lambda m, c, p: f"=IF({c}72>0.5,{m},999)", F_INT, None),
          (80, "Helper: month if debt is repaid after the peak", lambda m, c, p: f'=IF(AND({c}75<0.5,{m}>N($B$84)),{m},999)', F_INT, None)]
    for r, lab, fn, fmt, tot in lv:
        mrow(w, r, lab, fn, fmt, total=tot)
    kp = [(82, "First draw month", '=IF(MIN(C79:DR79)=999,"none",MIN(C79:DR79))', F_INT),
          (83, "Peak debt ($)", "=MAX(C75:DR75)", F_USD), (84, "Peak debt month", '=IF(B83>0,MATCH(B83,C75:DR75,0),"")', F_INT),
          (85, "Debt fully repaid by month", '=IF(B83>0,IF(MIN(C80:DR80)=999,"not repaid",MIN(C80:DR80)),"n.a.")', F_INT),
          (86, "Total interest paid ($)", "=B74", F_USD),
          (87, "Debt recap: distributions in the facility start month funded by the first draw ($)", f"={sp}B24", F_USD)]
    for r, lab, f, fmt in kp:
        w.label(r, lab, key=True)
        w.put(r, 2, f, fmt, key=True)
    w.label(88, "Overall program check (funding and debt roll)")
    w.put(88, 2, '=IF(AND(B58<1,B78<1),"OK","CHECK")'); B.name("ChkProgram", PD, "B88")
    w.widths(2, 7, 15)
    w.widths(8, 3 + NM, 13)


# ============================================================================ Sensitivity: advance rate
def build_sens_advance(B: Book, inp: Inputs):
    w = B.sheet(SA, 70)
    le = q(LE)
    blocks, bl = B.ref["LE"]["blocks"], B.ref["LE"]["bl"]
    w.title("Sensitivity: advance rate", "Investor and school returns per $1 of face. Each column shifts every term's advance rate by the change shown; "
                                         "upfront fees are recomputed. The middle column is the current case. Needs a single constant split.")
    w.label(4, "Advance-rate step size (Assumptions, section 8)"); w.put(4, 2, "=AdvStep", F_PCT)
    w.label(5, "Single constant investor / school split (required)"); w.put(5, 2, "=SingleSplit")
    w.label(6, "Investor % used"); w.put(6, 2, "=SplitPct", F_PCT)
    steps = list(range(-5, 6))
    w.label(8, "Step"); w.header_row(8, 3, steps)
    for j, i in enumerate(steps):
        c = L(3 + j)
        w.put(9, 3 + j, f"={i}*AdvStep", F_PCT)
        w.put(10, 3 + j, f"=AdvBlended+{c}9*SUM({q(AS)}$B$20:$E$20)", F_PCT)
        w.put(11, 3 + j, f"=1-{c}10", F_PCT)
    w.label(9, "Change in advance rate (pp)"); w.label(10, "Advance rate (blended)", key=True); w.label(11, "Discount taken at day 0 (1 - advance rate)")
    # helper: per term outlay and investor cash flows for every step
    H0 = 40
    w.section(H0 - 2, "Helper: shifted advance and outlay by term, and investor cash flows per $1 for every step", width=14)
    w.label(H0 - 1, "Loan age"); w.header_row(H0 - 1, AGE0 + 1, list(range(NA)))
    adv_r, out_r = {}, {}
    r = H0
    for k in range(4):
        tc = L(2 + k)
        adv_r[k], out_r[k] = r, r + 1
        w.label(r, f"{inp.terms[k]}m: shifted advance rate"); w.label(r + 1, f"{inp.terms[k]}m: shifted investor outlay per $1")
        for j in range(11):
            c = L(3 + j)
            w.put(r, 3 + j, f"={q(AS)}${tc}$58+{c}$9", F_PCT)
            w.put(r + 1, 3 + j, f"={c}{r}+{q(AS)}${tc}$72+{c}$9*UpfrontNetPaidRate", F_USD3)
        r += 2
    cf = {}
    r += 1
    for k in list(range(4)) + ["bl"]:
        for j in range(11):
            cf[(k, j)] = r
            nm = f"{inp.terms[k]}m" if k != "bl" else "Blended"
            w.label(r, f"{nm} investor cash flow, step {steps[j]:+d}")
            netrow = blocks[k]["net"] if k != "bl" else bl["net"]
            outl = f"{L(3 + j)}${out_r[k]}" if k != "bl" else f"(OutlayBlended+{L(3 + j)}$9*(1+UpfrontNetPaidRate)*SUM({q(AS)}$B$20:$E$20))"
            for a in range(NA):
                f = f"=SplitPct*{le}{ac(a)}{netrow}" + (f"-{outl}" if a == 0 else "")
                w.put(r, AGE0 + 1 + a, f, F_USD3)
            r += 1
    sumnet = lambda k: f"SUM({le}{ac(0)}{(blocks[k]['net'] if k != 'bl' else bl['net'])}:{ac(60)}{(blocks[k]['net'] if k != 'bl' else bl['net'])})"
    w.section(13, "A. Investor MOIC per $1 of face (net of all fees)", width=14)
    w.section(19, "B. Investor IRR per $1 of face (annualised, net of all fees)", width=14)
    w.section(25, "C. School (blended, per $1 of face)", width=14)
    names = [f"{t}m" for t in inp.terms] + ["Blended (mix-weighted)"]
    for idx, k in enumerate(list(range(4)) + ["bl"]):
        w.label(14 + idx, names[idx], key=(k == "bl")); w.label(20 + idx, names[idx], key=(k == "bl"))
        for j in range(11):
            c = L(3 + j)
            bad = f"{c}{adv_r[k]}<=0" if k != "bl" else "MIN(" + ",".join(f"{c}{adv_r[x]}" for x in range(4)) + ")<=0"
            outl = f"{c}{out_r[k]}" if k != "bl" else f"(OutlayBlended+{c}$9*(1+UpfrontNetPaidRate)*SUM({q(AS)}$B$20:$E$20))"
            row = cf[(k, j)]
            w.put(14 + idx, 3 + j, f'=IF(OR(SingleSplit<>"OK",{bad}),"n.a.",SplitPct*{sumnet(k)}/{outl})', F_X, key=(k == "bl"))
            w.put(20 + idx, 3 + j, f'=IF(OR(SingleSplit<>"OK",{bad}),"n.a.",IFERROR((1+IRR({L(AGE0 + 1)}{row}:{L(AGE0 + 61)}{row},0.01))^12-1,"n.a."))',
                  F_PCT, key=(k == "bl"))
    w.label(26, "School total received, undiscounted (% of face)", key=True)
    w.label(27, "School total received, PV at the school discount rate (% of face)", key=True)
    w.label(28, "School effective discount, PV (1 - PV received)")
    bln = bl["net"]
    for j in range(11):
        c = L(3 + j)
        w.put(26, 3 + j, f"={c}10+(1-SplitPct)*{sumnet('bl')}", F_PCT, key=True)
        w.put(27, 3 + j, f"={c}10+(1-SplitPct)*({le}{ac(0)}{bln}+NPV((1+SchoolPV)^(1/12)-1,{le}{ac(1)}{bln}:{ac(60)}{bln}))", F_PCT, key=True)
        w.put(28, 3 + j, f"=1-{c}27", F_PCT)
    S = B.ref["LE_sum"]
    w.label(30, "Check: middle column = Loan economics, blended MOIC, IRR and school (0)")
    w.put(30, 2, f'=IF(SingleSplit<>"OK",0,ROUND(H18-{le}F{S["moic"]},9)+ROUND(H24-{le}F{S["irr"]},9)+ROUND(H26-{le}F{S["s_tot"]},9))', F_NUM)
    w.label(31, "Check: middle column MOIC and IRR by term = Loan economics (0)")
    w.put(31, 2, '=IF(SingleSplit<>"OK",0,' + "+".join(f"ROUND(H{14 + k}-{le}{L(2 + k)}{S['moic']},9)+ROUND(H{20 + k}-{le}{L(2 + k)}{S['irr']},9)" for k in range(4)) + ")", F_NUM)
    w.label(32, "Overall sensitivity check", key=True)
    w.put(32, 2, '=IF(AND(ABS(B30)<0.000001,ABS(B31)<0.000001),"OK","CHECK")', key=True); B.name("ChkSensAdv", SA, "B32")
    w.widths(2, 14, 12)
    B.ref["SA"] = {"x": 10, "moic": 14, "irr": 20}


# ============================================================================ Sensitivity: investor split
def build_sens_split(B: Book, inp: Inputs):
    w = B.sheet(SS, 66)
    pl = q(PL)
    w.title("Sensitivity: investor split", "SPV, OCM and Loan America returns vs a single investor / school split. Each column re-runs the SPV "
                                           "(other fees, leverage, recycling, equity / distributions and the waterfall) in its block below.")
    w.label(4, "Split step size (Assumptions, section 8)"); w.put(4, 2, "=SplitStep", F_PCT)
    w.label(5, "Single constant investor / school split (required)"); w.put(5, 2, "=SingleSplit")
    steps = list(range(-5, 6))
    w.label(7, "Step"); w.header_row(7, 3, steps)
    w.label(8, "Change in investor % (pp)"); w.label(9, "Investor % of net collections", key=True); w.label(10, "School % of net collections")
    for j, i in enumerate(steps):
        w.put(8, 3 + j, f"={i}*SplitStep", F_PCT)
        w.put(9, 3 + j, f"=SplitPct+{L(3 + j)}8", F_PCT, key=True)
        w.put(10, 3 + j, f"=1-{L(3 + j)}9", F_PCT)
    # scenario blocks
    B0, BS = 40, 40
    blocks = {}
    for j in range(11):
        s = B0 + j * BS
        blocks[j] = s
        w.section(s, f"Scenario block {j + 1}: step {steps[j]:+d}, investor % in cell B{s + 1}", width=12)
        w.label(s + 1, "Investor % of net collections"); w.put(s + 1, 2, f"={L(3 + j)}$9", F_PCT)
        sp_ = f"$B${s + 1}"
        R = {nm: s + 2 + i for i, nm in enumerate(["inv", "fees", "open", "draws", "int", "repay", "end", "net", "op", "eq", "dist", "cumeq",
                                                   "t1", "cumt1", "hb2", "t12", "hb2a", "ocm12", "hb3", "t3", "hb4", "t4", "t5", "ocm", "la",
                                                   "hb3a", "hb4a", "oeq", "leq", "tnet", "onet", "lnet"])}
        labs = {"inv": "Investor collections ($)", "fees": "Other SPV fees ($)", "open": "Opening debt ($)", "draws": "Debt draws ($)",
                "int": "Interest ($)", "repay": "Repayments ($)", "end": "Ending debt ($)", "net": "Net SPV cash before equity ($)",
                "op": "Operating cash ($)", "eq": "Equity called ($)", "dist": "Distributions ($)", "cumeq": "Cumulative equity ($)",
                "t1": "Tier 1 ($)", "cumt1": "Cumulative tier 1 ($)", "hb2": "Pref hurdle balance before ($)", "t12": "Tiers 1 + 2 ($)",
                "hb2a": "Pref hurdle balance after ($)", "ocm12": "OCM share of tiers 1 + 2 ($)", "hb3": "OCM tier 3 hurdle balance before ($)",
                "t3": "Tier 3 ($)", "hb4": "OCM tier 4 hurdle balance before ($)", "t4": "Tier 4 ($)", "t5": "Tier 5 ($)", "ocm": "OCM distributions ($)",
                "la": "Loan America distributions ($)", "hb3a": "OCM tier 3 hurdle balance after ($)", "hb4a": "OCM tier 4 hurdle balance after ($)",
                "oeq": "OCM equity ($)", "leq": "Loan America equity ($)", "tnet": "Total net cash flow ($)", "onet": "OCM net cash flow ($)",
                "lnet": "Loan America net cash flow ($)"}
        sc = s + 2 + len(R)
        w.label(sc, "Loan America share of capital (this scenario)"); w.label(sc + 1, "Effective Loan America shares in tiers 3, 4, 5")
        LAs = f"$B${sc}"
        w.put(sc, 2, f'=IF(CapitalBasis="OCM %",1-OCMPct,IFERROR(MIN(1,LAFixed/B{R["eq"]}),0))', F_PCT)
        for jj, nm in enumerate(["LAShare3", "LAShare4", "LAShare5"]):
            w.put(sc + 1, 2 + jj, f'=IF(PromoteMethod="On top of pro rata",{nm}+(1-{nm})*{LAs},{nm})', F_PCT)
        S3, S4, S5 = f"$B${sc + 1}", f"$C${sc + 1}", f"$D${sc + 1}"
        OC = f"(1-{LAs})"
        for nm, r in R.items():
            w.label(r, labs[nm])
        for m in range(1, NM + 1):
            c, p = mc(m), (mc(m - 1) if m > 1 else None)
            g = lambda nm: f"{c}{R[nm]}"
            pv = lambda nm: f"{p}{R[nm]}"
            mode, cap, outlay, rate, act = f"{pl}{c}50", f"{pl}{c}53", f"{pl}{c}11", f"{pl}{c}55", f"{pl}{c}46"
            F = {
                "inv": f"={sp_}*{pl}{c}23",
                "fees": f"=IF({g('inv')}<=0,0,-({act}*MIN(IF(ROUNDUP({m}/12,0)<=FeeYearsHigh,FeeHigh,FeeLow)/12,MAX(0,{g('inv')}-{g('int')}))))",
                "open": "=0" if m == 1 else f"={pv('end')}",
                "draws": f'=IF({mode}="Origination",IF(Recycling="ON",MAX(0,{cap}-{g("open")}),MIN(MAX(0,{cap}-{g("open")}),{outlay})),0)',
                "int": f'=IF({mode}="Origination",({g("open")}+{g("open")}+{g("draws")}-MAX(0,{g("open")}-{cap}))/2,{g("open")})*{rate}/12',
                "repay": f'=IF({mode}="Origination",MAX(0,{g("open")}-{cap}),IF({mode}="Amortization",MIN({g("open")},MAX(0,{g("open")}-{cap},Sweep*({g("inv")}+{g("fees")}-{g("int")}))),0))',
                "end": f"={g('open')}+{g('draws')}-{g('repay')}",
                "net": f"={g('inv')}+{g('draws')}-{outlay}+{g('fees')}-{g('int')}-{g('repay')}",
                "op": f"={g('inv')}+{g('fees')}-{g('int')}-{g('repay')}",
                "eq": f'=IF(Recycling="ON",MAX(0,-{g("net")}),MAX(0,{outlay}-{g("draws")})+MAX(0,-{g("op")}))',
                "dist": f'=IF(Recycling="ON",MAX(0,{g("net")}),MAX(0,{g("op")})+MAX(0,{g("draws")}-{outlay}))',
                "cumeq": f"=SUM(${mc(1)}{R['eq']}:{c}{R['eq']})",
                "t1": f"=MIN({g('dist')},MAX(0,{g('cumeq')}-{'0' if m == 1 else pv('cumt1')}))",
                "cumt1": f"=SUM(${mc(1)}{R['t1']}:{c}{R['t1']})",
                "hb2": f"={g('eq')}" if m == 1 else f"={pv('hb2a')}*(1+H2m)+{g('eq')}",
                "t12": f"=MAX({g('t1')},MIN({g('dist')},MAX(0,{g('hb2')})))",
                "hb2a": f"={g('hb2')}-{g('dist')}",
                "ocm12": f"={g('t12')}*{OC}",
                "hb3": f"={g('eq')}*{OC}" if m == 1 else f"={pv('hb3a')}*(1+H3m)+{g('eq')}*{OC}",
                "t3": f'=IF(Promote="ON",IF({S3}<1,MIN({g("dist")}-{g("t12")},MAX(0,{g("hb3")}-{g("ocm12")})/(1-{S3})),{g("dist")}-{g("t12")}),0)',
                "hb4": f"={g('eq')}*{OC}" if m == 1 else f"={pv('hb4a')}*(1+H4m)+{g('eq')}*{OC}",
                "t4": f'=IF(Promote="ON",IF({S4}<1,MIN({g("dist")}-{g("t12")}-{g("t3")},MAX(0,{g("hb4")}-{g("ocm12")}-{g("t3")}*(1-{S3}))/(1-{S4})),{g("dist")}-{g("t12")}-{g("t3")}),0)',
                "t5": f"={g('dist')}-{g('t12')}-{g('t3')}-{g('t4')}",
                "ocm": f'=IF(Promote="ON",{g("ocm12")}+{g("t3")}*(1-{S3})+{g("t4")}*(1-{S4})+{g("t5")}*(1-{S5}),{g("dist")}*{OC})',
                "la": f"={g('dist')}-{g('ocm')}",
                "hb3a": f"={g('hb3')}-{g('ocm')}", "hb4a": f"={g('hb4')}-{g('ocm')}",
                "oeq": f"={g('eq')}*{OC}", "leq": f"={g('eq')}*{LAs}",
                "tnet": f"={g('dist')}-{g('eq')}", "onet": f"={g('ocm')}-{g('oeq')}", "lnet": f"={g('la')}-{g('leq')}"}
            for nm, f in F.items():
                w.put(R[nm], MON1 + m - 1, f, F_USD)
        for nm in ("eq", "dist", "oeq", "leq", "ocm", "la", "tnet", "onet", "lnet"):
            w.put(R[nm], 2, f"=SUM({mc(1)}{R[nm]}:{mc(NM)}{R[nm]})", F_USD, bold=True)
        blocks[j] = R
    # summary
    sec = [("A. SPV (all investors)", "eq", "dist", "tnet", 12), ("B. OCM", "oeq", "ocm", "onet", 18), ("C. Loan America (sponsor)", "leq", "la", "lnet", 24)]
    for head, e_, d_, n_, r0 in sec:
        w.section(r0, head, width=14)
        short = head.split(". ")[1].split(" (")[0]
        for i, (lab, fmt) in enumerate([("Equity called ($)", F_USD), ("Distributions ($)", F_USD), ("MOIC", F_X), ("IRR (annualised)", F_PCT)]):
            r = r0 + 1 + i
            key = lab in ("MOIC", "IRR (annualised)")
            w.label(r, f"{short}: {lab}", key=key)
            for j in range(11):
                R = blocks[j]
                f = {0: f"=$B${R[e_]}", 1: f"=$B${R[d_]}", 2: f"=IF({L(3 + j)}{r0 + 1}>0,{L(3 + j)}{r0 + 2}/{L(3 + j)}{r0 + 1},0)",
                     3: f'=IFERROR((1+IRR({mc(1)}{R[n_]}:{mc(NM)}{R[n_]},0.01))^12-1,"n.a.")'}[i]
                w.put(r, 3 + j, f'=IF(SingleSplit<>"OK","n.a.",{f[1:]})', fmt, key=key)
    w.label(30, "Loan America share of total distributions", key=True)
    for j in range(11):
        c = L(3 + j)
        w.put(30, 3 + j, f'=IF(SingleSplit<>"OK","n.a.",IF(N({c}14)>0,{c}26/{c}14,0))', F_PCT, key=True)
    sp = q(SP)
    w.label(32, "Check: middle column = SPV and waterfall (equity, distributions, MOIC, IRR for each party) (0)")
    parts = []
    for r0, col in [(12, "B"), (18, "C"), (24, "D")]:
        for i, rr in enumerate([17, 18, 20, 21]):
            parts.append(f"ROUND(N(H{r0 + 1 + i})-N({sp}{col}{rr}),3)")
    w.put(32, 2, '=IF(SingleSplit<>"OK",0,' + "+".join(parts) + ")", F_NUM)
    w.label(33, "Check: OCM + Loan America = SPV in every scenario, equity and distributions (0)")
    w.put(33, 2, '=IF(SingleSplit<>"OK",0,ROUND(SUMPRODUCT(ABS(N(+C19:M19)+N(+C25:M25)-N(+C13:M13)))+SUMPRODUCT(ABS(N(+C20:M20)+N(+C26:M26)-N(+C14:M14))),2))', F_NUM)
    w.put(33, 2, '=IF(SingleSplit<>"OK",0,ROUND(' + "+".join(f"ABS({L(3 + j)}19+{L(3 + j)}25-{L(3 + j)}13)+ABS({L(3 + j)}20+{L(3 + j)}26-{L(3 + j)}14)" for j in range(11)) + ",2))", F_NUM)
    w.label(34, "Overall split-sensitivity check", key=True)
    w.put(34, 2, '=IF(AND(ABS(B32)<0.01,ABS(B33)<1),"OK","CHECK")', key=True); B.name("ChkSensSplit", SS, "B34")
    w.widths(2, 14, 14)
    w.widths(15, 3 + NM, 12)
    B.ref["SS"] = {"x": 9, "spv": 13, "ocm": 19, "la": 25}


# ============================================================================ Checks
def build_checks(B: Book, inp: Inputs):
    w = B.sheet(CK, 70)
    S = B.ref["LE_sum"]
    w.title("Checks", "Every identity and tie-out in the workbook. Model status is OK only if every calculation check below is OK.")
    w.header_row(4, 1, ["Check", "Status", "Area"])
    items = [("Term mix sums to 100%", "=ChkMix", "Assumptions"), ("Effective loss curves never decrease", "=ChkCurves", "Assumptions"),
             ("Every active fee row valid; loan sizes positive", "=ChkFees", "Assumptions"), ("At least one valid split tier", "=ChkTiers", "Assumptions"),
             ("Purchase split sums to 100%; period 5 to 60 months", "=ChkSplit", "Assumptions"), ("Purchase cash adds up to the program", "=ChkCash", "Assumptions"),
             ("Ramp-up valid", "=ChkRamp", "Assumptions"), ("Last draw month on or after the start month", "=ChkLev", "Assumptions"),
             ("Last draw month covers the purchase period (warning only)", "=WarnDraw", "Assumptions"),
             ("All input checks", "=InputStatus", "Assumptions")]
    calc = []
    for k in range(4):
        calc.append((f"{inp.terms[k]}m: investor + school = net collections; all principal accounted for", f"={q(LE)}{L(2 + k)}{S['chk']}", "Loan economics"))
    calc += [("Blended: investor + school = net collections", f"={q(LE)}F{S['chk']}", "Loan economics"),
             ("Gross investor + school = gross collections", "=ChkGross", "Gross vs net"),
             ("Pool identities (collections, split, balance roll, SPV cash)", f'=IF({q(PL)}B29+{q(PL)}B30+{q(PL)}B31+{q(PL)}B74=0,"OK","CHECK")', "Pool cash flows"),
             ("Debt fully repaid by month 120", f'=IF(ABS({q(PL)}DR60)<1,"OK","CHECK")', "Pool cash flows"),
             ("Pool MOIC and IRR tie to Loan economics", f'=IF({q(PL)}B93=0,"OK","CHECK")', "Pool cash flows"),
             ("Vintage grids add up to the pool", f'=IF({q(PL)}B171+{q(PL)}B236<1,"OK","CHECK")', "Pool cash flows"),
             ("Pool overall", "=ChkPool", "Pool cash flows"),
             ("Waterfall tiers add to distributions; OCM + Loan America = total", "=ChkWaterfall", "SPV and waterfall"),
             ("Effective promote shares between 0% and 100%", "=ChkShares", "SPV and waterfall"),
             ("Yearly view: OCM + Loan America = total", f'=IF(SUMPRODUCT(ABS({q(SP)}C118:L118))=0,"OK","CHECK")', "SPV and waterfall"),
             ("Funding and debt roll", "=ChkProgram", "Program detail"),
             ("Advance-rate sensitivity: middle column ties to the live model", "=ChkSensAdv", "Sens - Advance rate"),
             ("Split sensitivity: middle column ties to the live model", "=ChkSensSplit", "Sens - Investor split")]
    r = 5
    for lab, f, area in items + calc:
        w.label(r, lab)
        w.put(r, 2, f)
        w.put(r, 3, area)
        r += 1
    first_calc = 5 + len(items)
    w.label(r + 1, "Model status (all calculation checks)", key=True)
    w.put(r + 1, 2, f'=IF(COUNTIF(B{first_calc}:B{r - 1},"CHECK")=0,"OK","CHECK")', key=True)
    B.name("ModelStatus", CK, f"B{r + 1}")
    w.label(r + 2, "Input status (all input checks)", key=True)
    w.put(r + 2, 2, "=InputStatus", key=True)
    w.ws.column_dimensions["B"].width = 12
    w.ws.column_dimensions["C"].width = 24
    from openpyxl.formatting.rule import CellIsRule
    rng = f"B5:B{r + 2}"
    w.ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'], font=Font(color="1E7B45", bold=True)))
    w.ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"CHECK"'], font=Font(color="B3261E", bold=True)))
    w.ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"WARNING"'], font=Font(color="B7791F", bold=True)))


# ============================================================================ Overview
def build_overview(B: Book, inp: Inputs):
    w = B.sheet(OV, 46)
    ws = w.ws
    sp, pl, le, gn = q(SP), q(PL), q(LE), q(GN)
    S = B.ref["LE_sum"]
    w.title("Fusion Loans: deal overview", "Every figure is a formula linked to the calculation tabs. MOIC is shown to 2 decimals here, 3 elsewhere.")
    w.label(3, "Status"); w.put(3, 2, '="Inputs "&InputStatus&"  |  Model "&ModelStatus', bold=True)
    w.put(3, 4, f'=TEXT(SplitPct,"0.00%")&" to investors, "&TEXT(AdvBlended,"0.00%")&" advance, "&TEXT(PurchaseAmt,"$#,##0")&" over "&Period&" months, recycling "&Recycling&", leverage "&Leverage')
    w.section(5, "1. Investor returns (SPV level: after all fees, recycling, leverage and the promote)", width=6)
    w.header_row(6, 2, ["MOIC", "IRR", "Equity in ($)", "Distributions out ($)", "Share of capital"])
    for i, (lab, col) in enumerate([("SPV (all investors)", "B"), ("OCM", "C"), ("Loan America", "D")]):
        r = 7 + i
        w.label(r, lab, key=True)
        for j, (row, fmt) in enumerate([(20, F_X2), (21, F_PCT), (17, F_USD), (18, F_USD), (22, F_PCT)]):
            w.put(r, 2 + j, f"={sp}{col}{row}", fmt, key=(j < 2))
    w.section(11, "2. Capital and cash (program level)", width=3)
    cap = [("Equity called ($)", f"={pl}B99", F_USD), ("Distributions ($)", f"={pl}B100", F_USD), ("Net profit to investors ($)", f"={pl}B100-{pl}B99", F_USD),
           ("Loan purchases: cash to schools + upfront fees ($)", f"={pl}B11", F_USD), ("Face purchased ($)", f"={pl}B8", F_USD),
           ("Funded by recycled collections ($)", f"={pl}B101", F_USD), ("Share of purchases funded by recycled collections", f"=IF({pl}B11>0,{pl}B101/{pl}B11,0)", F_PCT),
           ("Peak debt ($)", f"={pl}B96", F_USD), ("Interest on debt ($)", f"={pl}B98", F_USD)]
    for i, (lab, f, fmt) in enumerate(cap):
        key = lab.startswith(("Net profit", "Equity called", "Distributions"))
        w.label(12 + i, lab, key=key); w.put(12 + i, 2, f, fmt, key=key)
    w.section(22, "3. Asset level, per $1 of face (blended, before SPV costs)", width=3)
    asset = [("Gross MOIC (before fees)", f"={gn}F10", F_X2), ("Gross IRR (before fees)", f"={gn}F11", F_PCT),
             ("Net MOIC (after fees)", f"={le}F{S['moic']}", F_X2), ("Net IRR (after fees)", f"={le}F{S['irr']}", F_PCT),
             ("Investor outlay per $1 of face", "=OutlayBlended", F_USD3), ("Fee drag on MOIC (gross - net)", f"={gn}F14", F_X2),
             ("Charge-offs (% of face)", f"={le}F{S['co']}", F_PCT), ("Advance rate (blended)", "=AdvBlended", F_PCT),
             ("Investor / school split", '=TEXT(SplitPct,"0.00%")&" / "&TEXT(1-SplitPct,"0.00%")', None)]
    for i, (lab, f, fmt) in enumerate(asset):
        key = "MOIC" in lab or "IRR" in lab
        w.label(23 + i, lab, key=key); w.put(23 + i, 2, f, fmt, key=key)
    w.section(33, "4. School, % of loan face", width=3)
    school = [("Day-0 advance", f"={le}F{S['s_adv']}", F_PCT), ("Share of collections, undiscounted", f"={le}F{S['s_coll']}", F_PCT),
              ("Total received, undiscounted", f"={le}F{S['s_tot']}", F_PCT), ("Total received, PV", f"={le}F{S['s_pv']}", F_PCT),
              ("Effective discount, PV", f"={le}F{S['s_discpv']}", F_PCT), ("School cash, program ($)", f"={pl}B36", F_USD)]
    for i, (lab, f, fmt) in enumerate(school):
        key = "Total received" in lab
        w.label(34 + i, lab, key=key); w.put(34 + i, 2, f, fmt, key=key)
    w.section(41, "5. Deal settings and model status", width=3)
    sett = [("Program / purchase period", '=TEXT(PurchaseAmt,"$#,##0")&" / "&Period&" months"'),
            ("Term mix", f'=TEXT({q(AS)}B20,"0.00%")&" / "&TEXT({q(AS)}C20,"0.00%")&" / "&TEXT({q(AS)}D20,"0.00%")&" / "&TEXT({q(AS)}E20,"0.00%")'),
            ("Loss multiplier", '=TEXT(LossMult,"0.00")&"x"'), ("Recycling", "=Recycling"),
            ("Leverage", '=IF(Leverage="ON","ON, "&TEXT(LTV,"0.00%")&" LTV, months "&FacStart&"-"&LastDraw,"OFF")'),
            ("Promote", '=Promote&" ("&PromoteMethod&")"'),
            ("Loan America capital", f'=TEXT({sp}B7,"0.00%")&" ("&CapitalBasis&")"'),
            ("Input checks", "=InputStatus"), ("Model checks", "=ModelStatus")]
    for i, (lab, f) in enumerate(sett):
        w.label(42 + i, lab); w.put(42 + i, 2, f)
    ws.column_dimensions["B"].width = 18
    for c in "CDEF":
        ws.column_dimensions[c].width = 21
    B.ref["OV"] = True


# ============================================================================ Charts (openpyxl native charts on live ranges)
SERIES_COLOURS = [("equity called", "8E2C3A"), ("repayments", "8E2C3A"), ("debt draws", "8E2C3A"), ("capacity", "7A8696"),
                  ("recycled", "1A8C7E"), ("new equity", "24527A"), ("purchases", "C8962E"), ("promote", "1A8C7E"),
                  ("pro rata", "62B3A4"), ("loan america", "1A8C7E"), ("ocm", "24527A"), ("school", "C8962E"), ("total", "17202A"),
                  ("spv", "24527A"), ("ending debt", "17202A"), ("draws", "24527A"), ("distributions", "24527A"), ("blended", "17202A"),
                  ("24m", "4F7FA8"), ("36m", "62B3A4"), ("48m", "C8962E"), ("60m", "8E2C3A"), ("share", "17202A")]


def _colour(title):
    t = (title or "").lower()
    for key, col in SERIES_COLOURS:
        if key in t:
            return col
    return None


def _series(chart, ws_title, row, c1, c2, title_row_label=None):
    from openpyxl.chart import Reference, Series
    ws = B_CURRENT.wb[ws_title]
    data = Reference(ws, min_col=c1, max_col=c2, min_row=row, max_row=row)
    title = title_row_label or ws.cell(row, 1).value
    s = Series(data, title=title)
    col = _colour(title)
    if col:
        s.graphicalProperties.solidFill = col
        s.graphicalProperties.line.solidFill = col
    chart.series.append(s)
    return s


def _cats(chart, ws_title, row, c1, c2):
    from openpyxl.chart import Reference
    chart.set_categories(Reference(B_CURRENT.wb[ws_title], min_col=c1, max_col=c2, min_row=row, max_row=row))


B_CURRENT = None


def make_chart(kind, title, sheet, rows, c1, c2, cat_row, stacked=False, y2=None, y2_kind="line", num_fmt=None, y2_fmt=None,
               height=9, width=26, grouping=None):
    from openpyxl.chart import BarChart, LineChart
    ch = BarChart() if kind == "bar" else LineChart()
    if kind == "bar":
        ch.type = "col"
        ch.grouping = "stacked" if stacked else (grouping or "clustered")
        if stacked:
            ch.overlap = 100
        ch.gapWidth = 40
    ch.title = title
    ch.height, ch.width = height, width
    for r in rows:
        s = _series(ch, sheet, r, c1, c2)
        if kind == "line":
            s.smooth = False
            s.graphicalProperties.line.width = 22000
    _cats(ch, sheet, cat_row, c1, c2)
    if num_fmt:
        ch.y_axis.number_format = num_fmt
    ch.y_axis.majorGridlines = None if kind == "bar" and False else ch.y_axis.majorGridlines
    ch.legend.position = "b"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    if y2:
        c2_ = LineChart() if y2_kind == "line" else BarChart()
        for r in y2:
            s = _series(c2_, sheet, r, c1, c2)
            s.smooth = False
        c2_.y_axis.axId = 200
        c2_.y_axis.crosses = "max"
        c2_.y_axis.delete = False
        if y2_fmt:
            c2_.y_axis.number_format = y2_fmt
        ch += c2_
    return ch


def build_charts(B: Book, inp: Inputs):
    global B_CURRENT
    B_CURRENT = B
    w = B.sheet(CH, 4)
    ws = w.ws
    w.title("Charts", "All charts are linked to the calculation tabs and update with the inputs.")
    M120, M60, M96 = (MON1, MON1 + 119), (MON1, MON1 + 59), (MON1, MON1 + 95)
    usd, x3, pct = '"$"#,##0', '0.000"x"', "0%"
    charts = [
        ("1. Returns overview", [
            make_chart("line", "Running MOIC by party (cumulative distributions / cumulative equity)", PD, [40, 41, 42], *M120, 33, num_fmt=x3),
            make_chart("bar", "Monthly cash received: investors (distributions) vs school (advance + share of collections)", PD, [37, 38], *M120, 33, stacked=True, num_fmt=usd),
            make_chart("line", "School cumulative cash received (% of program face)", PD, [46], *M120, 33, num_fmt=pct)]),
        ("2. Funding and recycling", [
            make_chart("bar", "How each month's purchases are funded; line = % funded by recycled collections", PD, [51, 52, 53], *M60, 49, stacked=True,
                       y2=[54], num_fmt=usd, y2_fmt=pct),
            make_chart("line", "Cumulative purchases vs cumulative equity called", PD, [56, 57], *M60, 49, num_fmt=usd)]),
        ("3. Leverage", [
            make_chart("bar", "Debt facility: draws and repayments (columns); debt vs borrowing capacity (lines)", PD, [72, 73], *M60, 67,
                       y2=[75, 70], num_fmt=usd, y2_fmt=usd),
            make_chart("bar", "Investors' cash: equity called vs distributions", PD, [76, 77], *M60, 67, num_fmt=usd)]),
        ("4. Equity called vs distributions by party", [
            make_chart("bar", "SPV (all investors): equity called vs distributions; line = cumulative net cash", SP, [91, 92], *M96, 37, stacked=True,
                       y2=[93], num_fmt=usd, y2_fmt=usd),
            make_chart("bar", "OCM: equity called vs distributions; line = cumulative net cash", SP, [94, 95], *M96, 37, stacked=True,
                       y2=[96], num_fmt=usd, y2_fmt=usd),
            make_chart("bar", "Loan America: equity called vs distributions (pro rata + promote); line = cumulative net cash", SP, [97, 98, 99], *M96, 37,
                       stacked=True, y2=[100], num_fmt=usd, y2_fmt=usd)]),
        ("5. Promote split over time", [
            make_chart("bar", "Monthly distributions: OCM vs Loan America; line = Loan America's share of the month", SP, [83, 84, 85], *M96, 37,
                       stacked=True, y2=[86], num_fmt=usd, y2_fmt=pct),
            make_chart("line", "Cumulative distributions: OCM (left axis) vs Loan America (right axis)", SP, [87], *M96, 37, y2=[88],
                       num_fmt=usd, y2_fmt=usd)]),
        ("6. Advance-rate sensitivity (per $1 of face, net of fees)", [
            make_chart("line", "Investor MOIC vs advance rate (by term and blended)", SA, [14, 15, 16, 17, 18], 3, 13, 10, num_fmt=x3),
            make_chart("line", "Investor IRR vs advance rate (by term and blended)", SA, [20, 21, 22, 23, 24], 3, 13, 10, num_fmt=pct)]),
        ("7. Split sensitivity", [
            make_chart("line", "MOIC vs investor split: SPV and OCM (left), Loan America (right)", SS, [15, 21], 3, 13, 9, y2=[27], num_fmt=x3, y2_fmt=x3),
            make_chart("line", "IRR vs investor split: SPV and OCM (left), Loan America (right)", SS, [16, 22], 3, 13, 9, y2=[28], num_fmt=pct, y2_fmt=pct)]),
    ]
    r = 4
    for title, chs in charts:
        w.section(r, title, width=18)
        r += 1
        for ch in chs:
            ws.add_chart(ch, f"B{r}")
            r += 20
        r += 1
    # overview charts
    ov = B.wb[OV]
    c1 = make_chart("line", "Running MOIC by party (Loan America on the right axis)", PD, [40, 42], MON1, MON1 + 95, 33, y2=[41], num_fmt=x3,
                    y2_fmt=x3, height=8, width=17)
    c2 = make_chart("bar", "SPV: equity called (below 0) vs distributions (above 0)", SP, [91, 92], MON1, MON1 + 95, 37, stacked=True,
                    num_fmt=usd, height=8, width=17)
    ov.add_chart(c1, "D12")
    ov.add_chart(c2, "D29")


# ============================================================================ Read me, compare, notes
def build_readme(B: Book, inp: Inputs):
    w = B.sheet(RM, 120)
    ws = w.ws
    w.title("Fusion Loans deal model: investor workbook")
    lines = [
        ("What this is", ""),
        ("", "A formula-driven copy of the Fusion Loans deal model. Every number outside the Assumptions tab is a live Excel formula, so each figure can be traced back to its inputs."),
        ("", "It was generated from the dashboard, and its formulas were recalculated and tied cell by cell to the dashboard's Python engine before release."),
        ("Colour coding", ""),
        ("", "Blue text on a yellow background = an input you can change (Assumptions tab only). Black = a formula. Bold rows on a blue band = key results (totals, MOIC, IRR)."),
        ("How the model flows", ""),
        ("", "1. Assumptions: every input, with named cells (for example Recycling, LossMultEff, NTiers, AdvBlended) so formulas read in plain words."),
        ("", "2. Loan economics: one $1 of loan face per term (24, 36, 48, 60 months), month by month: amortisation, charge-offs from the loss curve, prepayments, payment stop, recoveries, fees, then the investor / school split by tier."),
        ("", "3. Gross vs net: the same split run on collections before any fees, to show the fee drag."),
        ("", "4. Pool cash flows: each month's purchases form a vintage; the pool is the sum of vintages (SUMPRODUCT of face purchased with the per-$1 pattern), then other SPV fees, the leverage facility and recycling."),
        ("", "5. SPV and waterfall: equity calls and distributions split between OCM and Loan America, through return of capital, the pref and the promote tiers."),
        ("", "6. Program detail: returns by party, the school's effective advance rate, monthly views, funding and leverage views."),
        ("", "7. Sensitivities: advance rate (per $1) and investor split (full SPV re-run in visible blocks)."),
        ("", "8. Checks: every identity in the workbook, rolled up into Model status on the Overview."),
        ("How to verify", ""),
        ("", "Start on Checks: Model status should read OK. Then trace any figure with Excel's Formulas > Trace Precedents, or press Ctrl + [ to jump to the cells it uses."),
        ("", "Spot-check one loan: on Loan economics, the scheduled balance row is standard level-payment amortisation; charge-offs equal the increase in the cumulative loss curve; collections = interest + scheduled principal + prepayments + recoveries - missed payments."),
        ("", "Change an input on Assumptions (for example the loss multiplier) and watch every tab update."),
        ("Conventions", ""),
        ("", "Month 0 is the purchase month. IRRs are monthly IRRs annualised as (1 + monthly)^12 - 1. MOIC = cash received / cash put in. Percentages to 2 decimals, MOIC to 3 (2 on the Overview)."),
        ("", "Sheet order: Overview, Assumptions, Loan economics, Gross vs net, Pool cash flows, SPV and waterfall, Program detail, Charts, the two sensitivity tabs, Checks, then reference tabs."),
    ]
    r = 3
    for a, b in lines:
        if a:
            c = ws.cell(r, 1, a); c.font = Font(bold=True, size=12, color=NAVY)
        else:
            c = ws.cell(r, 1, b); c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 30
        r += 1
    w.label(r + 1, "Scenario exported", key=True)
    ws.cell(r + 1, 2, dt.datetime.now().strftime("%d %b %Y %H:%M"))


def build_compare(B: Book, saved: dict | None):
    from .engine import run_model
    w = B.sheet("Compare scenarios", 40)
    w.title("Compare scenarios", "The first column is live (linked to this workbook). Other columns are saved scenarios from the dashboard, as values.")
    labels = ["Advance (blended)", "Investor split", "Loss multiplier", "Program ($)", "Net MOIC per $1", "Net IRR per $1",
              "School total received, PV", "Equity called ($)", "Distributions ($)", "SPV MOIC", "SPV IRR", "OCM MOIC", "OCM IRR",
              "Loan America MOIC", "Loan America IRR", "Peak debt ($)", "Checks"]
    fmts = [F_PCT, F_PCT, F_X2, F_USD, F_X, F_PCT, F_PCT, F_USD, F_USD, F_X, F_PCT, F_X, F_PCT, F_X, F_PCT, F_USD, None]
    S = B.ref["LE_sum"]
    sp, pl, le = q(SP), q(PL), q(LE)
    live = ["=AdvBlended", "=SplitPct", "=LossMult", "=PurchaseAmt", f"={le}F{S['moic']}", f"={le}F{S['irr']}", f"={le}F{S['s_pv']}",
            f"={pl}B99", f"={pl}B100", f"={sp}B20", f"={sp}B21", f"={sp}C20", f"={sp}C21", f"={sp}D20", f"={sp}D21", f"={pl}B96",
            '=InputStatus&" / "&ModelStatus']
    w.header_row(4, 2, ["This workbook (live)"] + list((saved or {}).keys()))
    for i, lab in enumerate(labels):
        key = "MOIC" in lab or "IRR" in lab
        w.label(5 + i, lab, key=key)
        w.put(5 + i, 2, live[i], fmts[i], key=key)
    for j, (nm, inp_) in enumerate((saved or {}).items()):
        r_ = run_model(inp_, sensitivities=False)
        R, PS, Su = r_["spv"]["returns"], r_["pool"]["summary"], r_["ue"]["summary"]
        vals = [r_["derived"]["F71"], r_["derived"]["tiers"][0][0] if r_["derived"]["tiers"] else None, inp_.loss_mult, inp_.purchase_amount,
                Su.loc["Investor MOIC", "Blended"], Su.loc["Investor IRR", "Blended"], Su.loc["School total (PV)", "Blended"], PS["Equity called"],
                PS["Distributions"], R.loc["MOIC", "Total investors"], R.loc["IRR", "Total investors"], R.loc["MOIC", "OCM"], R.loc["IRR", "OCM"],
                R.loc["MOIC", "Loan America"], R.loc["IRR", "Loan America"], PS["Peak debt"], f"{r_['input_status']} / {r_['model_status']}"]
        for i, v_ in enumerate(vals):
            try:
                v_ = float(v_) if not isinstance(v_, str) else v_
            except (TypeError, ValueError):
                pass
            w.put(5 + i, 3 + j, v_, fmts[i])
    w.widths(2, 2 + len(saved or {}), 20)


def build_notes(B: Book, sheet3: list | None, notes: list | None):
    w = B.sheet("Scenario notes", 14)
    w.title("Scenario notes", "Your own scenario table (typed values, not linked to the model).")
    for i, row in enumerate(sheet3 or []):
        for j, val in enumerate(row):
            if val not in (None, "", "None"):
                cell = w.ws.cell(4 + i, 2 + j, val)
                if isinstance(val, str) and val.startswith(("=", "+", "-", "@")):
                    cell.data_type = "s"
    w.widths(2, 14, 16)
    w2 = B.sheet("Model notes", 34)
    w2.title("Model notes", "Documentation carried over from the original Excel model (references like F1 or F3 point to tabs in that file).")
    r = 4

    def text(cell, val):
        cell.value = val
        if isinstance(val, str) and val.startswith(("=", "+", "-", "@")):
            cell.data_type = "s"          # notes are text, never formulas
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for n in notes or []:
        a, b, c = n.get("a", ""), n.get("b", ""), n.get("c", "")
        text(w2.ws.cell(r, 1), a)
        text(w2.ws.cell(r, 2), b)
        if c:
            text(w2.ws.cell(r, 3), c)
        if a and not b and a.split(".")[0].strip().isdigit():
            w2.ws.cell(r, 1).font = Font(bold=True, color=NAVY)
        r += 1
    w2.ws.column_dimensions["B"].width = 110
    w2.ws.column_dimensions["C"].width = 40
