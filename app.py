"""Fusion Loans deal dashboard: every tab of LoanAmerica_Model v4, live in the browser.
Run with:  streamlit run app.py"""
from __future__ import annotations
import json, tempfile
from pathlib import Path
import numpy as np
import pandas as pd
import plotly.graph_objects as go
import streamlit as st

from fusion import Inputs, load_inputs_from_excel, run_model
from fusion.inputs import Fee, Tier
import ui_kit
import charts as C

ROOT = Path(__file__).parent


# Display tables mix numbers and text in one column (e.g. a tier's "age hit" is 13 or "Not hit").
# Arrow needs one type per column, so text-and-number columns are shown as text.
def _arrow_safe(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for c in out.columns:
        col = out[c]
        if col.dtype == object:
            types = {type(v) for v in col if v is not None and not (isinstance(v, float) and np.isnan(v))}
            if len(types) > 1 or (types and not types <= {str}):
                out[c] = ["" if v is None or (isinstance(v, float) and np.isnan(v)) else str(v) for v in col]
    return out


_st_dataframe = getattr(st, "_original_dataframe", st.dataframe)   # keep the original across reruns
st._original_dataframe = _st_dataframe


def _dataframe(data=None, *args, **kwargs):
    if isinstance(data, pd.DataFrame):
        data = _arrow_safe(data)
    return _st_dataframe(data, *args, **kwargs)


st.dataframe = _dataframe


# Streamlit reads text between two $ signs as a LaTeX formula, which garbles labels like
# "Loan America fixed commitment ($) (used when the basis is Fixed $)". Every $ we show is money, so escape it.
def _no_math(s, html=False):
    if not isinstance(s, str) or "$" not in s:
        return s
    s = s.replace("\\$", "$")
    return s.replace("$", "&#36;" if html else "\\$")


def _escape_label(fn_name):
    orig = getattr(st, f"_original_{fn_name}", getattr(st, fn_name))
    setattr(st, f"_original_{fn_name}", orig)

    def wrapped(label=None, *args, **kwargs):
        if fn_name == "markdown":
            return orig(_no_math(label, kwargs.get("unsafe_allow_html", False)), *args, **kwargs)
        return orig(_no_math(label), *args, **kwargs)
    setattr(st, fn_name, wrapped)


for _fn in ["markdown", "number_input", "toggle", "selectbox", "radio", "date_input", "text_input",
            "warning", "info", "multiselect", "file_uploader", "button", "download_button"]:
    _escape_label(_fn)
BASE = ROOT / "scenarios" / "base_case_v4.json"
LAB = json.loads((ROOT / "fusion" / "excel_labels.json").read_text())
st.set_page_config(page_title="Fusion Loans deal model", layout="wide", initial_sidebar_state="expanded")
# light or dark, following the app theme (menu > Settings > Theme)
ui_kit.apply_mode(getattr(st.context.theme, "type", None) or "light")
from ui_kit import *  # noqa: E402  (palette for this run)
st.markdown(CSS, unsafe_allow_html=True)


# =========================================================================== state and run
def base_inputs() -> Inputs:
    return Inputs.from_json(BASE.read_text())


if "inp" not in st.session_state:
    st.session_state.inp = base_inputs()
    st.session_state.saved = {}
    st.session_state.sheet3 = json.loads((ROOT / "notes" / "sheet3.json").read_text())
    st.session_state.notes = json.loads((ROOT / "notes" / "model_notes.json").read_text())
inp: Inputs = st.session_state.inp


@st.cache_data(show_spinner=False, max_entries=64)
def _run(js: str, sens: bool):
    return run_model(Inputs.from_json(js), sensitivities=sens)


def results(sens=False):
    if sens:
        with st.spinner("Running the sensitivity scenarios"):
            return _run(inp.to_json(), True)
    return _run(inp.to_json(), False)


# =========================================================================== table helpers
LBL = lambda tab, r: LAB[tab][str(r)]


def fmt_cell(v, kind):
    if isinstance(v, str) or v is None:
        return "" if v is None else v
    if isinstance(v, (float, np.floating)) and not np.isfinite(v):
        return "n.a."
    return {"$": lambda x: f"{x:,.0f}", "x": lambda x: f"{x:,.2f}x", "%": lambda x: f"{x * 100:,.2f}%",
            "p": lambda x: f"{x:,.6f}", "i": lambda x: f"{x:,.0f}", "r": lambda x: f"{x:,.4f}"}[kind](v)


def wide(rows: list, index, total_col=True, height=None):
    """rows = [(label, values, kind, total?)]; shown like the workbook: rows down, months / ages across."""
    data, idx = [], []
    for item in rows:
        label, vals, kind = item[0], item[1], item[2]
        tot = item[3] if len(item) > 3 else None
        vals = list(vals)
        line = ([fmt_cell(tot, kind) if tot is not None else ""] if total_col else []) + [fmt_cell(v, kind) for v in vals]
        data.append(line); idx.append(label)
    cols = (["Total"] if total_col else []) + [str(i) for i in index]
    df = pd.DataFrame(data, index=idx, columns=cols)
    _show(df, height)


def kv(rows: list, cols=None, height=None):
    """rows = [(label, [values...], kind)] -> small table."""
    data = [[fmt_cell(v, r[2]) for v in r[1]] for r in rows]
    df = pd.DataFrame(data, index=[r[0] for r in rows], columns=cols or ["Value"])
    _show(df, height)


def _show(df, height=None):
    """Row labels as a wide, pinned first column so long Excel labels stay readable."""
    out = df.copy()
    out.insert(0, "Line", out.index.astype(str))
    kw = dict(width="stretch", hide_index=True,
              column_config={"Line": st.column_config.TextColumn("Line", width="large", pinned=True)})
    if height:
        kw["height"] = height
    st.dataframe(out, **kw)


def extras_header():
    st.markdown("---")
    st.markdown("## Additional views (not in the workbook)")


# =========================================================================== sidebar
def sidebar():
    with st.sidebar:
        st.markdown("### Scenario")
        c1, c2 = st.columns([3, 2])
        new = c1.text_input("Save current as", placeholder="Scenario name", label_visibility="collapsed")
        if c2.button("Save", width="stretch") and new:
            st.session_state.saved[new] = inp.to_json()
            st.toast(f"Saved {new}")
        if st.session_state.saved:
            pick = st.selectbox("Saved scenarios", list(st.session_state.saved))
            if st.button("Load scenario", width="stretch"):
                st.session_state.inp = Inputs.from_json(st.session_state.saved[pick])
                clear_widget_state(); st.rerun()
        if st.button("Reset to the base case", width="stretch"):
            st.session_state.inp = base_inputs(); clear_widget_state(); st.rerun()
        st.download_button("Download inputs (JSON)", inp.to_json(), file_name="fusion_inputs.json", width="stretch")
        up = st.file_uploader("Load inputs from JSON or the Excel model", type=["json", "xlsx"])
        if up is not None and st.session_state.get("_last_upload") != up.name + str(up.size):
            st.session_state._last_upload = up.name + str(up.size)
            if up.name.endswith(".json"):
                st.session_state.inp = Inputs.from_json(up.read().decode())
            else:
                with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
                    f.write(up.read()); path = f.name
                st.session_state.inp = load_inputs_from_excel(path)
                _load_workbook_notes(path)
            clear_widget_state(); st.rerun()
        st.markdown("---")
        r = results()
        st.markdown(f"<span class='dot' style='background:{status_color(r['input_status'])}'></span>Inputs "
                    f"<b>{r['input_status']}</b><br><span class='dot' style='background:{status_color(r['model_status'])}'>"
                    f"</span>Model <b>{r['model_status']}</b>", unsafe_allow_html=True)
        note("Every number is recalculated in Python each time an input changes.")
        note("Light or dark: menu (top right) > Settings > Choose app theme. By default it follows your computer's setting.")


def _load_workbook_notes(path):
    import openpyxl, datetime
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        if "Sheet3" in wb.sheetnames:
            s = wb["Sheet3"]
            st.session_state.sheet3 = [[s.cell(r, c).value for c in range(2, 13)] for r in range(2, 6)]
        if "Claude Handoff" in wb.sheetnames:
            ws = wb["Claude Handoff"]; out = []
            for r in range(1, ws.max_row + 1):
                a, b, c = [ws.cell(r, k).value for k in (1, 2, 3)]
                if a is None and b is None:
                    continue
                if isinstance(a, datetime.datetime):
                    a = a.date().isoformat()
                out.append({"row": r, "a": str(a or ""), "b": str(b or ""), "c": str(c or "")})
            st.session_state.notes = out
    except Exception:
        pass


# =========================================================================== Claude Handoff
def page_notes():
    st.title("Model notes")
    lede("Documentation carried over from the Excel model: what it does, settings, verification, open items, the original spec, "
         "open questions, build history and the change log. References like F1 or F3 in it point to tabs in the Excel file.")
    notes = st.session_state.notes
    section, rows, log = None, [], []

    def flush():
        if section and rows:
            st.markdown(f"### {section}")
            st.dataframe(pd.DataFrame(rows, columns=["Item", "Detail"]), hide_index=True, width="stretch")
    for n in notes:
        a, b, c = n["a"], n["b"], n["c"]
        if n["row"] <= 2:
            st.markdown(f"**{a}**"); continue
        if a and not b and a.split(".")[0].strip().isdigit():
            flush(); section, rows = a, []; continue
        if section and section.startswith("11"):
            if a == "Date":
                continue
            log.append([a, b, c]); continue
        rows.append([a, b])
    flush()
    if log:
        st.markdown(f"### {section}")
        st.dataframe(pd.DataFrame(log, columns=["Date", "Change", "Verified"]), hide_index=True, width="stretch", height=600)


# =========================================================================== F0
def page_f0():
    r = results()
    st.title("Fusion Loans: deal overview")
    status_rail(r)
    R, PS, S, A = r["spv"]["returns"], r["pool"]["summary"], r["ue"]["summary"], r["asset"]["table"]
    d = r["derived"]

    # 1. Investor returns: one card per party
    st.markdown("## 1. Investor returns")
    lede("SPV level: after all fees, recycling, leverage and the OCM / Loan America promote.")
    c = st.columns(3)
    la_dist = R.loc["Share of distributions", "Loan America"]
    cards = [("SPV (all investors)", "Total investors", INK, "100% of capital"),
             ("OCM", "OCM", OCM, f"{fp(R.loc['Share of capital','OCM'])} of capital"),
             ("Loan America", "Loan America", LA, f"{fp(R.loc['Share of capital','Loan America'])} of capital, "
                                                  f"{fp(la_dist,0)} of distributions (IRR incl. promote)")]
    for col, (name, party, colr, share) in zip(c, cards):
        with col:
            party_card(name, share, fx(R.loc["MOIC", party]), fp(R.loc["IRR", party]),
                       f"Equity {fm(R.loc['Equity called', party])} in, {fm(R.loc['Distributions', party])} out",
                       colr, R.loc["MOIC", party] < 1 or R.loc["IRR", party] < 0)

    # 2-4. Capital and cash, asset level, school: grouped lists
    st.markdown("## Program, asset and school figures")
    tp = PS["Cash paid to schools"] + PS["Upfront fees paid by investor"]
    c = st.columns(3)
    with c[0]:
        group_card("2. Capital and cash (program level)", [
            ("Equity called", fm(PS["Equity called"]), f"OCM {fm(R.loc['Equity called','OCM'])}, LA {fm(R.loc['Equity called','Loan America'])}"),
            ("Distributions", fm(PS["Distributions"]), f"OCM {fm(R.loc['Distributions','OCM'])}, LA {fm(R.loc['Distributions','Loan America'])}"),
            ("Net profit to investors", fm(PS["Distributions"] - PS["Equity called"]), "distributions minus equity"),
            ("Loan purchases", fm(tp), f"cash to schools + upfront fees; face {fm(PS['Total face purchased'],0)}"),
            ("Funded by recycled collections", fm(PS["Collections recycled into purchases"]),
             f"{fp(PS['Collections recycled into purchases']/tp if tp else 0,0)} of purchases"),
            ("Peak debt", fm(PS["Peak debt"]), f"interest {fm(PS['Total interest'])}"),
        ], INK)
    with c[1]:
        group_card("3. Asset level, per $1 of face (blended)", [
            ("Gross MOIC / IRR", f"{fx(A.loc['Investor MOIC, gross','Blended'])}  /  {fp(A.loc['Investor IRR, gross','Blended'])}", "investor, before fees"),
            ("Net MOIC / IRR", f"{fx(S.loc['Investor MOIC','Blended'])}  /  {fp(S.loc['Investor IRR','Blended'])}",
             f"after fees; outlay ${S.loc['Investor outlay','Blended']:.3f} per $1 of face"),
            ("Fee drag on MOIC", fx(A.loc["Fee drag on MOIC", "Blended"]), "gross minus net"),
            ("Charge-offs", fp(S.loc["Charge-offs", "Blended"]), f"of face; loss multiplier {inp.loss_mult:.2f}x; before SPV costs"),
            ("Advance rate (blended)", fp(d["F71"]), "cash paid to the school per $1 of face"),
            ("Investor split", f"{fp(d['tiers'][0][0],0)} / {fp(1-d['tiers'][0][0],0)}" if len(d["tiers"]) == 1 else f"{len(d['tiers'])} tiers",
             "investor / school share of net collections"),
        ], OCM)
    with c[2]:
        group_card("4. School, % of loan face", [
            ("Day-0 advance", fp(S.loc["School day-0 advance", "Blended"]), "cash at purchase"),
            ("Share of collections", fp(S.loc["School collections", "Blended"]), "undiscounted"),
            ("Total received", fp(S.loc["School total (undiscounted)", "Blended"]), "undiscounted"),
            ("Total received, PV", fp(S.loc["School total (PV)", "Blended"]), f"at a {fp(inp.school_pv_rate)} discount rate"),
            ("Effective discount, PV", fp(S.loc["Effective discount (PV)", "Blended"]), "1 minus total received (PV)"),
            ("School cash, program", fm(r["pool"]["monthly"].school_total.sum()), "advance + share of collections"),
        ], SCHOOL)

    # 5. Deal settings and model status
    st.markdown("## 5. Deal settings and model status")
    lev = f"ON, {fp(inp.ltv,0)} LTV, months {inp.facility_start}-{inp.last_draw}" if inp.leverage else "OFF"
    c = st.columns(3)
    with c[0]:
        group_card("Program", [
            ("Program / purchase period", f"{fm(inp.purchase_amount,0)} / {inp.purchase_period} months"),
            ("Term mix " + " / ".join(str(t) for t in inp.terms), " / ".join(fp(m, 0) for m in inp.mix)),
            ("Loss multiplier", f"{inp.loss_mult:.2f}x"),
            ("Advance rate (blended)", fp(d["F71"])),
        ], FEES)
    with c[1]:
        group_card("Structure", [
            ("Recycling", "ON" if inp.recycling else "OFF"), ("Leverage", lev),
            ("Promote", f"{'ON' if inp.promote else 'OFF'} ({inp.promote_method})"),
            ("Pref / LA tier shares", f"{fp(inp.hurdles[0],0)} pref; LA " + " / ".join(fp(s, 0) for s in inp.la_tier_shares)),
        ], FEES)
    with c[2]:
        group_card("Capital and checks", [
            ("Loan America capital", f"{fp(r['spv']['la_share'],1)} ({inp.capital_basis})"),
            ("Investor / school split", f"{fp(d['tiers'][0][0],0)} / {fp(1-d['tiers'][0][0],0)}" if len(d["tiers"]) == 1 else f"{len(d['tiers'])} tiers"),
            ("Input checks", f"<span class='status' style='color:{status_color(r['input_status'])}'>{r['input_status']}</span>"),
            ("Model checks", f"<span class='status' style='color:{status_color(r['model_status'])}'>{r['model_status']}</span>"),
        ], FEES)

    st.markdown("## 6. Key charts")
    c1, c2 = st.columns(2)
    c1.plotly_chart(C.running_moic(r, 96, la_right=True, title="Running MOIC by party"), width="stretch")
    c2.plotly_chart(C.party_equity_vs_dist(r, "SPV", 96, with_line=False, title="SPV: equity called (below 0) vs distributions (above 0)"), width="stretch")
    note("All 16 charts are on the Charts page. Switch light / dark with the menu (top right) > Settings > Choose app theme.")


# =========================================================================== F1
def page_f1():
    before = inp.to_json()
    r = results()
    d = r["derived"]; chk = d["checks"]
    st.title("Assumptions")
    lede("Every input to the deal, with the calculated values next to them. Percentages are entered as 0 to 100.")
    c1, c2 = st.columns(2)
    with c1:
        st.markdown(f"<span class='dot' style='background:{status_color(r['input_status'])}'></span>**Overall check: {r['input_status']}**, "
                    f"{'all 9 checks pass' if r['input_status']=='OK' else ', '.join(chk['input_failures']) or 'last draw month warning'}", unsafe_allow_html=True)
    with c2:
        st.markdown(f"<span class='dot' style='background:{status_color(r['model_status'])}'></span>**Model checks: "
                    f"{r['model_status']}**, {'unit economics, gross returns, pool and SPV identities all tie' if r['model_status']=='OK' else 'see the Checks page'}",
                    unsafe_allow_html=True)
    tabs = st.tabs(["1. Switches", "2. Loans & credit", "3. Price & fees", "4. Split targets", "5. Program & ramp-up",
                    "6. SPV structure", "7. School", "8. Sensitivity settings"])
    T = [f"{t}m" for t in inp.terms]

    with tabs[0]:
        c = st.columns(2)
        with c[0]:
            inp.recycling = w_toggle("Recycling (ON / OFF): collections fund new purchases", "w_rec", inp.recycling)
            inp.leverage = w_toggle("Leverage (ON / OFF): debt facility on the investor cost basis", "w_lev", inp.leverage)
            inp.losses = w_toggle("Loss curves (ON / OFF); OFF also turns off recoveries", "w_loss", inp.losses)
            inp.prepay = w_toggle("Prepayment curves (ON / OFF)", "w_pp", inp.prepay)
            inp.payment_stop = w_toggle("Payment stop before charge-off (ON / OFF)", "w_stop", inp.payment_stop)
            inp.promote = w_toggle("Promote waterfall OCM / Loan America (ON / OFF)", "w_prom", inp.promote)
        with c[1]:
            inp.target_mode = w_select("Target mode for the investor / school split", "w_mode", inp.target_mode, ["MOIC", "IRR"])
            inp.outlay_basis = w_select("Investor outlay used for the target test", "w_ob", inp.outlay_basis, ["Price + fee", "Price only"])
            inp.after_last_tier = w_select("After the last active tier", "w_alt", inp.after_last_tier, ["Continue split", "100% to school"])
            inp.capital_basis = w_select("Loan America commitment basis", "w_cb", inp.capital_basis, ["OCM %", "Fixed $"])

    with tabs[1]:
        st.markdown("### 2a. Product terms and loan size")
        df = pd.DataFrame({"Term (months)": inp.terms, "APR (annual, %)": [a * 100 for a in inp.apr],
                           "Term mix (% of face)": [m * 100 for m in inp.mix], "Curve family": inp.curve_family,
                           "Average loan size, face ($)": inp.loan_size}, index=T)
        ed = st.data_editor(df, key="w_terms", width="stretch", column_config={
            "Term (months)": st.column_config.NumberColumn(min_value=1, max_value=60, step=1),
            "Curve family": st.column_config.SelectboxColumn(options=inp.curve_names)})
        inp.terms = [int(v) for v in ed["Term (months)"]]
        inp.apr = [float(v) / 100 for v in ed["APR (annual, %)"]]
        inp.mix = [float(v) / 100 for v in ed["Term mix (% of face)"]]
        inp.curve_family = list(ed["Curve family"])
        inp.loan_size = [float(v) for v in ed["Average loan size, face ($)"]]
        kv([("Curve row in the base curves", [i + 1 for i in d["fam_idx"]] + [""], "i"),
            ("Monthly payment per $1 of face (level amortisation)", list(d["pmt"]) + [None], "p"),
            ("APR, weighted", [None] * 4 + [d["wa_apr"]], "%"),
            ("Weighted-average term (months)", [None] * 4 + [d["wa_term"]], "r"),
            ("Term mix total", [None] * 4 + [float(sum(inp.mix))], "%"),
            ("Average loan size, face-weighted ($)", [None] * 4 + [d["avg_loan"]], "$"),
            ("Number of loans in the program by term", list(d["loans_by_term"]) + [sum(d["loans_by_term"])], "$")], cols=T + ["Total / WA"])
        st.markdown(f"Check: term mix sums to 100%: **{chk['term_mix']}**")
        st.markdown("### 2b. Credit: loss / prepayment multipliers, recoveries, payment stop")
        c = st.columns(4)
        with c[0]: inp.loss_mult = w_num("Loss multiplier (x)", "w_lm", float(inp.loss_mult), step=0.05)
        with c[1]: inp.prepay_mult = w_num("Prepayment multiplier (x)", "w_pm", float(inp.prepay_mult), step=0.05)
        with c[2]: inp.recovery_rate = w_pct("Recovery rate (% of booked charge-offs)", "w_rr", inp.recovery_rate)
        with c[3]: inp.recovery_lag = int(w_num("Recovery lag (months)", "w_rl", int(inp.recovery_lag), step=1, min_value=0))
        note(f"Effective loss multiplier used: {d['lm_eff']:.2f}x; effective prepayment multiplier used: {d['pm_eff']:.2f}x "
             "(0 when that curve switch is OFF).")
        sdf = pd.DataFrame([[v * 100 for v in inp.stop_share_short], [v * 100 for v in inp.stop_share_long]],
                           index=["short (24 / 36)", "long (48 / 60)"],
                           columns=["1 month (%)", "2 months (%)", "3 months (%)", "4 months (%)"])
        st.markdown("Payment stop: share of the scheduled payment paid in the months before charge-off")
        sed = st.data_editor(sdf, key="w_stop_tbl", width="stretch")
        inp.stop_share_short = [float(v) / 100 for v in sed.iloc[0]]
        inp.stop_share_long = [float(v) / 100 for v in sed.iloc[1]]
        st.markdown("### 2c. Base curves (cumulative loss % of face and prepayment SMM by loan age; paste from Excel)")
        ages = list(range(61))
        cdf = pd.DataFrame({f"{inp.curve_names[0]} (%)": [v * 100 for v in inp.loss_curves[0]],
                            f"{inp.curve_names[1]} (%)": [v * 100 for v in inp.loss_curves[1]],
                            f"{inp.curve_names[0]} SMM (%)": [v * 100 for v in inp.smm_curves[0]],
                            f"{inp.curve_names[1]} SMM (%)": [v * 100 for v in inp.smm_curves[1]]}, index=pd.Index(ages, name="Loan age"))
        c1, c2 = st.columns([2, 3])
        with c1:
            ced = st.data_editor(cdf, key="w_curves", height=520, width="stretch")
        cols = list(ced.columns)
        inp.loss_curves = [[float(v) / 100 for v in ced[cols[0]]], [float(v) / 100 for v in ced[cols[1]]]]
        inp.smm_curves = [[float(v) / 100 for v in ced[cols[2]]], [float(v) / 100 for v in ced[cols[3]]]]
        with c2:
            st.plotly_chart(line_fig(ages, {inp.curve_names[0]: inp.loss_curves[0], inp.curve_names[1]: inp.loss_curves[1]},
                                     "Base cumulative loss curves", ".0%", 250), width="stretch")
            st.plotly_chart(line_fig(ages, {f"{inp.curve_names[0]} SMM": inp.smm_curves[0], f"{inp.curve_names[1]} SMM": inp.smm_curves[1]},
                                     "Base prepayment SMM", ".1%", 230), width="stretch")
        st.markdown("### 2d. Effective curves by term (after multipliers and the term cut-off)")
        wide([(f"{t}m cumulative loss (% of face)", d["loss"][k], "%") for k, t in enumerate(inp.terms)] +
             [(f"{t}m prepayment SMM", d["smm"][k], "%") for k, t in enumerate(inp.terms)], ages, total_col=False)
        st.plotly_chart(line_fig(ages, {f"{t}m": d["loss"][k] for k, t in enumerate(inp.terms)}, "Effective cumulative loss by term", ".0%", 260), width="stretch")
        st.markdown(f"Check: every effective loss curve is non-decreasing: **{chk['loss_curves']}**")

    with tabs[2]:
        st.markdown("### 3a. Advance rate and outlay by term")
        adf = pd.DataFrame({"Advance rate by term: cash paid to the school per $1 of face (%)": [a * 100 for a in inp.advance]}, index=T).T
        aed = st.data_editor(adf, key="w_adv", width="stretch")
        inp.advance = [float(v) / 100 for v in aed.iloc[0]]
        kv([("Advance rate", list(d["adv"]) + [d["F71"]], "p"),
            ("Upfront fees per $1 of face by term (active upfront fees)", list(d["upfront"]) + [d["F72"]], "p"),
            ("Investor outlay per $1 of face at month 0", list(d["outlay"]) + [d["F73"]], "p"),
            ("Investor outlay per $1 used for the target test", list(d["target_outlay"]) + [float(np.dot(d["target_outlay"], d["mix"]))], "p")],
           cols=T + ["Total / WA"])
        st.markdown("### 3b. Fee table (add a fee: fill a row and set it ON)")
        fdf = pd.DataFrame([{"Active": f.active, "Fee name": f.name, "Paid to": f.paid_to, "Type": f.type or "",
                             "Method": f.method or "", "Basis (if %)": f.basis or "",
                             "Rate (%)": None if f.rate is None else f.rate * 100, "Fixed $/loan": f.fixed} for f in inp.fees],
                           index=pd.Index(range(1, 11), name="Fee #"))
        fed = st.data_editor(fdf, key="w_fees", width="stretch", column_config={
            "Type": st.column_config.SelectboxColumn(options=["Upfront", "Across life", ""]),
            "Method": st.column_config.SelectboxColumn(options=["%", "Fixed $", ""]),
            "Basis (if %)": st.column_config.SelectboxColumn(options=["Net paid", "Face", "Collections", "Beginning balance",
                                                                      "Ending balance", "Average balance", ""])})
        inp.fees = [Fee(active=bool(rw["Active"]), name=str(rw["Fee name"] or ""), paid_to=str(rw["Paid to"] or ""),
                        type=str(rw["Type"] or ""), method=str(rw["Method"] or ""), basis=str(rw["Basis (if %)"] or ""),
                        rate=None if pd.isna(rw["Rate (%)"]) else float(rw["Rate (%)"]) / 100,
                        fixed=None if pd.isna(rw["Fixed $/loan"]) else float(rw["Fixed $/loan"])) for _, rw in fed.iterrows()]
        kdf = []
        for f, ok in zip(inp.fees, d["fee_ok"]):
            k_ = 0.0
            if ok == "OK" and f.type == "Upfront":
                k_ = f.rate * (1 if f.basis == "Face" else d["F71"]) if f.method == "%" else (f.fixed / d["avg_loan"] if d["avg_loan"] else 0)
            kdf.append([ok, f"{k_:.6f}"])
        st.dataframe(pd.DataFrame(kdf, columns=["Row check", "Upfront /$1 (at the mix-weighted advance and loan size)"],
                                  index=pd.Index(range(1, 11), name="Fee #")).T, width="stretch")
        kv([("Total upfront fees per $1 of face (paid by investor at month 0)", [d["F72"]], "p"),
            ("Number of active fees", [sum(1 for x in d["fee_ok"] if x == "OK")], "i"),
            ("Check: every active fee row valid; loan sizes positive", [chk["fee_table"]], "p"),
            ("Across-life coefficient: % of monthly collections", [d["c_coll"]], "%"),
            ("Across-life coefficient: annual % of beginning balance (charged / 12)", [d["c_beg"]], "%"),
            ("Across-life coefficient: annual % of ending balance (charged / 12)", [d["c_end"]], "%"),
            ("Across-life coefficient: annual % of average balance (charged / 12)", [d["c_avg"]], "%"),
            ("Across-life coefficient: fixed $ per active loan per month", [d["c_fix"]], "$")])

    with tabs[3]:
        st.markdown("### 4. Investor / school split targets (add a tier: fill a row and set it ON)")
        tdf = pd.DataFrame([{"Active (ON / OFF)": t.active, "Investor % of net collections": None if t.investor_pct is None else t.investor_pct * 100,
                             "MOIC target (x outlay)": t.moic_target, "IRR target (annual, %)": None if t.irr_target is None else t.irr_target * 100}
                            for t in inp.tiers], index=pd.Index(range(1, 11), name="Tier"))
        ted = st.data_editor(tdf, key="w_tiers", width="stretch")
        inp.tiers = [Tier(active=bool(rw["Active (ON / OFF)"]),
                          investor_pct=None if pd.isna(rw["Investor % of net collections"]) else float(rw["Investor % of net collections"]) / 100,
                          moic_target=None if pd.isna(rw["MOIC target (x outlay)"]) else float(rw["MOIC target (x outlay)"]),
                          irr_target=None if pd.isna(rw["IRR target (annual, %)"]) else float(rw["IRR target (annual, %)"]) / 100)
                     for _, rw in ted.iterrows()]
        helper, order = [], 0
        act = d["tiers"]
        for i, (t, ok) in enumerate(zip(inp.tiers, d["tier_ok"])):
            if ok == "OK":
                order += 1
            inv_if = act[i][0] if i < len(act) else None
            tgt_if = act[i][1] if i < len(act) else None
            helper.append([fp(1 - t.investor_pct, 1) if t.investor_pct is not None else "", ok, order if ok == "OK" else "", i + 1,
                           fp(inv_if, 1) if inv_if is not None else "",
                           ("" if tgt_if is None else (f"{tgt_if:.2f}x" if inp.target_mode == "MOIC" else fp(tgt_if, 1))),
                           "OK" if tgt_if is not None else ""])
        st.dataframe(pd.DataFrame(helper, index=pd.Index(range(1, 11), name="Tier"),
                                  columns=["School %", "Row check", "Active order", "Order #", "Investor % in force",
                                           f"Target in force ({inp.target_mode})", "Order check"]), width="stretch")
        kv([("Number of active tiers", [d["n_tiers"]], "i"),
            ("After the last active tier: investor % (continue = last tier's split; 100% to school = 0%)", [d["after_pct"]], "%"),
            ("Check: at least one active tier, every active row valid", [chk["target_tiers"]], "p")])
        note(f"Target mode {inp.target_mode} picks the column used: MOIC targets are a multiple of the outlay used for the "
             "target test; IRR targets are annual hurdle rates.")

    with tabs[4]:
        st.markdown("### 5. Program plan and ramp-up")
        c = st.columns(4)
        with c[0]:
            fmonth = st.date_input("First purchase month", value=pd.Timestamp(inp.first_month).date(), key="w_fm")
            inp.first_month = fmonth.replace(day=1).isoformat()
        with c[1]: inp.purchase_amount = w_num("Purchase amount: cash paid to schools ($)", "w_pa", float(inp.purchase_amount), step=10_000_000.0, format="%.0f")
        with c[2]: inp.purchase_period = int(w_num("Purchase period (months, 5 to 60)", "w_ppd", int(inp.purchase_period), min_value=1, max_value=60, step=1))
        with c[3]: inp.ramp_months = int(w_num("Ramp-up period (months)", "w_rm", int(inp.ramp_months), min_value=0, max_value=59, step=1))
        if inp.ramp_months > 0:
            rdf = pd.DataFrame([[v * 100 for v in inp.ramp_split[:inp.ramp_months]]], columns=[str(i) for i in range(1, inp.ramp_months + 1)],
                               index=["Ramp-up purchase split inputs (% of purchase amount)"])
            red = st.data_editor(rdf, key=f"w_ramp_{inp.ramp_months}", width="stretch")
            vals = [float(v) / 100 for v in red.iloc[0]]
            inp.ramp_split = vals + inp.ramp_split[len(vals):]
        kv([("Total face purchased ($)", [d["total_face"]], "$")])
        wide([("Monthly purchase split (% of purchase amount)", d["plan_split"], "%", float(d["plan_split"].sum())),
              ("Monthly purchase cash paid to schools ($)", d["plan_cash"], "$", float(d["plan_cash"].sum())),
              ("Monthly face purchased ($) = cash / mix-weighted advance", d["plan_face"], "$", float(d["plan_face"].sum()))],
             range(1, 61))
        st.markdown(f"Checks: purchase split **{chk['purchase_split']}**, purchase cash **{chk['purchase_cash']}**, ramp-up **{chk['ramp_up']}**")

    with tabs[5]:
        st.markdown("### 6a. Investor capital (OCM / Loan America)")
        c = st.columns(3)
        with c[0]: inp.la_fixed = w_num("Loan America fixed commitment ($) (used when the basis is Fixed $)", "w_laf", float(inp.la_fixed), step=250_000.0, format="%.0f")
        with c[1]: inp.ocm_pct = w_pct("OCM share of capital (used when the basis is OCM %)", "w_ocm", inp.ocm_pct)
        kv([("Loan America share of capital (OCM % basis)", [1 - inp.ocm_pct], "%"),
            ("OCM share of capital used", [1 - r["spv"]["la_share"]] if inp.capital_basis == "OCM %" else ["per SPV"], "%"),
            ("Loan America share of capital used", [r["spv"]["la_share"]] if inp.capital_basis == "OCM %" else ["per SPV"], "%")])
        if inp.capital_basis == "Fixed $":
            note(f"Fixed $ basis: Loan America share = commitment ÷ total equity called = {fp(r['spv']['la_share'],2)}.")
        st.markdown("### 6b. Leverage facility (used when the leverage switch is ON)")
        c = st.columns(4)
        with c[0]: inp.ltv = w_pct("Advance rate on eligible cost basis (LTV)", "w_ltv", inp.ltv)
        with c[1]: inp.spread = w_pct("Spread over SOFR", "w_spr", inp.spread, step=0.25)
        with c[2]: inp.ineligibles = w_pct("Ineligibles (% of cost basis)", "w_inel", inp.ineligibles)
        with c[3]: inp.sweep = w_pct("Cash sweep in amortisation (% of available cash)", "w_sw", inp.sweep, step=5.0)
        c = st.columns(4)
        with c[0]: inp.facility_start = int(w_num("Facility start month (month of the projection)", "w_fs", int(inp.facility_start), min_value=1, step=1))
        with c[1]: inp.last_draw = int(w_num("Last draw month; after it the facility amortises", "w_ld", int(inp.last_draw), min_value=1, step=1))
        kv([("Maximum facility ($) = LTV x purchase amount", [d["facility_max"]], "$"),
            ("Check: last draw month on or after the start month", [chk["leverage_months"]], "p"),
            ("Warning: last draw month covers the purchase period", [chk["last_draw_warning"] + ("" if chk["last_draw_warning"] == "OK" else
             ": last draw month is before the purchase period ends; later purchases need more equity")], "p")])
        st.markdown("SOFR curve (annual, by month of the projection)")
        sdf = pd.DataFrame([[v * 100 for v in inp.sofr]], columns=[str(i) for i in range(1, 121)], index=["SOFR (%)"])
        sed = st.data_editor(sdf, key="w_sofr", width="stretch")
        inp.sofr = [float(v) / 100 for v in sed.iloc[0]]
        st.markdown("### 6c. Other SPV fees")
        c = st.columns(3)
        with c[0]: inp.fee_years_high = int(w_num("Years at the higher fee", "w_fyh", int(inp.fee_years_high), step=1, min_value=0))
        with c[1]: inp.fee_high = w_num("Other fees per year, first years ($)", "w_fh", float(inp.fee_high), step=100_000.0, format="%.0f")
        with c[2]: inp.fee_low = w_num("Other fees per year, thereafter ($)", "w_fl", float(inp.fee_low), step=50_000.0, format="%.0f")
        st.markdown("### 6d. Promote waterfall OCM / Loan America (used when the promote switch is ON)")
        pdf = pd.DataFrame({"1. Return of capital": [None, None], "2. Pref return": [inp.hurdles[0] * 100, None],
                            "3. To OCM IRR": [inp.hurdles[1] * 100, inp.la_tier_shares[0] * 100],
                            "4. To OCM IRR": [inp.hurdles[2] * 100, inp.la_tier_shares[1] * 100],
                            "5. Residual": [None, inp.la_tier_shares[2] * 100]},
                           index=["Hurdle rate (IRR, %)", "Loan America share in the tier (%)"])
        ped = st.data_editor(pdf, key="w_prom_tbl", width="stretch")
        inp.hurdles = [float(ped.iloc[0, 1]) / 100, float(ped.iloc[0, 2]) / 100, float(ped.iloc[0, 3]) / 100]
        inp.la_tier_shares = [float(ped.iloc[1, 2]) / 100, float(ped.iloc[1, 3]) / 100, float(ped.iloc[1, 4]) / 100]
        s_eff = r["spv"]["effective_la_shares"]
        st.dataframe(pd.DataFrame([["Pro rata", "Pro rata", fp(1 - s_eff[0]), fp(1 - s_eff[1]), fp(1 - s_eff[2])]],
                                  columns=list(pdf.columns), index=["OCM share in the tier (effective, after the promote method)"]),
                     width="stretch")
        note("Tiers 1-2 pay pro rata by capital. Tiers 3-4 run until OCM reaches the hurdle IRR; tier 5 is the residual.")
        inp.promote_method = w_select("Promote method (Absolute / On top of pro rata)", "w_pmeth", inp.promote_method, ["Absolute", "On top of pro rata"])

    with tabs[6]:
        st.markdown("### 7. School")
        inp.school_pv_rate = w_pct("School discount rate for PV of school cash flows (annual)", "w_spv", inp.school_pv_rate, step=0.25)

    with tabs[7]:
        st.markdown("### 8. Sensitivity settings")
        c = st.columns(2)
        with c[0]: inp.advance_step = w_pct("Advance-rate step size (advance-rate sensitivity)", "w_as", inp.advance_step)
        with c[1]: inp.split_step = w_pct("Investor / school split step size (split sensitivity)", "w_ss", inp.split_step)

    if inp.to_json() != before:
        n = st.session_state.get("_rr", 0)
        if n < 2:
            st.session_state._rr = n + 1
            st.rerun()
    st.session_state._rr = 0


# =========================================================================== F2
BLOCK_ROWS = [  # (excel row, df column, kind)
    (89, "sched_balance", "p"), (90, "sched_factor", "p"), (91, "loss_curve", "p"), (92, "smm", "p"), (93, "opening", "p"),
    (94, "charge_offs", "p"), (95, "surviving", "p"), (96, "interest", "p"), (97, "sched_principal", "p"), (98, "prepaid", "p"),
    (99, "ending", "p"), (100, "active_loans", "p"), (101, "recoveries", "p"), (102, "missed", "p"), (103, "prepay_survival", "p"),
    (104, "survivors_stop_off", "p"), (105, "defaulting_balance", "p"), (106, "defaulted_face", "p"), (107, "removed", "p"),
    (108, "gross", "p"), (109, "fees", "p"), (110, "net", "p")] + \
    [(112 + i, f"need_tier_{i + 1}", "p") for i in range(10)] + [(122 + i, f"breakpoint_{i}", "p") for i in range(11)] + \
    [(133 + i, f"alloc_tier_{i + 1}", "p") for i in range(10)] + [(143, "alloc_after_last", "p"), (144, "investor", "p"), (145, "school", "p"),
    (146, "tiers_met", "i"), (147, "tier_in_force", "i"), (148, "tier_in_force_label", "p"), (149, "investor_pct_in_force", "%"),
    (150, "target_hit_flag", "p"), (151, "need_next", "p"), (153, "investor_cf", "p"), (154, "investor_cum", "p"),
    (155, "moic_running", "r"), (156, "irr_running", "%"), (157, "school_cf", "p"), (158, "school_cum", "p"),
    (159, "upfront_fees_paid", "p"), (160, "across_life_fees_paid", "p")]
SUM_ROWS = {"charge_offs", "interest", "sched_principal", "prepaid", "recoveries", "missed", "removed", "gross", "fees", "net",
            "investor", "school", "investor_cf", "school_cf", "upfront_fees_paid", "across_life_fees_paid"} | \
           {f"alloc_tier_{i}" for i in range(1, 11)} | {"alloc_after_last"}
LAST_ROWS = {"ending", "investor_cum", "school_cum", "moic_running", "tiers_met"}


def page_f2():
    r = results()
    st.title("Loan economics")
    status_rail(r)
    lede("Per $1 of face, by term and blended. Month 0 is the purchase month; the investor pays the advance plus upfront fees, "
         "across-life fees come off the top, and net collections are split investor / school tier by tier, tested per term.")
    S = r["ue"]["summary"]; names = list(S.columns)
    st.markdown("## Summary (per $1 of face)")
    g = lambda k: [S.loc[k, n] for n in names]
    L = lambda n: LBL("F2_summary", n)
    rows = [(L(5), g("Term (months)"), "r"), (L(6), g("Mix weight"), "%"), (L(7), g("APR"), "%"),
            (L(8), g("Advance paid to school"), "p"), (L(9), g("Upfront fees"), "p"), (L(10), g("Investor outlay"), "p"),
            (L(11), g("Gross collections"), "p"), (L(12), g("Across-life fees"), "p"), (L(13), g("Net collections"), "p"),
            (L(14), g("Charge-offs"), "p"), (L(15), g("Missed payments"), "p"), (L(16), g("Investor cash received"), "p"),
            (L(17), g("Investor MOIC"), "x"), (L(18), g("Investor IRR"), "%"), (L(19), g("Tiers hit"), "p"),
            (L(20), g("Month tier 1 hit"), "i"), (L(21), [""] * 5, "p"),
            (L(22), g("School day-0 advance"), "p"), (L(23), g("School collections"), "p"), (L(24), g("School total (undiscounted)"), "p"),
            (L(25), g("School PV of collections"), "p"), (L(26), g("School total (PV)"), "p"), (L(27), g("School day-0 advance"), "%"),
            (L(28), g("School collections"), "%"), (L(29), g("School total (undiscounted)"), "%"), (L(30), g("School PV of collections"), "%"),
            (L(31), g("School total (PV)"), "%"), (L(32), g("Effective discount (undiscounted)"), "%"), (L(33), g("Effective discount (PV)"), "%"),
            (L(34), [inp.school_pv_rate] * 5, "%"), (L(35), list(r["ue"]["checks"].values()), "p")]
    kv(rows, cols=names, height=1130)

    st.markdown("## Tier summary: loan age each target is hit, and cash within each tier")
    d = r["derived"]; blocks = r["ue"]["blocks"]
    cols = ["Investor %", "Target in force"]
    for t in inp.terms:
        cols += [f"{t}m age hit", f"{t}m investor $", f"{t}m school $", f"{t}m inv. MOIC at hit", f"{t}m inv. IRR at hit"]
    rows = []
    for i in range(10):
        if i < d["n_tiers"]:
            pct, tgt = d["tiers"][i]
            line = [fp(pct, 1), f"{tgt:.2f}x MOIC" if inp.target_mode == "MOIC" else f"{tgt:.1%} IRR"]
            for k in range(4):
                df = blocks[k]["df"]; met = df.tiers_met.values
                age = int(np.argmax(met >= i + 1)) if (met >= i + 1).any() else None
                a_ = df[f"alloc_tier_{i + 1}"].sum()
                line += [age if age is not None else "Not hit", f"{pct * a_:.6f}", f"{(1 - pct) * a_:.6f}",
                         fx(df.moic_running.values[age], 4) if age is not None else "", fp(df.irr_running.values[age], 2) if age is not None else ""]
            rows.append((f"Tier {i + 1}", line))
        else:
            rows.append((f"Tier {i + 1} (off)", [""] * len(cols)))
    after, tot = [fp(d["after_pct"], 1), ""], ["", ""]
    for k in range(4):
        df = blocks[k]["df"]; a_ = df["alloc_after_last"].sum()
        after += ["", f"{d['after_pct'] * a_:.6f}", f"{(1 - d['after_pct']) * a_:.6f}", "", ""]
        tot += ["", f"{df.investor.sum():.6f}", f"{df.school.sum():.6f}", "Ties" if r["ue"]["checks"][f"{inp.terms[k]}m"] == "OK" else "CHECK", ""]
    rows += [("After last tier", after), ("Total (ties to investor / school share)", tot)]
    st.dataframe(pd.DataFrame([x[1] for x in rows], index=[x[0] for x in rows], columns=cols), width="stretch")

    st.markdown("## " + LBL("F2_blended", 56))
    bl = r["ue"]["blended"]
    brow = [(58, "opening"), (59, "charge_offs"), (60, "interest"), (61, "sched_principal"), (62, "prepaid"), (63, "ending"),
            (64, "recoveries"), (65, "missed"), (66, "removed"), (67, "gross"), (68, "fees"), (69, "net"), (70, "investor"),
            (71, "school"), (72, "investor_cf"), (73, "investor_cum"), (74, "moic_running"), (75, "irr_running"),
            (76, "school_cf"), (77, "school_cum")]
    items = []
    for rr, col in brow:
        kind = "%" if col == "irr_running" else ("r" if col == "moic_running" else "p")
        tot = (float(bl[col].sum()) if col in SUM_ROWS else float(bl[col].values[-1]) if col in LAST_ROWS else
               S.loc["Investor IRR", "Blended"] if col == "irr_running" else None)
        items.append((LBL("F2_blended", rr), bl[col].values, kind, tot))
    c78 = (bl.investor + bl.school - bl.net).round(9)
    items.append((LBL("F2_blended", 78), c78.values, "p", float(c78.abs().sum())))
    items.append((LBL("F2_blended", 80), bl["cost_basis"].values, "p"))
    wide(items, bl.index, height=780)

    st.markdown("## Term blocks: the full monthly engine")
    pick = st.radio("Term", [f"{t}m" for t in inp.terms], horizontal=True, label_visibility="collapsed")
    k = [f"{t}m" for t in inp.terms].index(pick)
    df = blocks[k]["df"]
    kv([(LBL("F2_block", 85), [inp.terms[k]], "i"), (LBL("F2_block", 86), [d["r"][k]], "p"),
        (LBL("F2_block", 87), [d["pmt"][k]], "p"), (LBL("F2_block", 88), [inp.mix[k]], "%")])
    items = []
    for rr, col, kind in BLOCK_ROWS:
        tot = (float(df[col].sum()) if col in SUM_ROWS else float(df[col].values[-1]) if col in LAST_ROWS else
               S.loc["Investor IRR", pick] if col == "irr_running" else None)
        items.append((LBL("F2_block", rr), df[col].values, kind, tot))
    c161 = (df.investor + df.school - df.net).round(9)
    items.append((LBL("F2_block", 161), c161.values, "p", float(c161.abs().sum())))
    c162 = round(df.removed.sum() + df.sched_principal.sum() + df.prepaid.sum() + df.ending.values[-1] - 1, 9)
    items.append((LBL("F2_block", 162), [None] * len(df), "p", c162))
    wide(items, df.index, height=900)
    st.download_button("Download this term block (CSV)", df.to_csv().encode(), file_name=f"F2_{pick}_block.csv")

    extras_header()
    pk = st.radio("Charts for", [f"{t}m" for t in inp.terms] + ["Blended"], horizontal=True, key="f2_extra")
    df = bl if pk == "Blended" else blocks[[f"{t}m" for t in inp.terms].index(pk)]["df"]
    ages = df.index
    c1, c2 = st.columns(2)
    c1.plotly_chart(bar_fig(ages, {"Interest": df.interest, "Scheduled principal": df.sched_principal, "Prepaid": df.prepaid,
                                   "Missed payments": -df.missed}, f"Collections per $1, {pk}",
                            colors={"Interest": OCM, "Scheduled principal": OCM_LIGHT, "Prepaid": LA, "Missed payments": BAD_C}), width="stretch")
    c2.plotly_chart(bar_fig(ages, {"Investor": df.investor, "School": df.school, "Fees": df.fees}, f"Who gets the collections, {pk}",
                            colors={"Investor": OCM, "School": SCHOOL, "Fees": FEES}), width="stretch")
    c1, c2 = st.columns(2)
    fig = line_fig(ages, {"Performing balance": df.ending}, f"Balance and charge-offs, {pk}", height=320, colors={"Performing balance": OCM})
    fig.add_trace(go.Bar(x=ages, y=df.charge_offs, name="Charge-offs", marker_color=BAD_C, yaxis="y2", opacity=0.7))
    fig.update_layout(yaxis2=dict(overlaying="y", side="right", showgrid=False, tickformat=".1%"), yaxis=dict(tickformat=".0%"))
    c1.plotly_chart(fig, width="stretch")
    fig = line_fig(ages, {"Running MOIC": df.moic_running}, f"Investor running MOIC and IRR, {pk}", height=320, colors={"Running MOIC": OCM})
    fig.add_trace(go.Scatter(x=ages, y=pd.Series(df.irr_running).clip(lower=-0.5), name="Running IRR (floor -50%)", yaxis="y2",
                             line=dict(color=LA, width=2, dash="dot")))
    fig.update_layout(yaxis2=dict(overlaying="y", side="right", showgrid=False, tickformat=".0%"))
    c2.plotly_chart(fig, width="stretch")


# =========================================================================== F2a
def page_f2a():
    r = results()
    st.title("Gross vs net returns")
    status_rail(r)
    lede("GROSS (no fees at all; tier split re-run on gross collections, targets on price only) vs NET (after all Fusion fees, "
         "from unit economics). SPV-level costs (other SPV fees, leverage, recycling) are excluded from both.")
    A = r["asset"]["table"]; names = list(A.columns)
    g = lambda k: [A.loc[k, n] for n in names]
    kv([("Mix weight (% of face)", g("Mix weight"), "%"), ("Price paid to the school (advance)", g("Price paid (advance)"), "p"),
        ("Gross collections (no fees)", g("Gross collections"), "p"), ("Investor cash received, gross", g("Investor cash, gross"), "p"),
        ("Investor MOIC, GROSS (pre-fees)", g("Investor MOIC, gross"), "x"), ("Investor IRR, GROSS (pre-fees, annualised)", g("Investor IRR, gross"), "%"),
        ("Investor MOIC, NET (after fees; loan economics)", g("Investor MOIC, net"), "x"),
        ("Investor IRR, NET (after fees; loan economics)", g("Investor IRR, net"), "%"),
        ("Fee drag on investor MOIC (gross − net)", g("Fee drag on MOIC"), "x"),
        ("School total received, GROSS (advance + share of gross collections, % of face)", g("School total, gross"), "%"),
        ("School total received, NET (after fees; loan economics)", g("School total, net"), "%"),
        ("Check: gross investor + school share − gross collections (0)", [0.0] * 4 + [0.0 if r["asset"]["check"] == "OK" else 1.0], "p")],
       cols=names)
    st.markdown("### Blended gross investor cash flow (mix-weighted; for the blended IRR)")
    wide([("Investor cash flow, gross, blended", r["asset"]["gross_cf_blended"], "p", float(r["asset"]["gross_cf_blended"].sum()))], range(61))
    extras_header()
    c1, c2 = st.columns(2)
    c1.plotly_chart(bar_fig(names, {"Gross": A.loc["Investor MOIC, gross"], "Net": A.loc["Investor MOIC, net"]}, "Investor MOIC, gross vs net",
                            colors={"Gross": OCM_LIGHT, "Net": OCM}, barmode="group"), width="stretch")
    c2.plotly_chart(bar_fig(names, {"Gross": A.loc["Investor IRR, gross"], "Net": A.loc["Investor IRR, net"]}, "Investor IRR, gross vs net", ".0%",
                            colors={"Gross": OCM_LIGHT, "Net": OCM}, barmode="group"), width="stretch")


# =========================================================================== F2b
def page_f2b():
    r = results(sens=True)
    st.title("Sensitivity: advance rate")
    status_rail(r)
    lede("Investor and school asset returns per $1 of face. Each column shifts every term's advance rate by the change shown; "
         "upfront fees are recomputed. The middle column (0) is the current case.")
    kv([("Advance-rate step size (set on Assumptions, section 8)", [inp.advance_step], "%"), ("Steps either side of the current advance rate", [5], "i"),
        ("Check: single constant investor / school split (required for this sensitivity)", ["OK" if r.get("single_split") else "CHECK"], "p")])
    if not r.get("single_split"):
        st.warning("This sensitivity needs one constant investor / school split (one tier, continue split), as in the workbook.")
        return
    a = r["adv_sens"]; names = list(r["ue"]["summary"].columns)
    cols = [str(i) for i in a.columns]
    blend = lambda n: n if n != "Blended" else "Blended (mix-weighted)"
    rows = [("Change in advance rate (pp)", a.loc["Change (pp)"], "%"), ("Advance rate (blended; each term shifted by the same change)", a.loc["Blended advance"], "%"),
            ("Discount taken at day 0 (1 − advance rate, blended)", a.loc["Day-0 discount"], "%"), ("A. Investor MOIC per $1 of face (net of all fees)", [""] * 11, "p")]
    rows += [(f"MOIC {blend(n)}", a.loc[f"MOIC {n}"], "x") for n in names]
    rows += [("B. Investor IRR per $1 of face (annualised, net of all fees)", [""] * 11, "p")]
    rows += [(f"IRR {blend(n)}", a.loc[f"IRR {n}"], "%") for n in names]
    rows += [("C. School (blended, per $1 of face)", [""] * 11, "p"),
             ("School total received, undiscounted (% of face: day-0 advance + share of net collections)", a.loc["School total (undiscounted)"], "%"),
             ("School total received, PV at the school discount rate (% of face)", a.loc["School total (PV)"], "%"),
             ("School effective discount, PV (1 − PV received)", a.loc["School effective discount (PV)"], "%")]
    kv(rows, cols=cols, height=720)
    sc = r["sens_checks"]["F2b"]
    kv([("Check: middle column = loan economics (blended MOIC, IRR, school) (0)", [sc["middle = blended F2"]], "p"),
        ("Check: middle column MOIC by term = loan economics (0)", [sc["middle MOIC by term"]], "p"),
        ("Check: middle column IRR by term = loan economics (0)", [sc["middle IRR by term"]], "p"),
        ("Overall sensitivity check", [sc["status"]], "p")])
    note("The two charts for this table are on the Charts page, section 6.")


# =========================================================================== Sheet3
def page_sheet3():
    st.title("Scenario notes")
    lede("Your own scenario table (typed values, not linked to the model). Edit it here; edits last for this session.")
    g = st.session_state.sheet3
    df = pd.DataFrame(g, columns=[chr(66 + i) for i in range(len(g[0]))]).astype(str).replace("None", "")
    ed = st.data_editor(df, key="w_sheet3", width="stretch", hide_index=True)
    st.session_state.sheet3 = ed.values.tolist()


# =========================================================================== F3
def page_f3():
    r = results()
    st.title("Pool cash flows")
    status_rail(r)
    lede("Every monthly purchase is a vintage. Each vintage follows the per-$1 pattern from loan economics at its own age, and the pool is the sum of all vintages. Then SPV fees, the debt facility and recycling. All $ amounts, months 1 to 120.")
    p = r["pool"]["monthly"]; M = p.index; S = r["pool"]["summary"]
    tot = lambda s: float(np.sum(s))
    st.markdown("### 1. Purchases")
    wide([("Month end", [dd.isoformat() for dd in p.date], "p"), (LBL("F3", 7), p.face, "$", tot(p.face)),
          (LBL("F3", 8), p.cash_to_school, "$", tot(p.cash_to_school)), (LBL("F3", 9), p.upfront_fees, "$", tot(p.upfront_fees)),
          (LBL("F3", 10), p.outlay, "$", tot(p.outlay))], M)
    st.markdown("### 2. Pool loan cash flows (all vintages and terms)")
    c26 = (p.gross - (p.interest + p.sched_principal + p.prepaid + p.recoveries - p.missed)).round(4)
    c27 = (p.investor + p.school - p.net).round(4)
    prev = np.r_[0.0, p.ending.values[:-1]]
    c28 = pd.Series(p.ending.values - (prev + p.face.values - p.removed.values - p.sched_principal.values - p.prepaid.values), index=M).round(4)
    wide([(LBL("F3", 13), p.opening, "$"), (LBL("F3", 14), p.charge_offs, "$", tot(p.charge_offs)), (LBL("F3", 15), p.interest, "$", tot(p.interest)),
          (LBL("F3", 16), p.sched_principal, "$", tot(p.sched_principal)), (LBL("F3", 17), p.prepaid, "$", tot(p.prepaid)),
          (LBL("F3", 18), p.recoveries, "$", tot(p.recoveries)), (LBL("F3", 19), p.missed, "$", tot(p.missed)), (LBL("F3", 20), p.gross, "$", tot(p.gross)),
          (LBL("F3", 21), p.fees, "$", tot(p.fees)), (LBL("F3", 22), p.net, "$", tot(p.net)), (LBL("F3", 23), p.investor, "$", tot(p.investor)),
          (LBL("F3", 24), p.school, "$", tot(p.school)), (LBL("F3", 25), p.ending, "$", float(p.ending.values[-1])),
          (LBL("F3", 26), c26, "$", tot(c26.abs())), (LBL("F3", 27), c27, "$", tot(c27.abs())), (LBL("F3", 28), c28, "$", tot(c28.abs()))], M)
    st.markdown("### 3. School")
    wide([(LBL("F3", 31), p.cash_to_school, "$", tot(p.cash_to_school)), (LBL("F3", 32), p.school, "$", tot(p.school)),
          (LBL("F3", 33), p.school_total, "$", tot(p.school_total)), (LBL("F3", 34), p.school_cum, "$")], M)
    st.markdown("### 4. Investor cash before SPV structure (unlevered, no other fees)")
    wide([(LBL("F3", 37), p.investor, "$", tot(p.investor)), (LBL("F3", 38), -p.outlay, "$", -tot(p.outlay)),
          (LBL("F3", 39), p.unlev_cf, "$", tot(p.unlev_cf)), (LBL("F3", 40), p.unlev_cf.cumsum(), "$")], M)
    st.markdown("### 5. Other SPV fees")
    wide([(LBL("F3", 43), p.active, "i", tot(p.active)), (LBL("F3", 44), p.other_fees, "$", tot(p.other_fees))], M)
    st.markdown("### 6. Leverage facility")
    wide([(LBL("F3", 47), p["mode"], "p"), (LBL("F3", 48), p.cost_basis, "$"), (LBL("F3", 49), p.eligible, "$"), (LBL("F3", 50), p.capacity, "$"),
          (LBL("F3", 51), p.sofr, "%"), (LBL("F3", 52), p.rate, "%"), (LBL("F3", 53), p.debt_open, "$"),
          (LBL("F3", 54), p.interest_debt, "$", tot(p.interest_debt)), (LBL("F3", 55), p.draws, "$", tot(p.draws)),
          (LBL("F3", 56), p.repay, "$", tot(p.repay)), (LBL("F3", 57), p.debt_end, "$", float(p.debt_end.values[-1]))], M)
    st.markdown("### 7. SPV cash and recycling")
    c72 = (p.dist - p.equity - p.net_spv).round(4)
    wide([(LBL("F3", 61), p.investor, "$", tot(p.investor)), (LBL("F3", 62), p.draws, "$", tot(p.draws)), (LBL("F3", 63), -p.outlay, "$", -tot(p.outlay)),
          (LBL("F3", 64), p.other_fees, "$", tot(p.other_fees)), (LBL("F3", 65), -p.interest_debt, "$", -tot(p.interest_debt)),
          (LBL("F3", 66), -p.repay, "$", -tot(p.repay)), (LBL("F3", 67), p.net_spv, "$", tot(p.net_spv)), (LBL("F3", 68), p.operating, "$", tot(p.operating)),
          (LBL("F3", 69), p.equity, "$", tot(p.equity)), (LBL("F3", 70), p.dist, "$", tot(p.dist)), (LBL("F3", 71), p.recycled, "$", tot(p.recycled)),
          (LBL("F3", 72), c72, "$", tot(c72.abs())), (LBL("F3", 73), p.net_to_investors, "$", tot(p.net_to_investors))], M)
    st.markdown("### 8. Pool summary")
    ck = r["pool"]["checks"]
    kv([(LBL("F3", 76), [S["Total face purchased"]], "$"), (LBL("F3", 77), [S["Cash paid to schools"]], "$"),
        (LBL("F3", 78), [S["Upfront fees paid by investor"]], "$"), (LBL("F3", 79), [S["Gross collections"]], "$"),
        (LBL("F3", 80), [S["Across-life fees"]], "$"), (LBL("F3", 81), [S["Net collections"]], "$"),
        (LBL("F3", 82), [S["Investor share of net collections"]], "$"), (LBL("F3", 83), [S["School share of net collections"]], "$"),
        (LBL("F3", 84), [S["Charge-offs"]], "$"), (LBL("F3", 85), [S["Charge-offs (% of face)"]], "%"),
        (LBL("F3", 86), [S["Unlevered investor MOIC"]], "x"), (LBL("F3", 87), [S["Unlevered investor IRR"]], "%"),
        (LBL("F3", 88), [ck["tie_to_F2"]], "p"), (LBL("F3", 89), [S["School effective advance (undiscounted)"]], "%"),
        (LBL("F3", 90), [0.0], "p"), (LBL("F3", 91), [S["Peak debt"]], "$"), (LBL("F3", 92), [S["Total debt draws"]], "$"),
        (LBL("F3", 93), [S["Total interest"]], "$"), (LBL("F3", 94), [S["Equity called"]], "$"), (LBL("F3", 95), [S["Distributions"]], "$"),
        (LBL("F3", 96), [S["Collections recycled into purchases"]], "$"), (LBL("F3", 97), [S["Equity MOIC"]], "x"),
        (LBL("F3", 98), [S["Equity IRR"]], "%"), (LBL("F3", 99), [ck["status"]], "p")], height=890)
    nv = r["pool"]["vintage_investor"].shape[0]
    for title, grid, total, lab in [("9. Vintage view: investor share of net collections by monthly vintage", r["pool"]["vintage_investor"], p.investor,
                                     "Check: total of vintages - pool investor share (0)"),
                                    ("10. Vintage view: school cash by monthly vintage (advance at purchase + share of collections)", r["pool"]["vintage_school"],
                                     p.school_total, "Check: total of vintages - pool school cash (0)")]:
        st.markdown(f"### {title}")
        items = [(f"Vintage {v + 1}", grid[v], "$", float(grid[v].sum())) for v in range(nv)]
        items.append(("Total of vintages", grid.sum(0), "$", float(grid.sum())))
        chk = np.round(grid.sum(0) - total.values, 4)
        items.append((lab, chk, "$", float(np.abs(chk).sum())))
        wide(items, M, height=520)
    extras_header()
    last = int(np.max(np.nonzero(p.active.values)[0])) + 1 if p.active.sum() else 120
    x = p.index[:last]
    c1, c2 = st.columns(2)
    c1.plotly_chart(bar_fig(x, {"Investor": p.investor[:last], "School": p.school[:last], "Fees": p.fees[:last]}, "Net collections split by month ($)",
                            colors={"Investor": OCM, "School": SCHOOL, "Fees": FEES}), width="stretch")
    c2.plotly_chart(line_fig(x, {"Performing balance": p.ending[:last], "Debt": p.debt_end[:last], "Borrowing capacity": p.capacity[:last]},
                             "Balance, capacity and debt ($)", colors={"Performing balance": OCM, "Debt": DEBT, "Borrowing capacity": FEES},
                             dash={"Borrowing capacity": "dot"}), width="stretch")
    vi = r["pool"]["vintage_investor"]
    fig = go.Figure(go.Heatmap(z=vi[:, :last], x=list(x), y=[f"V{i}" for i in range(1, vi.shape[0] + 1)],
                               colorscale=[[0, HEAT0], [0.4, HEAT1], [1, OCM]], colorbar=dict(title="$")))
    fig.update_layout(height=520, yaxis=dict(autorange="reversed"), title="Investor cash by vintage (rows) and month (columns)")
    st.plotly_chart(fig, width="stretch")
    st.download_button("Download the pool table (CSV)", p.to_csv().encode(), file_name="F3_pool.csv")


# =========================================================================== F4
def page_f4():
    r = results()
    st.title("SPV and waterfall")
    status_rail(r)
    lede("Equity calls and distributions from the pool, split between OCM and Loan America by capital and the promote waterfall, with returns by party.")
    sp = r["spv"]; R, w = sp["returns"], sp["monthly"]; M = w.index
    st.markdown("### 1. Capital split and waterfall settings")
    kv([(LBL("F4", 7), [inp.capital_basis, "", ""], "p"), (LBL("F4", 8), [sp["la_share"], None, None], "%"),
        (LBL("F4", 9), [1 - sp["la_share"], None, None], "%"), (LBL("F4", 10), ["ON" if inp.promote else "OFF", "", ""], "p"),
        (LBL("F4", 11), sp["monthly_hurdles"], "%"), (LBL("F4", 12), sp["effective_la_shares"], "%"),
        ("Check: effective Loan America shares within 0-100%", [sp["checks"]["promote_shares"], "", ""], "p"),
        (LBL("F4", 13), [sp["checks"]["waterfall"], "", ""], "p")], cols=["B", "C", "D"])
    st.markdown("### 2. Returns by party")
    T = sp["tiers"]
    rows = [(LBL("F4", 15), R.loc["Share of capital"], "%"), (LBL("F4", 16), R.loc["Equity called"], "$"), (LBL("F4", 17), R.loc["Distributions"], "$"),
            (LBL("F4", 18), R.loc["Net profit"], "$"), (LBL("F4", 19), R.loc["MOIC"], "x"), (LBL("F4", 20), R.loc["IRR"], "%"),
            (LBL("F4", 21), R.loc["Share of distributions"], "%"), (LBL("F4", 22), R.loc["Debt recap"], "$"), (LBL("F4", 23), R.loc["Adjusted MOIC"], "x")]
    rows += [(LBL("F4", 24 + i), T.iloc[i], "$") for i in range(5)]
    kv(rows, cols=list(R.columns), height=540)
    st.markdown("### 3. Capital called (monthly)")
    wide([("Month end", [dd.isoformat() for dd in w.date], "p"), (LBL("F4", 30), w.equity, "$", float(w.equity.sum())),
          (LBL("F4", 31), w.ocm_equity, "$", float(w.ocm_equity.sum())), (LBL("F4", 32), w.la_equity, "$", float(w.la_equity.sum())),
          (LBL("F4", 33), w.equity.cumsum(), "$")], M)
    st.markdown("### 4. Promote waterfall on investor distributions (monthly)")
    ocm12 = (w.t1 + w.t2) * (1 - sp["la_share"])
    wide([(LBL("F4", 36), w.dist, "$", float(w.dist.sum())), (LBL("F4", 38), w.t1, "$", float(w.t1.sum())), (LBL("F4", 39), w.t1.cumsum(), "$"),
          (LBL("F4", 40), w.hb2, "$"), (LBL("F4", 41), w.t1 + w.t2, "$", float((w.t1 + w.t2).sum())), (LBL("F4", 42), w.t2, "$", float(w.t2.sum())),
          (LBL("F4", 43), w.hb2 - w.dist, "$"), (LBL("F4", 44), ocm12, "$", float(ocm12.sum())), (LBL("F4", 45), w.hb3, "$"),
          (LBL("F4", 46), w.t3, "$", float(w.t3.sum())), (LBL("F4", 47), w.hb4, "$"), (LBL("F4", 48), w.t4, "$", float(w.t4.sum())),
          (LBL("F4", 49), w.t5, "$", float(w.t5.sum())), (LBL("F4", 50), w.ocm, "$", float(w.ocm.sum())), (LBL("F4", 51), w.la, "$", float(w.la.sum())),
          (LBL("F4", 52), w.hb3 - w.ocm, "$"), (LBL("F4", 53), w.hb4 - w.ocm, "$"),
          (LBL("F4", 54), [("" if x == 0 else x) for x in w.highest_tier], "p")], M)
    st.markdown("### 5. Loan America distributions by tier")
    c61 = (w.la_t12 + w.la_t3 + w.la_t4 + w.la_t5 - w.la).round(4)
    wide([(LBL("F4", 57), w.la_t12, "$", float(w.la_t12.sum())), (LBL("F4", 58), w.la_t3, "$", float(w.la_t3.sum())),
          (LBL("F4", 59), w.la_t4, "$", float(w.la_t4.sum())), (LBL("F4", 60), w.la_t5, "$", float(w.la_t5.sum())),
          (LBL("F4", 61), c61, "$", float(c61.abs().sum()))], M)
    st.markdown("### 6. Net cash flows and running returns")
    c67 = (w.ocm_net + w.la_net - w.total_net).round(4)
    wide([(LBL("F4", 64), w.ocm_net, "$", float(w.ocm_net.sum())), (LBL("F4", 65), w.la_net, "$", float(w.la_net.sum())),
          (LBL("F4", 66), w.total_net, "$", float(w.total_net.sum())), (LBL("F4", 67), c67, "$", float(c67.abs().sum())),
          (LBL("F4", 68), w.ocm_moic_running, "r"), (LBL("F4", 69), w.la_moic_running, "r"),
          (LBL("F4", 70), w.ocm_irr_running, "%"), (LBL("F4", 71), w.la_irr_running, "%"),
          (LBL("F4", 73), w.debt_recap, "$", float(w.debt_recap.sum()))], M)
    st.markdown("### 7. Yearly view: equity in and distributions out by party (calendar years)")
    Y = sp["yearly"]
    groups = [("Total investors (OCM + Loan America)", ["Total: equity called", "Total: distributions", "Total: net cash flow", "Total: cumulative net"]),
              ("OCM", ["OCM: equity called", "OCM: distributions", "OCM: net cash flow", "OCM: cumulative net"]),
              ("Loan America (sponsor)", ["LA: equity called", "LA: distributions", "LA: of which promote above pro rata", "LA: net cash flow", "LA: cumulative net"])]
    items = []
    for head, cols in groups:
        items.append((head, [""] * len(Y), "p", ""))
        for c in cols:
            lbl = c.split(": ", 1)[1]
            lbl = {"equity called": "Equity called", "distributions": "Distributions", "net cash flow": "Net cash flow",
                   "cumulative net": "Cumulative net cash flow"}.get(lbl, LBL("F4", 91))
            tot = float(Y[c].values[-1]) if "cumulative" in c else float(Y[c].sum())
            items.append((f"{head.split(' (')[0]}: {lbl}", Y[c].values, "$", tot))
    tl = Y["Total: distributions"].sum()
    items.append((LBL("F4", 94), Y["LA share of distributions"].values, "%", float(Y["LA: distributions"].sum() / tl) if tl else 0.0))
    items.append((LBL("F4", 95), Y["Check: OCM + LA - total (0)"].values, "$", float(Y["Check: OCM + LA - total (0)"].abs().sum())))
    wide(items, Y.index, height=600)
    st.markdown("### 8. Distributions over time: OCM vs Loan America (pro rata vs promote)")
    wide([(LBL("F4", 99), w.ocm, "$", float(w.ocm.sum())), (LBL("F4", 100), w.la_pro_rata, "$", float(w.la_pro_rata.sum())),
          (LBL("F4", 101), w.la_promote, "$", float(w.la_promote.sum())), (LBL("F4", 102), w.la_share_month, "%"),
          (LBL("F4", 103), [("" if x == 0 else x) for x in w.highest_tier], "p"), (LBL("F4", 104), w.ocm_cum_dist, "$"),
          (LBL("F4", 105), w.la_cum_dist, "$")], M)
    st.markdown("### 9. Equity calls and distributions by party (SPV / OCM / Loan America)")
    c119 = ((-w.equity + w.ocm_equity + w.la_equity) + (w.dist - w.ocm - w.la_pro_rata - w.la_promote)).round(2)
    wide([(LBL("F4", 109), -w.equity, "$", -float(w.equity.sum())), (LBL("F4", 110), w.dist, "$", float(w.dist.sum())),
          (LBL("F4", 111), w.spv_cum_net, "$", float(w.spv_cum_net.values[-1])), (LBL("F4", 112), -w.ocm_equity, "$", -float(w.ocm_equity.sum())),
          (LBL("F4", 113), w.ocm, "$", float(w.ocm.sum())), (LBL("F4", 114), w.ocm_cum_net, "$", float(w.ocm_cum_net.values[-1])),
          (LBL("F4", 115), -w.la_equity, "$", -float(w.la_equity.sum())), (LBL("F4", 116), w.la_pro_rata, "$", float(w.la_pro_rata.sum())),
          (LBL("F4", 117), w.la_promote, "$", float(w.la_promote.sum())), (LBL("F4", 118), w.la_cum_net, "$", float(w.la_cum_net.values[-1])),
          (LBL("F4", 119), c119, "$", float(c119.abs().sum()))], M)
    note("The charts for sections 8 and 9 are on the Charts page, sections 4 and 5.")
    extras_header()
    last = int(np.max(np.nonzero(r["pool"]["monthly"].active.values)[0])) + 1
    x = w.index[:last]
    c1, c2 = st.columns(2)
    c1.plotly_chart(bar_fig(x, {"1 Return of capital": w.t1[:last], "2 Pref": w.t2[:last], "3 Promote": w.t3[:last], "4 Promote": w.t4[:last],
                                "5 Residual": w.t5[:last]}, "Monthly distributions by waterfall tier ($)",
                            colors=dict(zip(["1 Return of capital", "2 Pref", "3 Promote", "4 Promote", "5 Residual"], TIER_COLORS))), width="stretch")
    c2.plotly_chart(line_fig(x, {"Pref (investors)": w.hb2[:last], f"OCM to {fp(inp.hurdles[1],0)}": w.hb3[:last], f"OCM to {fp(inp.hurdles[2],0)}": w.hb4[:last]},
                             "Hurdle balances before distribution (below zero = hurdle met)", colors={"Pref (investors)": INK}), width="stretch")
    st.download_button("Download the waterfall table (CSV)", w.to_csv().encode(), file_name="F4_waterfall.csv")


# =========================================================================== F4a
def page_f4a():
    r = results(sens=True)
    st.title("Sensitivity: investor split")
    status_rail(r)
    lede("SPV, OCM and Loan America MOIC / IRR vs a single investor / school split. Each column re-runs the full SPV (other fees, "
         "leverage, recycling, equity / distributions and the waterfall). The middle column (0) is the current case.")
    kv([("Split step size (set on Assumptions, section 8)", [inp.split_step], "%"), ("Steps either side of the current investor %", [5], "i"),
        ("Check: single constant investor / school split (required)", ["OK" if r.get("single_split") else "CHECK"], "p")])
    if not r.get("single_split"):
        st.warning("This sensitivity needs one constant investor / school split, as in the workbook.")
        return
    s = r["split_sens"]; cols = [str(i) for i in s.columns]
    rows = [("Change in investor % (pp)", s.loc["Change (pp)"], "%"), ("Investor % of net collections", s.loc["Investor %"], "%"),
            ("School % of net collections", s.loc["School %"], "%")]
    for head, p_ in [("A. SPV (all investors)", "Total investors"), ("B. OCM", "OCM"), ("C. Loan America (sponsor)", "Loan America")]:
        short = head.split(". ")[1].split(" (")[0]
        rows += [(head, [""] * 11, "p"), (f"{short}: Equity called ($)", s.loc[f"{p_} Equity called"], "$"),
                 (f"{short}: Distributions ($)", s.loc[f"{p_} Distributions"], "$"), (f"{short}: MOIC", s.loc[f"{p_} MOIC"], "x"),
                 (f"{short}: IRR (annualised)", s.loc[f"{p_} IRR"], "%")]
    rows += [("Loan America share of total distributions", s.loc["LA share of distributions"], "%")]
    kv(rows, cols=cols, height=760)
    sc = r["sens_checks"]["F4a"]
    kv([("Check: middle column = SPV and waterfall page (SPV, OCM, LA MOIC and IRR, equity, distributions) (0)", [sc["middle = F4"]], "p"),
        ("Check: OCM + Loan America = SPV in every scenario (equity and distributions, 0)", [sc["OCM + LA = SPV"]], "p"),
        ("Overall split-sensitivity check", [sc["status"]], "p")])
    note("The two charts for this table are on the Charts page, section 7.")


# =========================================================================== F5
def page_f5():
    r = results()
    st.title("Program detail")
    status_rail(r)
    f5 = r["f5"]
    st.markdown("### 1. Deal returns by party (program level)")
    P = f5["parties"]
    kv([(P.index[0], P.iloc[0], "$"), (P.index[1], P.iloc[1], "$"), (P.index[2], P.iloc[2], "$"), (P.index[3], P.iloc[3], "x"),
        (P.index[4], P.iloc[4], "%"), (P.index[5], P.iloc[5], "%"),
        ("Loan America: origination / upfront fees received ($)", [None, f5["la_orig_fees"], None, None, None], "$")], cols=list(P.columns))
    st.markdown(f"**Key settings:** {f5['settings']}")
    note("School cash received = upfront advance + share of collections (it sells the loans rather than investing cash). "
         "Goal = across-life fees received.")
    st.markdown("### 2. School: effective advance rate (program level and per $1 by term)")
    Sch = f5["school"]
    kv([(Sch.index[0], Sch.iloc[0], "$"), (Sch.index[1], Sch.iloc[1], "%"), (Sch.index[2], Sch.iloc[2], "%"), (Sch.index[3], Sch.iloc[3], "%"),
        (Sch.index[4], Sch.iloc[4], "%"), (Sch.index[5], Sch.iloc[5], "%")], cols=list(Sch.columns))
    note("Program PV discounts each month's school cash to the program start (month 1 = first purchase month); per $1 PV discounts to the loan's month 0.")
    st.markdown("### 3. Investor per $1 of face by term (before SPV fees, leverage and recycling)")
    P1 = f5["per1"]
    kv([(P1.index[0], P1.iloc[0], "p"), (P1.index[1], P1.iloc[1], "p"), (P1.index[2], P1.iloc[2], "x"), (P1.index[3], P1.iloc[3], "%"),
        (P1.index[4], P1.iloc[4], "i")], cols=list(P1.columns))
    st.markdown("### 4. Monthly view (program level): cash by party and running returns")
    mv = f5["monthly"]
    kinds = ["$"] * 6 + ["r"] * 3 + ["%"] * 3 + ["%"]
    wide([("Month end", [dd.isoformat() for dd in r["pool"]["monthly"].date], "p")] +
         [(c, mv[c], k, float(mv[c].sum()) if k == "$" else None) for c, k in zip(mv.columns, kinds)], mv.index)
    note("Charts for sections 1-4 are on the Charts page, section 1.")
    st.markdown("### 5. Recycling: how each month's loan purchases are funded (collections vs debt vs new equity)")
    rc = f5["recycling"].iloc[:60]
    kinds = ["$", "$", "$", "$", "%", "$", "$", "$", "$"]
    s0 = rc.iloc[:, 0].sum()
    tots = [s0, rc.iloc[:, 1].sum(), rc.iloc[:, 2].sum(), rc.iloc[:, 3].sum(), rc.iloc[:, 1].sum() / s0 if s0 else 0,
            rc.iloc[:, 5].sum(), None, None, rc.iloc[:, 8].abs().sum()]
    wide([(c, rc[c], k, t) for c, k, t in zip(rc.columns, kinds, tots)], rc.index)
    kp = f5["recycling_kp"]
    kv([(k, [v], "%" if "÷" in k else "i") for k, v in kp.items()], cols=["Key points"])
    note("Recycling charts are on the Charts page, section 2.")
    st.markdown("### 6. Leverage: debt facility month by month (draws, interest, paydown) and its effect on equity and distributions")
    lv = f5["leverage"].iloc[:60]
    sums = ("Draws ($)", "Repayments ($)", "Interest paid ($)", "Equity called ($)", "Distributions to investors ($)")
    wide([(c, lv[c], "p" if c == "Facility mode" else "$",
           float(lv[c].sum()) if c in sums else (float(lv[c].abs().sum()) if c.startswith("Check") else None)) for c in lv.columns], lv.index)
    kp = f5["leverage_kp"]
    kv([("Leverage switch", [kp["Leverage switch"], ""], "p"), ("First draw month", [kp["First draw month"], ""], "i"),
        ("Peak debt ($) / month", [kp["Peak debt ($)"], kp["Peak debt month"]], "$"),
        ("Debt fully repaid by month", [kp["Debt fully repaid by month"], ""], "i"), ("Total interest paid ($)", [kp["Total interest paid ($)"], ""], "$"),
        ("Debt recap: distributions in the facility start month funded by the first draw ($)",
         [kp["Debt recap: start-month distributions funded by the first draw ($)"], ""], "$")], cols=["Key points", ""])
    note("Leverage charts are on the Charts page, section 3.")


# =========================================================================== F6
def page_f6():
    r = results(sens=True)
    st.title("Charts")
    lede("All Fusion charts in one place, live with every switch and input. Order: 1 returns overview, 2 funding and recycling, "
         "3 leverage, 4 equity vs distributions by party, 5 OCM / Loan America promote split, 6 advance-rate sensitivity, 7 split sensitivity.")
    st.markdown("## 1. Returns overview: how each party's money grows")
    st.plotly_chart(C.running_moic(r), width="stretch")
    st.plotly_chart(C.cash_investors_vs_school(r), width="stretch")
    st.plotly_chart(C.school_cumulative(r), width="stretch")
    st.markdown("## 2. Funding and recycling: who pays for each month's purchases")
    st.plotly_chart(C.funding_mix(r), width="stretch")
    st.plotly_chart(C.cum_equity_vs_purchases(r), width="stretch")
    st.markdown("## 3. Leverage: debt facility (blank when leverage is OFF)")
    st.plotly_chart(C.debt_facility(r), width="stretch")
    st.plotly_chart(C.leverage_investor_cash(r), width="stretch")
    st.markdown("## 4. Equity called vs distributions by party")
    lede("Equity below zero, distributions above; line = cumulative net cash (right axis).")
    for party in ["SPV", "OCM", "Loan America"]:
        st.plotly_chart(C.party_equity_vs_dist(r, party), width="stretch")
    st.markdown("## 5. Promote: how distributions split between OCM and Loan America over time")
    st.plotly_chart(C.promote_monthly(r), width="stretch")
    st.plotly_chart(C.promote_cumulative(r), width="stretch")
    st.markdown("## 6. Advance-rate sensitivity (per $1 of face, net of fees)")
    if r.get("single_split"):
        st.plotly_chart(C.adv_sens(r, "MOIC"), width="stretch")
        st.plotly_chart(C.adv_sens(r, "IRR"), width="stretch")
    else:
        st.warning("Needs one constant investor / school split, as in the workbook.")
    st.markdown("## 7. Split sensitivity: SPV / OCM / Loan America MOIC and IRR vs a single investor / school split")
    if r.get("single_split"):
        st.plotly_chart(C.split_sens(r, "MOIC"), width="stretch")
        st.plotly_chart(C.split_sens(r, "IRR"), width="stretch")
    else:
        st.warning("Needs one constant investor / school split, as in the workbook.")


# =========================================================================== extras: checks and compare
def page_checks():
    r = results(sens=True)
    st.title("Checks")
    status_rail(r)
    d = r["derived"]["checks"]
    rows = [("Assumptions", "Term mix sums to 100%", d["term_mix"]), ("Assumptions", "Effective loss curves never decrease", d["loss_curves"]),
            ("Assumptions", "Every active fee row valid; loan sizes positive", d["fee_table"]), ("Assumptions", "At least one valid split tier", d["target_tiers"]),
            ("Assumptions", "Purchase split sums to 100%, period 5 to 60", d["purchase_split"]), ("Assumptions", "Purchase cash adds to the program", d["purchase_cash"]),
            ("Assumptions", "Ramp-up valid", d["ramp_up"]), ("Assumptions", "Last draw month on or after the start month", d["leverage_months"]),
            ("Assumptions", "Last draw month covers the purchase period (warning)", d["last_draw_warning"]),
            ("Assumptions", "All input checks", r["input_status"]), ("Assumptions", "All calculation checks", r["model_status"])]
    for k, v in r["ue"]["checks"].items():
        rows.append(("Loan economics", f"{k}: split and balance identities", v))
    rows.append(("Gross vs net", "Gross investor + school = gross collections", r["asset"]["check"]))
    pc = r["pool"]["checks"]
    rows += [("Pool", "Identities", "OK" if pc["pool_identities"] == 0 else "CHECK"),
             ("Pool", "Debt fully repaid", "OK" if abs(pc["debt_repaid"]) < 1 else "CHECK"),
             ("Pool", "/ B90 tie to unit economics", "OK" if pc["tie_to_F2"] == 0 else "CHECK"),
             ("Pool", "Vintage grids tie", "OK" if pc["vintage_ties"] < 1 else "CHECK"), ("Pool", "Overall", pc["status"]),
             ("SPV and waterfall", "Waterfall ties, OCM + LA = total", r["spv"]["checks"]["waterfall"]),
             ("SPV and waterfall", "Promote shares 0-100%", r["spv"]["checks"]["promote_shares"]),
             ("SPV and waterfall", "Yearly OCM + LA = total", "OK" if r["spv"]["yearly"]["Check: OCM + LA - total (0)"].abs().sum() == 0 else "CHECK")]
    f5 = r["f5"]
    rows += [("Program detail", "Funding check", "OK" if f5["recycling"].iloc[:, 8].abs().sum() < 1 else "CHECK"),
             ("Program detail", "Debt roll check", "OK" if f5["leverage"].iloc[:, -1].abs().sum() < 1 else "CHECK")]
    for k, v in (r.get("sens_checks") or {}).items():
        rows.append(({"F2b": "Sensitivity: advance rate", "F4a": "Sensitivity: investor split"}.get(k, k), "Middle column ties to the live model", v["status"]))
    df = pd.DataFrame(rows, columns=["Area", "Check", "Status"])
    st.dataframe(df.style.map(lambda v: f"color: {status_color(v)}; font-weight: 600", subset=["Status"]), width="stretch", hide_index=True,
                 height=38 * len(df) + 40)
    note("The Python engine was tied to the workbook cell by cell (about 38,000 cells per workbook across every page) on the "
         "base case and five stress scenarios. Run tests/tie_out.py against any saved copy of the Excel model to repeat it.")


def page_compare():
    st.title("Compare scenarios")
    lede("Save scenarios from the sidebar, then compare them side by side with the current inputs.")
    pool_ = {"Current": inp.to_json(), **st.session_state.saved}
    pick = st.multiselect("Scenarios", list(pool_), default=list(pool_)[:4])
    out = {}
    for n in pick:
        r = _run(pool_[n], False)
        R, PS, S = r["spv"]["returns"], r["pool"]["summary"], r["ue"]["summary"]
        out[n] = {"Advance (blended)": fp(r["derived"]["F71"]), "Investor split": fp(r["derived"]["tiers"][0][0]) if r["derived"]["tiers"] else "n.a.",
                  "Loss multiplier": f"{r['inputs'].loss_mult:.2f}x", "Program": fm(r["inputs"].purchase_amount),
                  "Net MOIC per $1": fx(S.loc["Investor MOIC", "Blended"]), "Net IRR per $1": fp(S.loc["Investor IRR", "Blended"]),
                  "School total, PV": fp(S.loc["School total (PV)", "Blended"]), "Equity called": fm(PS["Equity called"]),
                  "Distributions": fm(PS["Distributions"]), "SPV MOIC": fx(R.loc["MOIC", "Total investors"]), "SPV IRR": fp(R.loc["IRR", "Total investors"]),
                  "OCM MOIC": fx(R.loc["MOIC", "OCM"]), "OCM IRR": fp(R.loc["IRR", "OCM"]), "Loan America MOIC": fx(R.loc["MOIC", "Loan America"]),
                  "Loan America IRR": fp(R.loc["IRR", "Loan America"]), "Peak debt": fm(PS["Peak debt"]),
                  "Checks": f"{r['input_status']} / {r['model_status']}"}
    if out:
        st.dataframe(pd.DataFrame(out), width="stretch", height=660)


# =========================================================================== navigation (workbook order, then extras)
sidebar()
pg = st.navigation({
    "Deal model": [
        st.Page(page_f0, title="Overview", url_path="overview", default=True),
        st.Page(page_f1, title="Assumptions", url_path="assumptions"),
        st.Page(page_f2, title="Loan economics", url_path="loan-economics"),
        st.Page(page_f2a, title="Gross vs net returns", url_path="gross-vs-net"),
        st.Page(page_f3, title="Pool cash flows", url_path="pool"),
        st.Page(page_f4, title="SPV and waterfall", url_path="spv"),
        st.Page(page_f5, title="Program detail", url_path="program"),
        st.Page(page_f6, title="Charts", url_path="charts"),
    ],
    "Sensitivities": [
        st.Page(page_f2b, title="Advance rate", url_path="sensitivity-advance"),
        st.Page(page_f4a, title="Investor split", url_path="sensitivity-split"),
    ],
    "Tools": [
        st.Page(page_checks, title="Checks", url_path="checks"),
        st.Page(page_compare, title="Compare scenarios", url_path="compare"),
        st.Page(page_sheet3, title="Scenario notes", url_path="scenario-notes"),
        st.Page(page_notes, title="Model notes", url_path="notes"),
    ],
}, expanded=True)
pg.run()
